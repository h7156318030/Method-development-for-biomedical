"""Regression tests for the backend hardening list (docs/后端加固清单.md).

Every finding of the independent review has at least one test here that REPRODUCES the reviewer's scenario
(the scenario scripts are not part of the repository) and asserts the hardened behaviour.  Test names carry the
URS id (URS-80 ... URS-97, see docs/URS_需求与测试追溯.md).  Run:  python -m unittest discover -s tests -v
"""
import contextlib
import io
import json
import os
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import zipfile
from unittest import mock

import test_workflow as tw                       # sets METHODFLOW_DATA and imports the app (shared helpers)
from test_workflow import PW, client, fill_required, login, make_project, project_id_of, sign, uid, approve_stage  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402

from app import audit, auth, config, db, docgen, main, workflow  # noqa: E402

HTMP = tempfile.mkdtemp(prefix="mf_hard_")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IDS = {}
USERS = [("alice", ["author"]), ("bob", ["author"]), ("carol", ["author"]), ("dora", ["author", "reviewer"]),
         ("rita", ["reviewer"]), ("rob", ["reviewer"]), ("quinn", ["approver"]), ("vic", ["viewer"]),
         ("multi", ["reviewer", "approver"])]


def setUpModule():
    config.set_data(HTMP)
    config.override("pdf_enabled", False)      # PDF output has its own tests (test_docgen.py); here it would only slow things down
    auth.reset_limits()
    db.init_db()
    main._seed_kb()
    c = client()
    r = c.post("/api/setup", json={"username": "admin", "display_name": "管理员", "password": PW})
    assert r.status_code == 200, r.text
    for u, roles in USERS:
        r = c.post("/api/users", json={"username": u, "display_name": u.capitalize(), "password": PW, "roles": roles})
        assert r.status_code == 200, r.text
    with db.tx() as conn:
        conn.execute("UPDATE users SET must_change=0")
    for u in c.get("/api/users").json():
        IDS[u["username"]] = u["id"]


def tearDownModule():
    shutil.rmtree(HTMP, ignore_errors=True)


# ====================================================================== helpers
@contextlib.contextmanager
def cfg(**kw):
    old = {k: config.get(k) for k in kw}
    for k, v in kw.items():
        config.override(k, v)
    try:
        yield
    finally:
        for k, v in old.items():
            config.override(k, v)


def rclient(username=None):
    """Client that turns server exceptions into 500 responses instead of re-raising them."""
    c = TestClient(main.app, raise_server_exceptions=False)
    c.headers.update({"X-Requested-With": "MethodFlow"})
    if username:
        r = c.post("/api/login", json={"username": username, "password": PW})
        assert r.status_code == 200, r.text
    return c


def raw(sql, args=(), drop_triggers=False):
    """Direct SQL on the live database - what somebody with file access could do."""
    con = sqlite3.connect(str(config.DB_PATH), timeout=15)
    con.row_factory = sqlite3.Row
    try:
        if drop_triggers:
            for n in db.TRIGGER_NAMES:
                con.execute("DROP TRIGGER IF EXISTS %s" % n)
        cur = con.execute(sql, args)
        rows = cur.fetchall()
        con.commit()
        return rows
    finally:
        con.close()
        if drop_triggers:
            db.init_db()            # CREATE TRIGGER IF NOT EXISTS: puts the protection back for the other tests


def stage(c, mid, key):
    r = c.get("/api/methods/%d/stages/%s" % (mid, key))
    assert r.status_code == 200, r.text
    return r.json()


def exps(c, mid):
    return c.get("/api/methods/%d/experiments" % mid).json()


def exp_put(c, eid, **kw):
    cur = c.get("/api/experiments/%d" % eid).json()
    return c.put("/api/experiments/%d" % eid, json=dict(version=cur["version"], **kw))


def members_of(pid_roles):
    return [{"user_id": IDS[n], "role": r} for n, r in pid_roles]


BASE = [("alice", "owner"), ("alice", "author"), ("rita", "reviewer"), ("quinn", "approver"), ("rob", "reviewer")]


def approve_through_plan(a, mid):
    for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
        approve_stage(mid, k, a, with_approver=w)


def lock_explore(a, mid):
    """Close every explore experiment and submit the explore stage (-> in_review)."""
    for e in exps(a, mid):
        if e["stage_key"] == "explore" and e["status"] in ("planned", "running"):
            r = a.post("/api/experiments/%d/status" % e["id"], json={"status": "cancelled", "reason": "测试取消"})
            assert r.status_code == 200, r.text
    fill_required(a, mid, "explore")
    r = a.post("/api/methods/%d/stages/explore/submit" % mid, json={})
    assert r.status_code == 200, r.text


def docx_text(c, did):
    from docx import Document
    d = Document(io.BytesIO(c.get("/api/docs/%d/download" % did).content))
    return "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(c_.text for t in d.tables for row in t.rows for c_ in row.cells)


def upload(c, mid, name, data=b"x", **params):
    p = dict(params)
    p["filename"] = name
    return c.post("/api/methods/%d/attachments" % mid, params=p, content=data)


def copy_db():
    """Private copy of the live database (consistent snapshot, WAL merged) for destructive audit tests."""
    dst = os.path.join(HTMP, "copy_%d.db" % int(time.time() * 1000))
    s = sqlite3.connect(str(config.DB_PATH))
    d = sqlite3.connect(dst)
    s.backup(d)
    s.close()
    d.row_factory = sqlite3.Row
    return d


# ====================================================================== H1  separation of duties
class TestSeparationOfDuties(unittest.TestCase):
    def test_URS80_project_roles_must_match_global_roles_and_not_mix(self):
        """t1 scenario A: an owner could name any account (a viewer, another author) as reviewer / approver."""
        a = login("alice")
        pid = a.post("/api/projects", json={"name": "P80", "members": [{"user_id": i, "role": "member"} for n, i in IDS.items() if n not in ("admin", "alice")]}).json()["id"]

        def create(members):
            return a.post("/api/projects/%d/methods" % pid, json={"name": "M", "template_key": "flow", "members": members})
        # roles that the account does not hold globally
        self.assertEqual(create(members_of([("vic", "approver")])).status_code, 400)
        self.assertEqual(create(members_of([("bob", "reviewer")])).status_code, 400)
        self.assertEqual(create(members_of([("bob", "approver")])).status_code, 400)
        self.assertEqual(create(members_of([("rita", "author")])).status_code, 400)
        self.assertEqual(create(members_of([("rita", "owner")])).status_code, 400)
        # the creator (owner + author) cannot also sign
        self.assertEqual(create(members_of([("alice", "reviewer")])).status_code, 400)
        # an account that holds both global roles still cannot hold an editing AND a signing role
        self.assertEqual(create(members_of([("dora", "author"), ("dora", "reviewer")])).status_code, 400)
        # reviewer and approver may not be one person
        self.assertEqual(create(members_of([("multi", "reviewer"), ("multi", "approver")])).status_code, 400)
        # malformed
        for bad in ("abc", [1], [{"role": "reviewer"}], [{"user_id": "x", "role": "reviewer"}], [{"user_id": IDS["rita"], "role": "king"}],
                    [{"user_id": 99999, "role": "viewer"}], [{"user_id": True, "role": "viewer"}]):
            self.assertEqual(create(bad).status_code, 400, bad)
        # a valid method team
        self.assertEqual(create(members_of([("dora", "reviewer"), ("quinn", "approver"), ("vic", "viewer")])).status_code, 200)

    def test_URS80_set_members_same_rules_and_frozen_while_in_review(self):
        """t1 scenario B: after somebody else did the work the owner made himself reviewer; and members of an
        in-review stage could be swapped."""
        a, mid, ids = make_project()
        url = "/api/methods/%d/members" % mid
        self.assertEqual(a.put(url, json={"members": members_of(BASE + [("alice", "reviewer")])}).status_code, 400)
        self.assertEqual(a.put(url, json={"members": members_of(BASE + [("vic", "approver")])}).status_code, 400)
        self.assertEqual(a.put(url, json={"members": members_of(BASE + [("dora", "author"), ("dora", "reviewer")])}).status_code, 400)
        self.assertEqual(a.put(url, json={"members": "abc"}).status_code, 400)
        self.assertEqual(a.put(url, json={"members": members_of([("rita", "reviewer")])}).status_code, 400)   # no owner
        # with an in-review stage the signing roles are frozen
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        self.assertEqual(a.post("/api/methods/%d/stages/atp/submit" % mid, json={}).status_code, 200)
        self.assertEqual(a.put(url, json={"members": members_of([m for m in BASE if m[0] != "rob"])}).status_code, 409)       # remove a reviewer
        self.assertEqual(a.put(url, json={"members": members_of(BASE + [("dora", "reviewer")])}).status_code, 409)             # add a reviewer
        self.assertEqual(a.put(url, json={"members": members_of([m for m in BASE if m[0] != "quinn"])}).status_code, 409)      # remove approver
        # editing-side changes are still possible
        self.assertEqual(a.put(url, json={"members": members_of(BASE + [("carol", "author"), ("vic", "viewer")])}).status_code, 200)

    def test_URS81_contributors_are_cumulative_across_withdraw_and_role_changes(self):
        """t1 scenario C: withdraw reset `editors`, so the author could later become reviewer of his own work."""
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        approve_stage(mid, "charter", a)
        self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE + [("dora", "author")])}).status_code, 200)
        fill_required(a, mid, "atp")
        dora = login("dora")
        ver = stage(a, mid, "atp")["version"]
        self.assertEqual(dora.put(url + "/stages/atp", json={"version": ver, "sections": {"specificity": "dora 写的"}}).status_code, 200)
        st = stage(a, mid, "atp")
        self.assertIn(IDS["dora"], st["editors"])
        self.assertIn(IDS["dora"], st["contributors"])
        # dora is moved from author to reviewer: nothing but the history stops her from signing her own work
        self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE + [("dora", "reviewer")])}).status_code, 200)
        self.assertEqual(a.post(url + "/stages/atp/submit", json={}).status_code, 200)
        d = stage(dora, mid, "atp")
        self.assertIsNone(d["can_sign_as"])
        self.assertIn("职责分离", d["cannot_sign_reason"])
        self.assertEqual(sign(dora, mid, "atp").status_code, 403)
        # withdraw empties `editors` (current revision) but not `contributors`
        self.assertEqual(a.post(url + "/stages/atp/withdraw", json={"reason": "改一下"}).status_code, 200)
        st = stage(a, mid, "atp")
        self.assertEqual(st["editors"], [])
        self.assertIn(IDS["dora"], st["contributors"])
        self.assertIn("Dora", st["contributors_names"])
        self.assertEqual(a.post(url + "/stages/atp/submit", json={}).status_code, 200)
        self.assertEqual(sign(dora, mid, "atp").status_code, 403)
        # a clean reviewer can sign; a reject also keeps the contributors
        self.assertEqual(sign(login("rita"), mid, "atp", "reject", "退回").status_code, 200)
        st = stage(a, mid, "atp")
        self.assertEqual(st["status"], "returned")
        self.assertIn(IDS["dora"], st["contributors"])

    def test_URS81_submitter_who_never_saved_is_a_contributor(self):
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        approve_stage(mid, "charter", a)
        self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE + [("dora", "author")])}).status_code, 200)
        fill_required(a, mid, "atp")                                   # alice did all the work
        dora = login("dora")
        self.assertEqual(dora.post(url + "/stages/atp/submit", json={}).status_code, 200)   # dora only submitted
        st = stage(a, mid, "atp")
        self.assertNotIn(IDS["dora"], st["editors"])
        self.assertIn(IDS["dora"], st["contributors"])
        self.assertEqual(a.post(url + "/stages/atp/withdraw", json={"reason": "x"}).status_code, 200)
        self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE + [("dora", "reviewer")])}).status_code, 200)
        self.assertEqual(a.post(url + "/stages/atp/submit", json={}).status_code, 200)
        self.assertEqual(sign(dora, mid, "atp").status_code, 403)

    def test_URS81_experiment_and_experiment_attachment_work_counts_for_the_stage(self):
        """t2 F: an attachment tied to an experiment (no stage_key) used to leave the stage's editors untouched."""
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE + [("dora", "author"), ("carol", "author")])}).status_code, 200)
        ex = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        dora, carol = login("dora"), login("carol")
        self.assertEqual(exp_put(dora, ex["id"], conclusion="dora 的结论").status_code, 200)
        r = upload(carol, mid, "raw.fcs", b"abc", experiment_id=ex["id"])
        self.assertEqual(r.status_code, 200, r.text)
        st = stage(a, mid, "explore")
        self.assertIn(IDS["dora"], st["contributors"])
        self.assertIn(IDS["carol"], st["contributors"])
        self.assertIn(IDS["carol"], st["editors"])


# ====================================================================== H2  experiments / attachments vs. locks
class TestExperimentLocks(unittest.TestCase):
    def test_URS82_stage_key_cannot_be_cleared_or_moved_into_locked_stage(self):
        """t2 D/E: detaching an experiment (stage_key=None) removed the submit blocker; moving an experiment
        into an approved / in-review stage rewrote that stage's editors."""
        a, mid, ids = make_project()
        # before the plan is approved an experiment may move between two editable stages ...
        ex = [e for e in exps(a, mid) if e["stage_key"] == "explore"]
        self.assertEqual(exp_put(a, ex[0]["id"], stage_key="robust").status_code, 200)
        self.assertEqual(exp_put(a, ex[0]["id"], stage_key="explore").status_code, 200)
        self.assertEqual(exp_put(a, ex[0]["id"], stage_key="nonexistent").status_code, 404)
        self.assertEqual(exp_put(a, ex[0]["id"], stage_key=["x"]).status_code, 400)
        approve_through_plan(a, mid)
        # ... but never to "no stage"
        for e in exps(a, mid):
            if e["stage_key"] == "explore":
                r = exp_put(a, e["id"], stage_key=None)
                self.assertEqual(r.status_code, 409, r.text)
        fill_required(a, mid, "explore")
        r = a.post("/api/methods/%d/stages/explore/submit" % mid, json={})
        self.assertEqual(r.status_code, 422)                      # the blocker is still there
        self.assertTrue(any("尚未完成或取消" in b for b in r.json()["detail"]["blocks"]))
        # an experiment of the signed plan can't be re-assigned any more; one added later can
        plan_scope = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        self.assertEqual(exp_put(a, plan_scope["id"], stage_key="robust").status_code, 409)
        late = a.post("/api/methods/%d/experiments" % mid, json={"title": "late", "stage_key": "robust", "design": "d", "acceptance": "a"}).json()["id"]
        lock_explore(a, mid)
        before = stage(a, mid, "plan")
        for target in ("explore", "plan", "charter"):               # in_review / approved / completed
            r = exp_put(a, late, stage_key=target)
            self.assertEqual(r.status_code, 409, (target, r.text))
        after = stage(a, mid, "plan")
        self.assertEqual((before["editors"], before["contributors"]), (after["editors"], after["contributors"]))
        self.assertEqual(before["hash"], after["hash"])

    def test_URS82_delete_attachment_checks_the_experiment_lock(self):
        """t2 F: DELETE of an attachment tied only to an experiment of a locked stage succeeded."""
        a, mid, ids = make_project()
        ex = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        r = upload(a, mid, "raw.fcs", b"abc" * 100, experiment_id=ex["id"])
        self.assertEqual(r.status_code, 200, r.text)
        aid = r.json()["id"]
        approve_through_plan(a, mid)
        lock_explore(a, mid)
        self.assertTrue(a.get("/api/experiments/%d" % ex["id"]).json()["locked"])
        self.assertEqual(upload(a, mid, "late.fcs", b"x", experiment_id=ex["id"]).status_code, 409)
        r = a.delete("/api/attachments/%d?reason=because" % aid)
        self.assertEqual(r.status_code, 409, r.text)
        self.assertEqual(len(a.get("/api/experiments/%d" % ex["id"]).json()["attachments_list"]), 1)


# ====================================================================== H3  what a signature covers
class TestSignatureCoverage(unittest.TestCase):
    def test_URS83_plan_experiments_frozen_after_approval_and_late_ones_flagged(self):
        """t2 G: the acceptance criteria of an approved plan could be rewritten; new experiments changed nothing
        visible.  Now the plan's experiments are frozen (409) and later additions are listed apart."""
        a, mid, ids = make_project()
        approve_through_plan(a, mid)
        plan0 = stage(a, mid, "plan")
        self.assertTrue(plan0["signatures_ok"])
        ex = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        for field in ("acceptance", "design", "objective", "title"):
            r = exp_put(a, ex["id"], **{field: "anything goes now"})
            self.assertEqual(r.status_code, 409, (field, r.text))
            self.assertIn("重新打开", r.json()["error"])
        self.assertEqual(exp_put(a, ex["id"], conclusion="可以补结论").status_code, 200)       # not part of the plan hash
        # a new experiment is allowed, flagged, and does not disturb the plan signature
        r = a.post("/api/methods/%d/experiments" % mid, json={"title": "追加实验", "stage_key": "explore", "design": "d", "acceptance": "a"})
        self.assertEqual(r.status_code, 200)
        late = r.json()["id"]
        self.assertEqual(a.get("/api/experiments/%d" % late).json()["added_after_plan"], 1)
        self.assertEqual(exps(a, mid)[0]["added_after_plan"], 0)
        self.assertEqual(exp_put(a, late, design="late edit").status_code, 200)               # free to change until its own stage is signed
        r = upload(a, mid, "run.fcs", b"1234", experiment_id=ex["id"])                        # raw data arrives while running
        self.assertEqual(r.status_code, 200, r.text)
        plan1 = stage(a, mid, "plan")
        self.assertEqual(plan1["hash"], plan0["hash"])
        self.assertTrue(plan1["signatures_ok"])
        self.assertTrue(all(x["intact"] for x in plan1["reviews"]))
        # the Word plan lists the late experiment separately
        did = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).json()["id"]
        self.assertIn("方案外追加", docx_text(a, did))
        # plan-scope experiments can't be deleted, late ones can
        self.assertEqual(a.delete("/api/experiments/%d?reason=x" % ex["id"]).status_code, 409)
        self.assertEqual(a.delete("/api/experiments/%d?reason=不需要" % late).status_code, 200)
        self.assertEqual(stage(a, mid, "plan")["hash"], plan0["hash"])
        # after re-opening the plan the experiments can be edited again
        r = a.post("/api/methods/%d/stages/plan/reopen" % mid, json={"reason": "改方案", "password": PW})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(exp_put(a, ex["id"], acceptance="新的接受标准 ≥ 90%").status_code, 200)

    def test_URS83_experiment_added_after_the_plan_needs_design_and_acceptance(self):
        """Testing found that an experiment added to an approved plan without design / acceptance criterion made the
        plan stage unable to be recertified (blocker '缺少设计或接受标准'); creation now refuses it up front (400)."""
        a, mid, ids = make_project()
        # plan still open: a bare title is fine (the plan gate will ask for the details later)
        r = a.post("/api/methods/%d/experiments" % mid, json={"title": "草稿实验", "stage_key": "explore"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(exp_put(a, r.json()["id"], design="d", acceptance="a").status_code, 200)   # needed for the plan gate itself
        for k in ("charter", "atp", "tech", "risk"):
            approve_stage(mid, k, a, with_approver=False)
        fill_required(a, mid, "plan")
        r = a.post("/api/methods/%d/stages/plan/submit" % mid, json={})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(stage(a, mid, "plan")["status"], "in_review")

        def create(**kw):
            body = {"title": "追加", "stage_key": "explore"}
            body.update(kw)
            return a.post("/api/methods/%d/experiments" % mid, json=body)

        n0 = len(exps(a, mid))
        for kw in ({}, {"design": "d"}, {"acceptance": "a"}, {"design": "   ", "acceptance": "a"}, {"design": "d", "acceptance": "\n "}):
            r = create(**kw)
            self.assertEqual(r.status_code, 400, (kw, r.text))
            self.assertIn("设计", r.json()["error"])
            self.assertIn("接受标准", r.json()["error"])
        self.assertEqual(len(exps(a, mid)), n0)                       # nothing was created
        self.assertEqual(create(design="d", acceptance="a").status_code, 200)
        # the same once the plan is approved
        a.post("/api/methods/%d/stages/plan/withdraw" % mid, json={"reason": "继续审批"})
        approve_stage(mid, "plan", a, with_approver=True)
        self.assertEqual(stage(a, mid, "plan")["status"], "approved")
        self.assertEqual(create().status_code, 400)
        r = create(design="d", acceptance="a")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(a.get("/api/experiments/%d" % r.json()["id"]).json()["added_after_plan"], 1)

    def test_URS83_signed_hash_covers_experiments_and_attachments(self):
        """A change made behind the application's back to an experiment / attachment that a signature vouches
        for is now visible: `intact` turns false (it used to stay true because only stage text was hashed)."""
        a, mid, ids = make_project()
        ex = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        self.assertEqual(upload(a, mid, "raw.fcs", b"abc" * 10, experiment_id=ex["id"]).status_code, 200)
        approve_through_plan(a, mid)
        plan = stage(a, mid, "plan")
        lock_explore(a, mid)
        self.assertEqual(sign(login("rita"), mid, "explore").status_code, 200)
        explore = stage(a, mid, "explore")
        self.assertEqual(explore["status"], "approved")
        self.assertTrue(explore["signatures_ok"] and all(r["intact"] for r in explore["reviews"]))
        # (1) plan: an experiment's acceptance criteria rewritten in the database
        old = raw("SELECT acceptance FROM experiments WHERE id=?", (ex["id"],))[0][0]
        try:
            raw("UPDATE experiments SET acceptance=? WHERE id=?", ("tampered", ex["id"]))
            p2 = stage(a, mid, "plan")
            self.assertFalse(p2["signatures_ok"])
            self.assertFalse(all(r["intact"] for r in p2["reviews"]))
        finally:
            raw("UPDATE experiments SET acceptance=? WHERE id=?", (old, ex["id"]))
        self.assertTrue(stage(a, mid, "plan")["signatures_ok"])
        # (2) explore: conclusion / results of an experiment
        exp2 = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        old = raw("SELECT conclusion FROM experiments WHERE id=?", (exp2["id"],))[0][0]
        try:
            raw("UPDATE experiments SET conclusion=? WHERE id=?", ("改过的结论", exp2["id"]))
            e2 = stage(a, mid, "explore")
            self.assertFalse(e2["signatures_ok"])
            self.assertFalse(any(r["intact"] for r in e2["reviews"]))
        finally:
            raw("UPDATE experiments SET conclusion=? WHERE id=?", (old, exp2["id"]))
        # (3) explore: an attachment's recorded hash
        old = raw("SELECT sha256 FROM attachments WHERE experiment_id=?", (ex["id"],))[0][0]
        try:
            raw("UPDATE attachments SET sha256=? WHERE experiment_id=?", ("0" * 64, ex["id"]))
            self.assertFalse(stage(a, mid, "explore")["signatures_ok"])
        finally:
            raw("UPDATE attachments SET sha256=? WHERE experiment_id=?", (old, ex["id"]))
        self.assertTrue(stage(a, mid, "explore")["signatures_ok"])
        self.assertEqual(plan["hash"], stage(a, mid, "plan")["hash"])

    def test_URS84_method_metadata_frozen_after_approval_and_covered_by_hash(self):
        """t6: title / product / attribute / intended_use could be edited after the stages were approved.  Now:
        the method's name / analyte / matrix / intended_use and the product of its project."""
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        purl = "/api/projects/%d" % project_id_of(a, mid)
        h0 = stage(a, mid, "charter")["hash"]
        self.assertEqual(a.patch(url, json={"analyte": "新分析物", "intended_use": "放行"}).status_code, 200)   # nothing approved yet
        h1 = stage(a, mid, "charter")["hash"]
        self.assertNotEqual(h1, h0)                                                                          # metadata is part of the charter hash
        self.assertEqual(a.patch(purl, json={"product": "新产品"}).status_code, 200)
        h2 = stage(a, mid, "charter")["hash"]
        self.assertNotEqual(h2, h1)                                                                          # ... and so is the project's product
        self.assertEqual(a.patch(url, json={"phase": "II 期", "method_version": "V0.2"}).status_code, 200)    # not covered -> free
        self.assertEqual(a.patch(purl, json={"indication": "B-ALL", "phase": "I 期", "name": "改名的项目"}).status_code, 200)
        self.assertEqual(stage(a, mid, "charter")["hash"], h2)
        approve_stage(mid, "charter", a)
        for body in ({"name": "改名"}, {"analyte": "x"}, {"matrix": "血浆"}, {"intended_use": "changed after approval"}):
            r = a.patch(url, json=body)
            self.assertEqual(r.status_code, 409, (body, r.text))
        r = a.patch(purl, json={"product": "其他"})
        self.assertEqual(r.status_code, 409, r.text)                                                         # the project's product is frozen too
        self.assertEqual(a.patch(url, json={"phase": "III 期", "analyte": "新分析物"}).status_code, 200)       # unchanged values are fine
        self.assertEqual(a.patch(purl, json={"product": "新产品", "indication": "DLBCL"}).status_code, 200)
        self.assertEqual(a.get(url).json()["phase"], "III 期")
        # direct database edit after the plan was approved is detected by the plan signature
        approve_through_plan_rest = [("atp", False), ("tech", False), ("risk", False), ("plan", True)]
        for k, w in approve_through_plan_rest:
            approve_stage(mid, k, a, with_approver=w)
        self.assertTrue(stage(a, mid, "plan")["signatures_ok"])
        pid = project_id_of(a, mid)
        raw("UPDATE projects SET product='sneaky' WHERE id=?", (pid,))
        try:
            self.assertFalse(stage(a, mid, "plan")["signatures_ok"])
            self.assertFalse(stage(a, mid, "charter")["signatures_ok"])
        finally:
            raw("UPDATE projects SET product='新产品' WHERE id=?", (pid,))
        self.assertTrue(stage(a, mid, "plan")["signatures_ok"])
        raw("UPDATE methods SET matrix='sneaky' WHERE id=?", (mid,))
        try:
            self.assertFalse(stage(a, mid, "charter")["signatures_ok"])
        finally:
            raw("UPDATE methods SET matrix='全血' WHERE id=?", (mid,))
        self.assertTrue(stage(a, mid, "charter")["signatures_ok"])


# ====================================================================== H4 / H5
class TestTamperDetection(unittest.TestCase):
    def test_URS85_intact_flag_documents_and_audit_verify_detect_tampering(self):
        """t7: a raw database edit of an approved stage left every `intact` flag true, the Word file said
        '已批准版本' and /api/audit/verify said ok."""
        a, mid, ids = make_project()
        admin = login("admin")
        approve_through_plan(a, mid)
        st = stage(a, mid, "atp")
        self.assertTrue(st["signatures_ok"] and all(r["intact"] for r in st["reviews"]))
        did = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).json()["id"]
        self.assertEqual(a.get("/api/methods/%d/docs" % mid).json()[0]["status_label"], "已批准版本")
        self.assertIn("状态：已批准版本", docx_text(a, did))
        v = admin.get("/api/audit/verify").json()
        self.assertTrue(v["signatures_ok"])
        original = raw("SELECT data FROM stages WHERE method_id=? AND key='atp'", (mid,))[0][0]
        d = json.loads(original)
        d["sections"]["use"] = "TAMPERED DIRECTLY IN DB"
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='atp'", (json.dumps(d, ensure_ascii=False), mid))
        try:
            st = stage(a, mid, "atp")
            self.assertEqual(st["status"], "approved")
            self.assertFalse(any(r["intact"] for r in st["reviews"]))
            self.assertFalse(st["signatures_ok"])
            self.assertNotEqual(st["hash"], st["reviews"][-1]["data_hash"])
            proj = a.get("/api/methods/%d" % mid).json()
            self.assertFalse(proj["signatures_ok"])
            self.assertEqual(proj["bad_signatures"][0]["stage"], "atp")
            self.assertFalse(a.get("/api/methods/%d/export" % mid).json()["signatures_ok"])
            # the plan covers the ATP: its signature is no longer valid either
            self.assertFalse(stage(a, mid, "plan")["signatures_ok"])
            # document generation refuses to call this an approved version
            r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
            self.assertEqual(r.status_code, 200, r.text)
            did2 = r.json()["id"]
            lst = a.get("/api/methods/%d/docs" % mid).json()
            self.assertEqual(lst[0]["status_label"], "完整性校验失败")
            self.assertFalse(lst[0]["signatures_ok"])
            self.assertFalse(lst[0]["generated_signatures_ok"])
            self.assertTrue(lst[1]["generated_signatures_ok"])           # the earlier file was generated while intact
            text = docx_text(a, did2)
            self.assertIn("完整性校验失败", text)
            self.assertNotIn("状态：已批准版本", text)
            # the audit verification reports it
            v = admin.get("/api/audit/verify").json()
            self.assertFalse(v["ok"])
            self.assertFalse(v["signatures_ok"])
            self.assertTrue(v["chain_ok"])                               # the audit log itself is untouched
            self.assertIn("atp", [b["stage"] for b in v["bad_signatures"] if b["method_id"] == mid])
        finally:
            raw("UPDATE stages SET data=? WHERE method_id=? AND key='atp'", (original, mid))
        self.assertTrue(stage(a, mid, "atp")["signatures_ok"])
        self.assertTrue(admin.get("/api/audit/verify").json()["signatures_ok"])

    def test_URS85_history_is_checked_against_the_snapshot_table(self):
        a, mid, ids = make_project()
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        a.post("/api/methods/%d/stages/atp/submit" % mid, json={})
        self.assertEqual(sign(login("rita"), mid, "atp", "reject", "判据缺依据").status_code, 200)
        st = stage(a, mid, "atp")
        self.assertEqual((st["status"], st["revision"]), ("returned", 2))
        self.assertTrue(st["reviews"][0]["intact"])
        self.assertTrue(st["signatures_ok"])
        # the old revision's snapshot hash is rewritten (the append-only trigger has to be dropped first)
        rows = raw("SELECT id, data_hash FROM stage_revisions WHERE revision=1 AND stage_id=?", (st["id"],))
        try:
            raw("UPDATE stage_revisions SET data_hash=? WHERE id=?", ("f" * 64, rows[0]["id"]), drop_triggers=True)
            st = stage(a, mid, "atp")
            self.assertFalse(st["reviews"][0]["intact"])
            self.assertFalse(st["signatures_ok"])
        finally:
            raw("UPDATE stage_revisions SET data_hash=? WHERE id=?", (rows[0]["data_hash"], rows[0]["id"]), drop_triggers=True)
        self.assertTrue(stage(a, mid, "atp")["signatures_ok"])

    def test_URS86_signature_is_bound_to_what_the_reviewer_saw(self):
        """t6: a reviewer who loaded revision 1 signed revision 2 (completely different content) with a click."""
        a, mid, ids = make_project()
        base = "/api/methods/%d/stages/atp" % mid
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        a.post(base + "/submit", json={})
        rita = login("rita")
        seen = stage(rita, mid, "atp")
        a.post(base + "/withdraw", json={"reason": "tweak"})
        a.put(base, json={"version": stage(a, mid, "atp")["version"], "sections": {"use": "COMPLETELY DIFFERENT CONTENT"}})
        a.post(base + "/submit", json={})
        now = stage(rita, mid, "atp")
        self.assertEqual(now["revision"], seen["revision"] + 1)

        def post(**kw):
            body = {"decision": "approve", "comment": "ok", "password": PW}
            body.update(kw)
            return rita.post(base + "/review", json=body)
        r = post(revision=seen["revision"], expected_hash=seen["hash"])             # the stale view
        self.assertEqual(r.status_code, 409, r.text)
        self.assertIn("刷新", r.json()["error"])
        self.assertEqual(post(revision=now["revision"], expected_hash=seen["hash"]).status_code, 409)   # right revision, wrong hash
        self.assertEqual(post().status_code, 400)                                                         # nothing sent
        self.assertEqual(post(revision=now["revision"]).status_code, 400)
        self.assertEqual(post(expected_hash=now["hash"]).status_code, 400)
        self.assertEqual(post(revision=str(now["revision"]), expected_hash=now["hash"]).status_code, 400)
        self.assertEqual(post(revision=True, expected_hash=now["hash"]).status_code, 400)
        self.assertEqual(post(revision=now["revision"], expected_hash=123).status_code, 400)
        self.assertEqual(post(revision=now["revision"], expected_hash=now["hash"], decision=["approve"]).status_code, 400)
        self.assertEqual(stage(rita, mid, "atp")["reviews"], [])                     # nothing was signed
        self.assertEqual(post(revision=now["revision"], expected_hash=now["hash"]).status_code, 200)
        self.assertEqual(stage(rita, mid, "atp")["reviews"][-1]["revision"], now["revision"])


# ====================================================================== M1  evidence tables / audit chain
class TestAuditProtection(unittest.TestCase):
    def test_URS87_append_only_triggers(self):
        a, mid, ids = make_project()
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        a.post("/api/methods/%d/stages/atp/submit" % mid, json={})
        sign(login("rita"), mid, "atp")
        con = sqlite3.connect(str(config.DB_PATH))
        try:
            for sql in ("UPDATE audit_log SET details='{}' WHERE id=1", "DELETE FROM audit_log WHERE id=1",
                        "UPDATE stage_revisions SET data_hash='x'", "DELETE FROM stage_revisions",
                        "UPDATE reviews SET data_hash='x'", "DELETE FROM reviews", "UPDATE reviews SET decision='reject'"):
                with self.assertRaises(sqlite3.DatabaseError, msg=sql):
                    con.execute(sql)
                    con.commit()
        finally:
            con.close()
        self.assertTrue(login("admin").get("/api/audit/verify").json()["triggers_ok"])

    def test_URS87_keyed_chain_rewrite_truncation_and_anchor(self):
        """t6: with the old unkeyed chain an attacker could rewrite one record and recompute the whole chain; and
        deleting the newest records was undetectable."""
        a, mid, ids = make_project()
        admin = login("admin")
        v = admin.get("/api/audit/verify").json()
        self.assertTrue(v["ok"], v)
        self.assertEqual(v["anchor_status"], "ok")
        self.assertTrue(os.path.exists(str(config.AUDIT_KEY)))
        head = audit.read_head()
        self.assertTrue(audit.head_valid(head))
        last = raw("SELECT id, hash, hv FROM audit_log ORDER BY id DESC LIMIT 1")[0]
        self.assertEqual((head["id"], head["hash"]), (last["id"], last["hash"]))
        self.assertEqual(last["hv"], 3)
        # --- rewrite one record and recompute every hash WITHOUT the key (plain SHA-256, as in the review script)
        c = copy_db()
        for n in db.TRIGGER_NAMES:
            c.execute("DROP TRIGGER IF EXISTS %s" % n)
        rows = c.execute("SELECT * FROM audit_log ORDER BY id").fetchall()
        prev = audit.GENESIS
        for r in rows:
            body = {"ts": r["ts"], "user_id": r["user_id"], "username": "mallory" if r["id"] == rows[3]["id"] else r["username"],
                    "action": r["action"], "entity_type": r["entity_type"], "entity_id": r["entity_id"], "project_id": r["project_id"],
                    "details": json.loads(r["details"]), "reason": r["reason"]}
            h = db.sha256(prev + db.canon(body))
            c.execute("UPDATE audit_log SET username=?, prev_hash=?, hash=? WHERE id=?", (body["username"], prev, h, r["id"]))
            prev = h
        res = audit.verify_chain(c)
        self.assertFalse(res["ok"])
        self.assertIsNotNone(res["first_bad_id"])
        c.close()
        # --- delete the last 20 records: the bare chain still verifies, the anchor does not
        c = copy_db()
        for n in db.TRIGGER_NAMES:
            c.execute("DROP TRIGGER IF EXISTS %s" % n)
        c.execute("DELETE FROM audit_log WHERE id > (SELECT MAX(id)-20 FROM audit_log)")
        self.assertTrue(audit.verify_chain(c)["ok"])                       # what the old code concluded
        res = audit.verify_chain(c, anchor=audit.read_head())
        self.assertFalse(res["ok"])
        self.assertIn(res["anchor_status"], ("shorter", "mismatch"))
        c.close()
        # --- anchor file deleted / altered: reported, and NOT silently re-created from the live log
        with open(str(config.AUDIT_HEAD), encoding="utf-8") as f:
            saved = f.read()
        try:
            os.remove(str(config.AUDIT_HEAD))
            login("alice")                                                  # an audited commit while the anchor is missing
            self.assertFalse(os.path.exists(str(config.AUDIT_HEAD)))
            v = admin.get("/api/audit/verify").json()
            self.assertFalse(v["ok"])
            self.assertEqual(v["anchor_status"], "missing")
            h = json.loads(saved)
            h["count"] = h["count"] - 5
            with open(str(config.AUDIT_HEAD), "w", encoding="utf-8") as f:
                json.dump(h, f)
            self.assertEqual(admin.get("/api/audit/verify").json()["anchor_status"], "tampered")
        finally:
            with open(str(config.AUDIT_HEAD), "w", encoding="utf-8") as f:
                f.write(saved)
        self.assertTrue(admin.get("/api/audit/verify").json()["ok"])


# ====================================================================== M2  project status
class TestProjectStatus(unittest.TestCase):
    def test_URS88_nothing_can_be_written_in_a_non_active_method(self):
        """t2 H: after archiving, stages, experiments and attachments could still be changed."""
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        self.assertEqual(a.post(url + "/stages/atp/submit", json={}).status_code, 200)     # atp in review
        up = upload(a, mid, "a.txt", b"hello", stage_key="tech")
        self.assertEqual(up.status_code, 200)
        aid = up.json()["id"]
        free = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        rita = login("rita")
        self.assertEqual(a.patch(url, json={"status": "archived"}).status_code, 200)
        ver = stage(a, mid, "tech")["version"]
        checks = {
            "save stage": a.put(url + "/stages/tech", json={"version": ver, "planned_end": "2031-01-01"}),
            "submit": a.post(url + "/stages/tech/submit", json={}),
            "withdraw": a.post(url + "/stages/atp/withdraw", json={"reason": "x"}),
            "recertify": a.post(url + "/stages/tech/recertify", json={}),
            "upload": upload(a, mid, "arch.txt", b"x"),
            "delete attachment": a.delete("/api/attachments/%d?reason=x" % aid),
            "exp update": exp_put(a, free["id"], conclusion="changed"),
            "exp create": a.post(url + "/experiments", json={"title": "z"}),
            "exp status": a.post("/api/experiments/%d/status" % free["id"], json={"status": "cancelled", "reason": "r"}),
            "exp delete": a.delete("/api/experiments/%d?reason=x" % free["id"]),
            "members": a.put(url + "/members", json={"members": members_of(BASE + [("vic", "viewer")])}),
        }
        for label, r in checks.items():
            self.assertEqual(r.status_code, 409, (label, r.text))
        self.assertIn(sign(rita, mid, "atp").status_code, (403, 409))           # a signature cannot be given either
        self.assertEqual(stage(a, mid, "atp")["reviews"], [])
        d = rita.get("/api/dashboard").json()
        self.assertFalse([t for t in d["to_review"] if t["method_id"] == mid])

    def test_URS88_reactivating_needs_reason_and_password(self):
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        self.assertEqual(a.patch(url, json={"status": "archived"}).status_code, 200)
        self.assertEqual(a.patch(url, json={"status": "active"}).status_code, 400)                              # no reason
        self.assertEqual(a.patch(url, json={"status": "active", "reason": "   ", "password": PW}).status_code, 400)
        self.assertEqual(a.patch(url, json={"status": "active", "reason": "误归档"}).status_code, 403)            # no password
        self.assertEqual(a.patch(url, json={"status": "active", "reason": "误归档", "password": "wrong-pass-9"}).status_code, 403)
        self.assertEqual(a.get(url).json()["status"], "archived")
        self.assertEqual(a.patch(url, json={"status": "active", "reason": "误归档", "password": PW}).status_code, 200)
        self.assertEqual(a.get(url).json()["status"], "active")
        row = [r for r in a.get(url + "/audit").json() if r["action"] == "method.update" and r["details"].get("status") == ["archived", "active"]]
        self.assertEqual(row[0]["reason"], "误归档")
        self.assertEqual(a.patch(url, json={"status": "cancelled"}).status_code, 200)                           # forward moves stay simple
        self.assertEqual(a.patch(url, json={"status": "active", "reason": "又要做了", "password": PW}).status_code, 200)

    def test_URS88_update_method_validates_fields(self):
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        for body in ({"name": ""}, {"name": None}, {"name": "   "}, {"name": 5}, {"analyte": {"a": 1}}, {"intended_use": ["x"]}, {"matrix": 3},
                     {"planned_start": {"a": 1}}, {"planned_end": "garbage"}, {"status": ["completed"]}, {"status": "bogus"}, {"name": "x" * 500},
                     {"method_version": ""}, {"method_version": ["V1"]}, {"planned_start": "2031-05-01", "planned_end": "2031-01-01"},
                     {"lifecycle": "routine"}, {"project_id": 99999}, {"template_key": "hplc"}):
            self.assertEqual(a.patch(url, json=body).status_code, 400, body)
        self.assertEqual(a.get(url).json()["name"], "CAR-T 流式测试")

    def test_URS88_a_non_active_project_blocks_writes_to_its_methods(self):
        """The project's status gates its methods as well: paused / completed / archived = nothing can change."""
        a, mid, ids = make_project()
        pid = project_id_of(a, mid)
        purl, url = "/api/projects/%d" % pid, "/api/methods/%d" % mid
        free = next(e for e in exps(a, mid) if e["stage_key"] == "explore")
        for st in ("paused", "archived"):
            self.assertEqual(a.patch(purl, json={"status": st}).status_code, 200, st)
            ver = stage(a, mid, "charter")["version"]
            self.assertEqual(a.put(url + "/stages/charter", json={"version": ver, "sections": {"background": "x"}}).status_code, 409, st)
            self.assertEqual(a.post(url + "/experiments", json={"title": "z"}).status_code, 409, st)
            self.assertEqual(exp_put(a, free["id"], conclusion="c").status_code, 409, st)
            self.assertEqual(upload(a, mid, "p.txt", b"x").status_code, 409, st)
            self.assertEqual(a.put(url + "/members", json={"members": members_of(BASE)}).status_code, 409, st)
            self.assertEqual(a.post(purl + "/methods", json={"name": "new"}).status_code, 409, st)
            self.assertEqual(a.put(purl + "/members", json={"members": members_of([("alice", "owner")])}).status_code, 409, st)
            self.assertEqual(a.get(url).status_code, 200)                                   # reading stays possible
            self.assertEqual(a.patch(purl, json={"status": "active"}).status_code, 400, st)  # reason + password needed
            self.assertEqual(a.patch(purl, json={"status": "active", "reason": "继续", "password": "wrong-pass-9"}).status_code, 403, st)
            self.assertEqual(a.patch(purl, json={"status": "active", "reason": "继续", "password": PW}).status_code, 200, st)
        self.assertEqual(a.post(url + "/experiments", json={"title": "z"}).status_code, 200)
        row = [r for r in a.get(purl + "/audit").json() if r["action"] == "project.update" and r["details"].get("status") == ["archived", "active"]]
        self.assertEqual(row[0]["reason"], "继续")

    def test_URS88_update_project_validates_fields(self):
        a, mid, ids = make_project()
        purl = "/api/projects/%d" % project_id_of(a, mid)
        for body in ({"name": ""}, {"name": None}, {"name": "   "}, {"name": 5}, {"product": {"a": 1}}, {"indication": ["x"]}, {"description": 5},
                     {"planned_start": {"a": 1}}, {"planned_end": "garbage"}, {"status": ["completed"]}, {"status": "bogus"}, {"status": "cancelled"},
                     {"name": "x" * 500}, {"planned_start": "2031-05-01", "planned_end": "2031-01-01"}):
            self.assertEqual(a.patch(purl, json=body).status_code, 400, body)
        self.assertEqual(a.get(purl).json()["name"], "项目-CAR-T 流式测试")


# ====================================================================== M3 / M4  accounts
def _mk_user(admin, name, roles=("author",)):
    r = admin.post("/api/users", json={"username": name, "display_name": name, "password": PW, "roles": list(roles)})
    assert r.status_code == 200, r.text
    return r.json()["id"]


class TestAccounts(unittest.TestCase):
    def test_URS89_must_change_is_enforced_by_the_api(self):
        """t6: a user with the initial password could call every API; changing a password left other sessions open."""
        admin = login("admin")
        _mk_user(admin, "newbie1")
        c = rclient("newbie1")
        self.assertTrue(c.get("/api/me").json()["user"]["must_change"])
        for method, url, body in (("post", "/api/projects", {"title": "mc"}), ("get", "/api/dashboard", None), ("get", "/api/users", None),
                                  ("get", "/api/templates", None), ("get", "/api/kb", None), ("get", "/api/projects/1", None)):
            r = getattr(c, method)(url, **({"json": body} if body is not None else {}))
            self.assertEqual(r.status_code, 403, (url, r.text))
            self.assertEqual(r.json(), {"error": "须先修改密码", "must_change": True})
        other = rclient("newbie1")                                          # a second session
        self.assertEqual(c.post("/api/me/password", json={"old": PW, "new": "short1"}).status_code, 400)
        self.assertEqual(c.post("/api/me/password", json={"old": PW, "new": PW}).status_code, 400)        # same password
        r = c.post("/api/me/password", json={"old": PW, "new": "NewPassw0rd!!"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(c.get("/api/dashboard").status_code, 200)         # this session continues, unrestricted now
        self.assertFalse(c.get("/api/me").json()["user"]["must_change"])
        self.assertEqual(other.get("/api/me").status_code, 401)            # every other session is gone
        self.assertEqual(c.post("/api/logout").status_code, 200)

    def test_URS89_change_password_counts_failures_and_locks(self):
        admin = login("admin")
        _mk_user(admin, "newbie2")
        c = rclient("newbie2")
        for i in range(config.get("max_failed_logins")):
            self.assertEqual(c.post("/api/me/password", json={"old": "bad-pass-%d" % i, "new": "Whatever123!"}).status_code, 400)
        r = c.post("/api/me/password", json={"old": PW, "new": "Whatever123!"})           # right old password, but locked
        self.assertEqual(r.status_code, 423, r.text)
        self.assertEqual(rclient().post("/api/login", json={"username": "newbie2", "password": PW}).status_code, 423)
        fails = raw("SELECT COUNT(*) FROM audit_log WHERE action='user.password_check_fail'")[0][0]
        self.assertGreaterEqual(fails, config.get("max_failed_logins"))

    def test_URS90_login_does_not_reveal_which_accounts_exist(self):
        """t4: unknown user, wrong password, locked and deactivated accounts answered differently."""
        admin = login("admin")
        _mk_user(admin, "lockme")
        _mk_user(admin, "inact1")
        admin.patch("/api/users/%d" % IDS_of(admin, "inact1"), json={"active": False})
        c = rclient()
        unknown = c.post("/api/login", json={"username": "no-such-user", "password": "x" * 12})
        known = c.post("/api/login", json={"username": "lockme", "password": "x" * 12})
        inactive = c.post("/api/login", json={"username": "inact1", "password": "x" * 12})
        self.assertEqual((unknown.status_code, unknown.json()), (known.status_code, known.json()))
        self.assertEqual((inactive.status_code, inactive.json()), (known.status_code, known.json()))
        self.assertEqual(known.status_code, 401)
        # a deactivated account is only named as such to somebody who knows the password
        r = c.post("/api/login", json={"username": "inact1", "password": PW})
        self.assertEqual(r.status_code, 403)
        # lock-out looks the same for an existing and a non-existing name (and the right password does not help)
        c.post("/api/login", json={"username": "ghost-lock", "password": "x" * 12})      # 'lockme' has had one failure above as well
        for _ in range(config.get("max_failed_logins") - 1):
            c.post("/api/login", json={"username": "lockme", "password": "x" * 12})
            c.post("/api/login", json={"username": "ghost-lock", "password": "x" * 12})
        locked_known = c.post("/api/login", json={"username": "lockme", "password": PW})
        locked_unknown = c.post("/api/login", json={"username": "ghost-lock", "password": PW})
        self.assertEqual(locked_known.status_code, 423)
        self.assertEqual((locked_known.status_code, locked_known.json()), (locked_unknown.status_code, locked_unknown.json()))
        # unknown names cost a bcrypt comparison as well
        t0 = time.time()
        c.post("/api/login", json={"username": "ghost-timing", "password": "x" * 12})
        self.assertGreater(time.time() - t0, 0.05)

    def test_URS90_rate_limit_and_audit_dedupe_for_unknown_names(self):
        auth.reset_limits()
        c = rclient()
        with cfg(login_rate_max=3, max_failed_logins=50):
            codes = [c.post("/api/login", json={"username": "rate-ghost", "password": "x" * 12}).status_code for _ in range(5)]
            self.assertEqual(codes, [401, 401, 401, 429, 429])
            # another user name from the same address is not limited
            self.assertEqual(c.post("/api/login", json={"username": "rate-ghost2", "password": "x" * 12}).status_code, 401)
        auth.reset_limits()
        n0 = raw("SELECT COUNT(*) FROM audit_log WHERE action='login.fail'")[0][0]
        with cfg(login_rate_max=100, max_failed_logins=50):
            for _ in range(8):
                c.post("/api/login", json={"username": "ghost-dedupe", "password": "x" * 12})
        rows = raw("SELECT COUNT(*) FROM audit_log WHERE action='login.fail' AND details LIKE '%ghost-dedupe%'")[0][0]
        self.assertEqual(rows, 1)
        self.assertEqual(raw("SELECT COUNT(*) FROM audit_log WHERE action='login.fail'")[0][0], n0 + 1)
        auth.reset_limits()

    def test_URS90_login_clears_expired_sessions(self):
        raw("INSERT INTO sessions(token_hash,user_id,created_at,last_seen,expires_at) VALUES('deadbeef',1,'2020-01-01T00:00:00Z','2020-01-01T00:00:00Z','2020-01-02T00:00:00Z')")
        self.assertEqual(raw("SELECT COUNT(*) FROM sessions WHERE token_hash='deadbeef'")[0][0], 1)
        login("alice")
        self.assertEqual(raw("SELECT COUNT(*) FROM sessions WHERE token_hash='deadbeef'")[0][0], 0)


def IDS_of(admin, name):
    return next(u["id"] for u in admin.get("/api/users").json() if u["username"] == name)


# ====================================================================== M5  uploads
class TestUploads(unittest.TestCase):
    def test_URS91_upload_checks_before_reading_and_enforces_limits(self):
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        vmid = make_project(title="V", team=["vic"], method_team=members_of([("vic", "viewer")]))[1]
        # a non-member / a read-only member is refused (403) and nothing is stored
        self.assertEqual(upload(login("bob"), mid, "x.bin", b"x" * 1000).status_code, 403)
        self.assertEqual(upload(login("vic"), vmid, "x.bin", b"x" * 1000).status_code, 403)
        self.assertEqual(upload(a, mid, "x.bin", b"").status_code, 400)                              # empty
        self.assertEqual(upload(a, mid, "x.bin", b"x", stage_key="no-such-stage").status_code, 404)
        self.assertEqual(upload(a, mid, "x.bin", b"x", experiment_id=99999).status_code, 404)
        self.assertEqual(raw("SELECT COUNT(*) FROM attachments WHERE method_id=?", (mid,))[0][0], 0)
        with cfg(max_upload_mb=1):
            r = upload(a, mid, "big.bin", b"x" * (1024 * 1024 + 10))
            self.assertEqual(r.status_code, 413, r.text)
            self.assertEqual(upload(a, mid, "ok.bin", b"x" * 1000).status_code, 200)
        tmpdir = config.UPLOADS / ".tmp"
        self.assertEqual(os.listdir(str(tmpdir)) if tmpdir.exists() else [], [])                     # no leftovers
        self.assertEqual(raw("SELECT COUNT(*) FROM attachments WHERE method_id=?", (mid,))[0][0], 1)
        # JSON bodies are capped as well
        with cfg(max_json_mb=1):
            ver = stage(a, mid, "charter")["version"]
            r = a.put(url + "/stages/charter", json={"version": ver, "sections": {"background": "x" * (2 * 1024 * 1024)}})
            self.assertEqual(r.status_code, 413, r.text)
            self.assertIn("error", r.json())
        self.assertEqual(a.put(url + "/stages/charter", json={"version": ver, "sections": {"background": "ok"}}).status_code, 200)

    def test_URS91_upload_endpoint_is_streaming_and_does_not_block_the_event_loop(self):
        import inspect
        self.assertTrue(inspect.iscoroutinefunction(main.att_upload))
        src = inspect.getsource(main.att_upload)
        self.assertIn("request.stream()", src)
        self.assertNotIn("await request.body()", src)
        self.assertIn("run_in_threadpool(current_user", src)       # authentication runs in a worker thread

    def test_URS97_filenames_are_not_decoded_twice(self):
        """Coordinator item: the framework decodes the query string once; an extra unquote() mangled a literal '%'."""
        a, mid, ids = make_project()
        names = ["100%_测试 file.txt", "%41%42.txt", "a%2Fb.txt", "50% off.csv"]
        for n in names:
            r = upload(a, mid, n, ("data of " + n).encode("utf-8"), stage_key="tech")
            self.assertEqual(r.status_code, 200, (n, r.text))
        got = [x["filename"] for x in stage(a, mid, "tech")["attachments"]]
        for n in names:
            self.assertIn(n, got)
        att = next(x for x in stage(a, mid, "tech")["attachments"] if x["filename"] == "100%_测试 file.txt")
        self.assertEqual(a.get("/api/attachments/%d/download" % att["id"]).content, "data of 100%_测试 file.txt".encode("utf-8"))
        self.assertTrue(a.post("/api/attachments/%d/verify" % att["id"]).json()["intact"])
        # path tricks are still neutralised, and nothing is written outside uploads/<project>/
        for evil in ("../../evil.txt", "..\\..\\evil2.txt", "C:\\Windows\\evil3.txt"):
            r = upload(a, mid, evil, b"evil", stage_key="tech")
            self.assertEqual(r.status_code, 200, r.text)
        for row in raw("SELECT path FROM attachments WHERE method_id=?", (mid,)):
            full = (config.DATA / row["path"]).resolve()
            self.assertEqual(full.parent, (config.UPLOADS / str(mid)).resolve())
        self.assertFalse(os.path.exists(os.path.join(HTMP, "evil.txt")))

    def test_URS91_existing_file_is_never_overwritten(self):
        a, mid, ids = make_project()
        r1 = upload(a, mid, "same.bin", b"payload", stage_key="tech")
        r2 = upload(a, mid, "same.bin", b"payload", stage_key="tech")
        self.assertEqual((r1.status_code, r2.status_code), (200, 200))
        p1, p2 = [x["path"] for x in raw("SELECT path FROM attachments WHERE method_id=? ORDER BY id", (mid,))]
        self.assertEqual(p1, p2)                                           # identical content shares the file
        (config.DATA / p1).write_bytes(b"altered on disk")                  # then the stored copy is damaged
        r3 = upload(a, mid, "same.bin", b"payload", stage_key="tech")
        self.assertEqual(r3.status_code, 200)
        paths = [x["path"] for x in raw("SELECT path FROM attachments WHERE method_id=? ORDER BY id", (mid,))]
        self.assertNotEqual(paths[2], paths[0])                            # the damaged evidence is kept, the new copy is separate
        self.assertEqual((config.DATA / paths[0]).read_bytes(), b"altered on disk")
        ids_ = [x["id"] for x in raw("SELECT id FROM attachments WHERE method_id=? ORDER BY id", (mid,))]
        self.assertFalse(a.post("/api/attachments/%d/verify" % ids_[0]).json()["intact"])
        self.assertTrue(a.post("/api/attachments/%d/verify" % ids_[2]).json()["intact"])


# ====================================================================== M6  backups
class TestBackups(unittest.TestCase):
    def test_URS92_backup_is_checked_and_comes_with_a_manifest(self):
        a, mid, ids = make_project()
        r = upload(a, mid, "m.bin", b"manifest me", stage_key="tech")
        self.assertEqual(r.status_code, 200)
        a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        name = login("admin").post("/api/admin/backup").json()["name"]
        mp = config.BACKUPS / (name + ".manifest.json")
        self.assertTrue(mp.exists())
        m = json.loads(mp.read_text(encoding="utf-8"))
        self.assertEqual(m["integrity_check"], "ok")
        self.assertEqual(m["schema_version"], db.SCHEMA_VERSION)
        self.assertEqual(m["db"], name)
        # the audit head in the manifest is the backup's own last record and is verifiable
        con = sqlite3.connect("file:%s?mode=ro&immutable=1" % (config.BACKUPS / name).as_posix(), uri=True)
        con.row_factory = sqlite3.Row
        try:
            last = con.execute("SELECT id, hash FROM audit_log ORDER BY id DESC LIMIT 1").fetchone()
            self.assertEqual((m["audit_head"]["id"], m["audit_head"]["hash"]), (last["id"], last["hash"]))
            self.assertTrue(audit.head_valid(m["audit_head"]))
            self.assertTrue(audit.verify_chain(con, anchor=m["audit_head"])["ok"])
        finally:
            con.close()
        path = raw("SELECT path FROM attachments WHERE method_id=?", (mid,))[0]["path"].replace("\\", "/")
        self.assertEqual(m["files"][path]["sha256"], db.sha256(b"manifest me"))
        self.assertTrue(any(k.startswith("documents/") for k in m["files"]))
        listed = [b for b in workflow.list_backups() if b["name"] == name][0]
        self.assertTrue(listed["manifest"])

    def test_URS92_startup_and_daily_backups_are_retained_separately(self):
        for f in list(config.BACKUPS.glob("methodflow_*")):
            f.unlink()
        with cfg(backup_keep=3, backup_keep_startup=2):
            for _ in range(4):
                workflow.backup_now("startup")
            for _ in range(5):
                workflow.backup_now("daily")
            workflow.backup_now("manual")
        names = sorted(f.name for f in config.BACKUPS.glob("methodflow_*.db"))
        self.assertEqual(len([n for n in names if "_startup" in n]), 2)
        self.assertEqual(len([n for n in names if "_daily" in n]), 3)
        self.assertEqual(len([n for n in names if "_manual" in n]), 1)
        manifests = sorted(f.name for f in config.BACKUPS.glob("methodflow_*.manifest.json"))
        self.assertEqual(manifests, sorted(n + ".manifest.json" for n in names))        # pruned together

    def test_URS92_backup_that_fails_integrity_check_is_rejected(self):
        real = sqlite3.connect
        calls = {"n": 0}

        class FakeCur:
            def fetchone(self):
                return ("*** in database main ***\nPage 2: btree error",)

        class FakeCon:
            row_factory = None

            def execute(self, *a, **k):
                return FakeCur()

            def close(self):
                pass

        def fake(*a, **k):
            calls["n"] += 1
            return FakeCon() if calls["n"] == 3 else real(*a, **k)
        with mock.patch.object(sqlite3, "connect", fake):
            with self.assertRaises(RuntimeError):
                workflow.backup_now("integ")
        self.assertTrue(list(config.BACKUPS.glob("methodflow_*_integ.db.corrupt")))
        self.assertFalse(list(config.BACKUPS.glob("methodflow_*_integ.db")))

    def test_URS92_backup_tool_with_files_still_works(self):
        a, mid, ids = make_project()
        upload(a, mid, "t.bin", b"tool", stage_key="tech")
        data = os.path.join(HTMP, "tool_copy")
        shutil.rmtree(data, ignore_errors=True)
        os.makedirs(data)
        s = sqlite3.connect(str(config.DB_PATH)); d = sqlite3.connect(os.path.join(data, "methodflow.db")); s.backup(d); s.close(); d.close()
        for item in (".audit_key", "audit_head.json"):
            shutil.copy(str(config.DATA / item), os.path.join(data, item))
        shutil.copytree(str(config.UPLOADS), os.path.join(data, "uploads"))
        shutil.copytree(str(config.DOCS), os.path.join(data, "documents"))
        p = subprocess.run([sys.executable, os.path.join(ROOT, "tools", "backup.py"), "--data", data, "--files"], capture_output=True)
        out = p.stdout.decode("utf-8", "replace")
        self.assertEqual(p.returncode, 0, out + p.stderr.decode("utf-8", "replace"))
        self.assertIn("已创建备份", out)
        zips = [f for f in os.listdir(os.path.join(data, "backups")) if f.startswith("files_") and f.endswith(".zip")]
        self.assertEqual(len(zips), 1)
        with zipfile.ZipFile(os.path.join(data, "backups", zips[0])) as z:
            self.assertTrue(any(n.replace("\\", "/").startswith("uploads/") for n in z.namelist()))
            self.assertFalse(any(".tmp" in n for n in z.namelist()))


# ====================================================================== M7  documents
class TestDocuments(unittest.TestCase):
    def test_URS93_documents_survive_control_characters_and_junk_numbers(self):
        config.override("enforce_stage_order", False); self.addCleanup(config.override, "enforce_stage_order", True)
        """t5: a vertical tab / NUL in a text, a weight like '5%', or odd cells crashed document generation."""
        a, mid, ids = make_project()
        url = "/api/methods/%d" % mid
        ver = stage(a, mid, "charter")["version"]
        r = a.put(url + "/stages/charter", json={"version": ver, "sections": {"background": "line1\x0bline2 & <b>x</b> \x00 \x1f end 😀"}})
        self.assertEqual(r.status_code, 200, r.text)
        r = a.post(url + "/docs", json={"kind": "plan"})
        self.assertEqual(r.status_code, 200, r.text)
        text = docx_text(a, r.json()["id"])
        self.assertIn("line1", text)
        self.assertIn("line2", text)
        self.assertNotIn("\x0b", text)
        self.assertIn("<b>x</b>", text)
        # such a value can no longer be saved through the API ...
        tech = stage(a, mid, "tech")
        m = tech["data"]["sections"]["matrix"]
        m["criteria"][0]["weight"] = "5%"
        self.assertEqual(a.put(url + "/stages/tech", json={"version": tech["version"], "sections": {"matrix": m}}).status_code, 400)
        # ... but data that is already stored (older versions) must not break the document
        data = json.loads(raw("SELECT data FROM stages WHERE method_id=? AND key='tech'", (mid,))[0][0])
        data["sections"]["matrix"]["criteria"][0]["weight"] = "5%"
        data["sections"]["matrix"]["criteria"][1]["scores"] = {data["sections"]["matrix"]["options"][0]: "abc", data["sections"]["matrix"]["options"][1]: None}
        data["sections"]["matrix"]["options"] = [1, 2]
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='tech'", (json.dumps(data), mid))
        a2 = a.post(url + "/docs", json={"kind": "plan"})
        self.assertEqual(a2.status_code, 200, a2.text)
        data["sections"]["matrix"] = "oops"
        data["sections"]["decision"] = {"a": 1}
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='tech'", (json.dumps(data), mid))
        self.assertEqual(a.post(url + "/docs", json={"kind": "report"}).status_code, 200)

    def test_URS93_fmea_scores_are_validated_on_save(self):
        """t3/t5: S = 10**400 or 'abc' made the stage / document unusable."""
        a, mid, ids = make_project()
        approve_stage(mid, "charter", a)
        approve_stage(mid, "atp", a)
        url = "/api/methods/%d" % mid

        def put(rows):
            return a.put(url + "/stages/risk", json={"version": stage(a, mid, "risk")["version"], "sections": {"fmea": rows}})
        for bad in ({"mode": "m", "S": "abc", "O": 1, "D": 1}, {"mode": "m", "S": 10 ** 400, "O": 1, "D": 1}, {"mode": "m", "S": 11, "O": 1, "D": 1},
                    {"mode": "m", "S": 0, "O": 1, "D": 1}, {"mode": "m", "S": 5.5, "O": 1, "D": 1}, {"mode": "m", "S": True, "O": 1, "D": 1},
                    {"mode": "m", "S": -3, "O": 1, "D": 1}, {"mode": "m", "S": [5], "O": 1, "D": 1}, {"mode": "m", "S": "1e3", "O": 1, "D": 1},
                    {"mode": 5, "S": 1, "O": 1, "D": 1}):
            r = put([bad])
            self.assertEqual(r.status_code, 400, (bad, r.text[:100]))
        r = put([{"mode": "ok", "S": "7", "O": 5.0, "D": 3, "control": None, "cause": None, "effect": "e"},
                 {"mode": "unscored", "S": "", "O": None, "D": ""}])
        self.assertEqual(r.status_code, 200, r.text)
        rows = stage(a, mid, "risk")["data"]["sections"]["fmea"]
        self.assertEqual((rows[0]["S"], rows[0]["O"], rows[0]["D"]), (7, 5, 3))
        self.assertEqual(rows[0]["control"], "")
        self.assertEqual(a.post(url + "/docs", json={"kind": "plan"}).status_code, 200)

    def test_URS93_charts_use_the_figure_api_not_pyplot(self):
        self.assertFalse(hasattr(docgen, "plt"))
        import inspect
        import re as _re
        src = inspect.getsource(docgen)
        self.assertFalse(_re.search(r"^\s*(import matplotlib\.pyplot|from matplotlib import pyplot|from matplotlib\.pyplot)", src, _re.M))
        self.assertIn("from matplotlib.figure import Figure", src)
        a, mid, ids = make_project()
        with db.read() as conn:
            png = docgen._gantt_png(workflow.list_stages(conn, mid))
            self.assertEqual(png.read(4), b"\x89PNG")
            png2 = docgen._fmea_png(workflow.get_stage(conn, mid, "risk")["data"]["sections"]["fmea"])
            self.assertEqual(png2.read(4), b"\x89PNG")

    def test_URS93_rendering_does_not_hold_the_database_write_lock(self):
        """t5: the whole Word document, charts included, was built inside the write transaction, so every other
        save waited for it."""
        a, mid, ids = make_project()
        b = login("alice")
        orig = docgen.render
        started = threading.Event()

        def slow(ctx):
            started.set()
            time.sleep(2.0)
            return orig(ctx)
        res = {}

        def gen():
            r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
            res["gen"] = r.status_code
        with mock.patch.object(docgen, "render", slow):
            th = threading.Thread(target=gen)
            th.start()
            self.assertTrue(started.wait(10))
            ver = stage(b, mid, "charter")["version"]
            t0 = time.time()
            r = b.put("/api/methods/%d/stages/charter" % mid, json={"version": ver, "sections": {"background": "written while rendering"}})
            elapsed = time.time() - t0
            self.assertTrue(th.is_alive())
            self.assertEqual(r.status_code, 200)
            self.assertLess(elapsed, 1.2)
            th.join(30)
        # the stage changed while the document was being built -> the document is not recorded from stale data
        self.assertEqual(res["gen"], 409)
        self.assertEqual(a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).status_code, 200)

    def test_URS93_parallel_generation_in_two_projects(self):
        a, p1, _ = make_project(title="par1")
        b, p2, _ = make_project(title="par2")
        out = []

        def gen(c, mid):
            out.append(c.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).status_code)
        ts = [threading.Thread(target=gen, args=(a, p1)), threading.Thread(target=gen, args=(b, p2))]
        [t.start() for t in ts]
        [t.join(60) for t in ts]
        self.assertEqual(sorted(out), [200, 200])


# ====================================================================== M8 / L1  validation
class TestValidation(unittest.TestCase):
    def test_URS94_malformed_sections_are_rejected_and_never_brick_the_stage(self):
        config.override("enforce_stage_order", False); self.addCleanup(config.override, "enforce_stage_order", True)
        """t3: PUT matrix='oops' stored fine, then GET stage / project / dashboard / documents returned 500."""
        mid = make_project(title="fuzz")[1]
        a = rclient("alice")
        base = "/api/methods/%d" % mid

        def ver(k):
            return a.get(base + "/stages/" + k).json()["version"]
        cases = [
            ("tech", {"matrix": "oops"}), ("tech", {"matrix": [1]}), ("tech", {"matrix": {"options": "x"}}),
            ("tech", {"matrix": {"options": [1, 2]}}), ("tech", {"matrix": {"criteria": [{"name": "a", "weight": "5%"}]}}),
            ("tech", {"matrix": {"criteria": [{"name": "a", "scores": [1]}]}}), ("tech", {"matrix": {"criteria": "x"}}),
            ("risk", {"fmea": ["abc"]}), ("risk", {"fmea": {"a": 1}}), ("risk", {"fmea": [{"mode": "m", "S": "abc"}]}),
            ("atp", {"criteria": ["abc"]}), ("atp", {"criteria": {"a": 1}}), ("atp", {"criteria": [{"criterion": {"a": 1}}]}),
            ("atp", {"criteria": [{"criterion": ["x"]}]}), ("charter", {"background": ["a", "b"]}), ("charter", {"background": 5}),
            ("charter", {"context": "abc"}), ("charter", {"context": {"nope": "x"}}), ("charter", {"context": {"phase": {"a": 1}}}),
            ("charter", {"nonexistent": "x"}),
        ]
        for key, secs in cases:
            r = a.put(base + "/stages/" + key, json={"version": ver(key), "sections": secs})
            self.assertEqual(r.status_code, 400, (key, secs, r.text[:120]))
            self.assertNotIn("Traceback", r.text)
        for label, payload in [
            ("version str", {"version": "abc", "sections": {}}), ("version None", {"version": None}), ("version float", {"version": 1.5, "sections": {}}),
            ("sections list", {"sections": [1]}), ("checklist str", {"checklist": {"c1": "done"}}), ("checklist note int", {"checklist": {"c1": {"status": "done", "note": 5}}}),
            ("checklist status", {"checklist": {"c1": {"status": "weird"}}}), ("checklist unknown", {"checklist": {"zz": {"status": "done"}}}),
            ("planned_start dict", {"planned_start": {"a": 1}}), ("planned_start garbage", {"planned_start": "garbage"}),
        ]:
            body = dict(payload)
            body.setdefault("version", ver("charter"))
            r = a.put(base + "/stages/charter", json=body)
            self.assertEqual(r.status_code, 400, (label, r.text[:120]))
        self.assertEqual(a.put(base + "/stages/charter", json=[1, 2]).status_code, 400)
        self.assertEqual(a.put(base + "/stages/charter", content=b"{bad", headers={"content-type": "application/json"}).status_code, 400)
        for url in (base + "/stages/tech", base + "/stages/risk", base, "/api/dashboard", base + "/stages/risk/checks"):
            self.assertEqual(a.get(url).status_code, 200, url)
        self.assertEqual(a.post(base + "/docs", json={"kind": "plan"}).status_code, 200)

    def test_URS94_data_stored_before_validation_existed_cannot_brick_the_stage(self):
        a, mid, ids = make_project()
        c = rclient("alice")
        base = "/api/methods/%d" % mid
        tech = json.loads(raw("SELECT data FROM stages WHERE method_id=? AND key='tech'", (mid,))[0][0])
        tech["sections"]["matrix"] = "oops"
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='tech'", (json.dumps(tech), mid))
        risk = json.loads(raw("SELECT data FROM stages WHERE method_id=? AND key='risk'", (mid,))[0][0])
        risk["sections"]["fmea"] = ["abc", {"mode": "m", "S": "abc", "O": 1, "D": 1}, {"mode": "n", "S": 10 ** 400, "O": 1, "D": 1}, None]
        raw("UPDATE stages SET data=? WHERE method_id=?  AND key='risk'", (json.dumps(risk), mid))
        atp = json.loads(raw("SELECT data FROM stages WHERE method_id=? AND key='atp'", (mid,))[0][0])
        atp["sections"]["criteria"] = [{"characteristic": None, "criterion": {"a": 1}, "rationale": ["x"], "source": 5}, "str"]
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='atp'", (json.dumps(atp), mid))
        for url in (base, base + "/stages/tech", base + "/stages/risk", base + "/stages/atp", base + "/stages/plan", base + "/stages/risk/checks",
                    base + "/stages/atp/checks", "/api/dashboard", base + "/export"):
            r = c.get(url)
            self.assertEqual(r.status_code, 200, (url, r.text[:200]))
        self.assertEqual(c.post(base + "/docs", json={"kind": "report"}).status_code, 200)

    def test_URS94_bad_requests_get_400_json_never_500_or_a_stack_trace(self):
        """t3: ~25 requests returned HTTP 500 with a plain-text error page."""
        a, mid, ids = make_project(title="fuzz3")
        admin = rclient("admin")
        a = rclient("alice")
        base = "/api/methods/%d" % mid
        pid = project_id_of(a, mid)
        pbase = "/api/projects/%d" % pid
        eid = exps(a, mid)[0]["id"]
        reqs = [
            ("submit comment dict", a.post(base + "/stages/charter/submit", json={"comment": {"x": 1}})),
            ("withdraw reason int", a.post(base + "/stages/charter/withdraw", json={"reason": 5})),
            ("review decision list", a.post(base + "/stages/atp/review", json={"decision": [1], "password": PW, "revision": 1, "expected_hash": "x"})),
            ("review password int", a.post(base + "/stages/atp/review", json={"decision": "approve", "password": 12345})),
            ("reopen reason int", a.post(base + "/stages/charter/reopen", json={"reason": 5, "password": PW})),
            ("method patch name int", a.patch(base, json={"name": 5})),
            ("method patch dates dict", a.patch(base, json={"planned_start": {"a": 1}})),
            ("method patch status list", a.patch(base, json={"status": ["active"]})),
            ("method patch reason int", a.patch(base, json={"status": "archived", "reason": 5})),
            ("project patch name int", a.patch(pbase, json={"name": 5})),
            ("project patch dates dict", a.patch(pbase, json={"planned_start": {"a": 1}})),
            ("project patch status int", a.patch(pbase, json={"status": 5})),
            ("project members non-list", a.put(pbase + "/members", json={"members": "abc"})),
            ("project members [str]", a.put(pbase + "/members", json={"members": ["abc"]})),
            ("project members role int", a.put(pbase + "/members", json={"members": [{"user_id": IDS["rita"], "role": 5}]})),
            ("project members user_id str", a.put(pbase + "/members", json={"members": [{"user_id": "x", "role": "owner"}]})),
            ("project members unknown user", a.put(pbase + "/members", json={"members": [{"user_id": 99999, "role": "member"}]})),
            ("members non-list", a.put(base + "/members", json={"members": "abc"})),
            ("members [str]", a.put(base + "/members", json={"members": ["abc"]})),
            ("members missing user_id", a.put(base + "/members", json={"members": [{"role": "owner"}]})),
            ("members user_id str", a.put(base + "/members", json={"members": [{"role": "owner", "user_id": "x"}]})),
            ("project create name int", a.post("/api/projects", json={"name": 5})),
            ("project create no name", a.post("/api/projects", json={})),
            ("project create planned_start bad", a.post("/api/projects", json={"name": "x", "planned_start": "garbage"})),
            ("project create planned_start dict", a.post("/api/projects", json={"name": "x", "planned_start": {"a": 1}})),
            ("project create members str", a.post("/api/projects", json={"name": "x", "members": "abc"})),
            ("project create members [1]", a.post("/api/projects", json={"name": "x", "members": [1]})),
            ("project create product dict", a.post("/api/projects", json={"name": "x", "product": {"a": 1}})),
            ("method create name int", a.post(pbase + "/methods", json={"name": 5})),
            ("method create no name", a.post(pbase + "/methods", json={})),
            ("method create template list", a.post(pbase + "/methods", json={"name": "x", "template_key": ["a"]})),
            ("method create template unknown", a.post(pbase + "/methods", json={"name": "x", "template_key": "nope"})),
            ("method create planned_start bad", a.post(pbase + "/methods", json={"name": "x", "planned_start": "garbage"})),
            ("method create members str", a.post(pbase + "/methods", json={"name": "x", "members": "abc"})),
            ("method create members [1]", a.post(pbase + "/methods", json={"name": "x", "members": [1]})),
            ("method create analyte dict", a.post(pbase + "/methods", json={"name": "x", "analyte": {"a": 1}})),
            ("method create on missing project", a.post("/api/projects/99999/methods", json={"name": "x"})),
            ("method list bad project_id", a.get("/api/methods?project_id=abc")),
            ("clone title int", a.post(base + "/clone", json={"title": 5})),
            ("clone project_id str", a.post(base + "/clone", json={"project_id": "x"})),
            ("clone project_id missing", a.post(base + "/clone", json={"project_id": 99999})),
            ("lifecycle list", a.post(base + "/lifecycle", json={"lifecycle": ["routine"], "reason": "r", "password": PW})),
            ("lifecycle bogus", a.post(base + "/lifecycle", json={"lifecycle": "bogus", "reason": "r", "password": PW})),
            ("lifecycle reason int", a.post(base + "/lifecycle", json={"lifecycle": "validation", "reason": 5, "password": PW})),
            ("lifecycle password int", a.post(base + "/lifecycle", json={"lifecycle": "validation", "reason": "r", "password": 5})),
            ("exp create title int", a.post(base + "/experiments", json={"title": 5})),
            ("exp create code int", a.post(base + "/experiments", json={"title": "t", "code": 5})),
            ("exp create owner str", a.post(base + "/experiments", json={"title": "t", "owner_id": "zz"})),
            ("exp create owner 9999", a.post(base + "/experiments", json={"title": "t", "owner_id": 9999})),
            ("exp create dates dict", a.post(base + "/experiments", json={"title": "t", "planned_start": {"a": 1}})),
            ("exp create stage list", a.post(base + "/experiments", json={"title": "t", "stage_key": ["a"]})),
            ("exp create design dict", a.post(base + "/experiments", json={"title": "t", "design": {"a": 1}})),
            ("exp put results str", a.put("/api/experiments/%d" % eid, json={"version": 1, "results": "x"})),
            ("exp put results [str]", a.put("/api/experiments/%d" % eid, json={"version": 1, "results": ["x"]})),
            ("exp put version str", a.put("/api/experiments/%d" % eid, json={"version": "x"})),
            ("exp status list", a.post("/api/experiments/%d/status" % eid, json={"status": ["x"]})),
            ("exp status reason int", a.post("/api/experiments/%d/status" % eid, json={"status": "cancelled", "reason": 5})),
            ("login username dict", rclient().post("/api/login", json={"username": {"a": 1}, "password": "x"})),
            ("login username int", rclient().post("/api/login", json={"username": 5, "password": "x"})),
            ("login password int", rclient().post("/api/login", json={"username": "alice", "password": 5})),
            ("login password None", rclient().post("/api/login", json={"username": "alice", "password": None})),
            ("login empty", rclient().post("/api/login", json={})),
            ("setup after init", rclient().post("/api/setup", json={})),
            ("change pw int", a.post("/api/me/password", json={"old": 1, "new": 2})),
            ("users create username int", admin.post("/api/users", json={"username": 5, "display_name": "x", "password": PW})),
            ("users create roles str", admin.post("/api/users", json={"username": "zzz1", "display_name": "x", "password": PW, "roles": "admin"})),
            ("users create roles int", admin.post("/api/users", json={"username": "zzz2", "display_name": "x", "password": PW, "roles": 5})),
            ("users create roles [int]", admin.post("/api/users", json={"username": "zzz2", "display_name": "x", "password": PW, "roles": [5]})),
            ("users create password None", admin.post("/api/users", json={"username": "zzz3", "display_name": "x", "password": None})),
            ("users create display int", admin.post("/api/users", json={"username": "zzz4", "display_name": 5, "password": PW})),
            ("users patch display int", admin.patch("/api/users/%d" % IDS["bob"], json={"display_name": 5})),
            ("users patch roles dict", admin.patch("/api/users/%d" % IDS["bob"], json={"roles": {"admin": 1}})),
            ("users patch active str", admin.patch("/api/users/%d" % IDS["bob"], json={"active": "no"})),
            ("users patch reset int", admin.patch("/api/users/%d" % IDS["bob"], json={"reset_password": 123456789012})),
            ("kb create tags int", a.post("/api/kb", json={"title": "t", "body": "b", "tags": 5})),
            ("kb create title int", a.post("/api/kb", json={"title": 5, "body": "b"})),
            ("lesson title int", a.post(base + "/lessons", json={"title": 5, "body": "b"})),
            ("docs kind list", a.post(base + "/docs", json={"kind": ["plan"]})),
            ("upload exp id huge", upload(a, mid, "a", b"x", experiment_id=99999999999999999999)),
            ("body is a list", a.post(base + "/stages/charter/submit", content=b"[1]", headers={"content-type": "application/json"})),
        ]
        self.assertEqual(a.get(base + "/audit?limit=99999999999999999999").status_code, 200)       # clamped, not an error
        for label, r in reqs:
            self.assertLess(r.status_code, 500, (label, r.status_code, r.text[:150]))
            self.assertGreaterEqual(r.status_code, 400, (label, r.status_code))
            self.assertNotIn("Traceback", r.text)
            self.assertIsInstance(r.json(), dict, label)           # always a JSON error object
        # bob was not changed by the rejected requests
        self.assertEqual([u["roles"] for u in admin.get("/api/users").json() if u["username"] == "bob"], [["author"]])


# ====================================================================== low findings
class TestMisc(unittest.TestCase):
    def test_URS95_experiment_owner_must_be_a_project_member(self):
        a, mid, ids = make_project()
        base = "/api/methods/%d" % mid
        self.assertEqual(a.post(base + "/experiments", json={"title": "t", "owner_id": IDS["bob"]}).status_code, 400)
        r = a.post(base + "/experiments", json={"title": "t", "owner_id": IDS["rita"]})
        self.assertEqual(r.status_code, 200, r.text)
        eid = r.json()["id"]
        self.assertEqual(exp_put(a, eid, owner_id=IDS["bob"]).status_code, 400)
        self.assertEqual(exp_put(a, eid, owner_id=IDS["alice"]).status_code, 200)
        self.assertEqual(exp_put(a, eid, owner_id="alice").status_code, 400)

    def test_URS95_audit_limit_is_clamped(self):
        a, mid, ids = make_project()
        base = "/api/methods/%d" % mid
        self.assertEqual(len(a.get(base + "/audit?limit=-5").json()), 1)
        self.assertEqual(len(a.get(base + "/audit?limit=0").json()), 1)
        self.assertEqual(len(a.get(base + "/audit?limit=-1").json()), 1)
        self.assertLessEqual(len(login("admin").get("/api/audit?limit=100000").json()), 2000)

    def test_URS95_staleness_propagates_transitively(self):
        """t6: re-opening the charter only staled its direct dependents."""
        a, mid, ids = make_project()
        approve_through_plan(a, mid)
        r = a.post("/api/methods/%d/stages/charter/reopen" % mid, json={"reason": "r", "password": PW})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreaterEqual({x["stage"] for x in r.json()["affected"]}, {"atp", "tech", "risk", "plan"})
        flags = {s["key"]: s["stale"] for s in a.get("/api/methods/%d" % mid).json()["stages"]}
        for k in ("atp", "tech", "risk", "plan"):
            self.assertTrue(flags[k], k)
        self.assertFalse(flags["charter"])

    def test_URS95_cookie_csp_and_host_header(self):
        c = rclient()
        r = c.post("/api/login", json={"username": "alice", "password": PW})
        sc = r.headers["set-cookie"]
        self.assertIn("HttpOnly", sc)
        self.assertIn("Max-Age=%d" % (config.get("session_hours") * 3600), sc)
        self.assertNotIn("Secure", sc)                                       # loopback deployment
        csp = r.headers["content-security-policy"]
        self.assertIn("default-src 'self'", csp)
        self.assertIn("style-src 'self' 'unsafe-inline'", csp)
        self.assertIn("img-src 'self' data:", csp)
        self.assertEqual(c.get("/", headers={}).headers["content-security-policy"], csp)
        with cfg(host="0.0.0.0"):
            self.assertIn("Secure", rclient().post("/api/login", json={"username": "alice", "password": PW}).headers["set-cookie"])
        with cfg(cookie_secure=True):
            self.assertIn("Secure", rclient().post("/api/login", json={"username": "alice", "password": PW}).headers["set-cookie"])
        with cfg(host="0.0.0.0", cookie_secure=False):
            self.assertNotIn("Secure", rclient().post("/api/login", json={"username": "alice", "password": PW}).headers["set-cookie"])
        # Host header validation (DNS rebinding)
        self.assertEqual(c.get("/api/state", headers={"Host": "evil.example"}).status_code, 400)
        self.assertEqual(c.get("/api/state", headers={"Host": "evil.example:8780"}).status_code, 400)
        self.assertEqual(c.get("/api/state", headers={"Host": "localhost:8780"}).status_code, 200)
        self.assertEqual(c.get("/api/state", headers={"Host": "127.0.0.1:8780"}).status_code, 200)
        self.assertEqual(c.get("/api/state", headers={"Host": "[::1]:8780"}).status_code, 200)
        with cfg(allowed_hosts=["mf.corp.example"]):
            self.assertEqual(c.get("/api/state", headers={"Host": "mf.corp.example:443"}).status_code, 200)
            self.assertEqual(c.get("/api/state", headers={"Host": "localhost"}).status_code, 400)
        with cfg(host="0.0.0.0"):                                              # LAN: anything unless configured
            self.assertEqual(c.get("/api/state", headers={"Host": "192.168.1.20:8780"}).status_code, 200)

    def test_URS95_user_list_hides_inactive_accounts_from_ordinary_users(self):
        admin = login("admin")
        uid_ = _mk_user(admin, "gone2")
        admin.patch("/api/users/%d" % uid_, json={"active": False})
        names = [u["username"] for u in login("alice").get("/api/users").json()]
        self.assertNotIn("gone2", names)
        self.assertIn("rita", names)
        self.assertIn("gone2", [u["username"] for u in admin.get("/api/users").json()])
        u = login("alice").get("/api/users").json()[0]
        self.assertGreaterEqual(set(u), {"id", "display_name", "roles"})
        self.assertNotIn("pw_hash", u)

    def test_URS95_authentication_reads_without_the_write_lock_and_throttles_last_seen(self):
        c = login("alice")
        tok = c.cookies.get("mf_session")
        h = db.sha256(tok)
        before = raw("SELECT last_seen FROM sessions WHERE token_hash=?", (h,))[0][0]
        self.assertEqual(c.get("/api/me").status_code, 200)
        self.assertEqual(raw("SELECT last_seen FROM sessions WHERE token_hash=?", (h,))[0][0], before)       # < 60 s: no write
        old = "2020-01-01T00:00:00Z"
        stamp = (time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime(time.time() - 300)))
        raw("UPDATE sessions SET last_seen=? WHERE token_hash=?", (stamp, h))
        self.assertEqual(c.get("/api/me").status_code, 200)
        self.assertNotEqual(raw("SELECT last_seen FROM sessions WHERE token_hash=?", (h,))[0][0], stamp)    # > 60 s: touched
        raw("UPDATE sessions SET last_seen=? WHERE token_hash=?", (old, h))
        self.assertEqual(c.get("/api/me").status_code, 401)                                                 # idle timeout still applies
        self.assertEqual(raw("SELECT COUNT(*) FROM sessions WHERE token_hash=?", (h,))[0][0], 0)
        # a long write transaction elsewhere no longer stalls authentication
        c = login("alice")
        holder = sqlite3.connect(str(config.DB_PATH), timeout=15, isolation_level=None)
        holder.execute("BEGIN IMMEDIATE")
        try:
            t0 = time.time()
            self.assertEqual(c.get("/api/me").status_code, 200)
            self.assertEqual(c.get("/api/dashboard").status_code, 200)
            self.assertLess(time.time() - t0, 2.0)
        finally:
            holder.execute("ROLLBACK")
            holder.close()

    def test_URS95_document_generation_needs_edit_or_review_role_and_is_rate_limited(self):
        a, mid, ids = make_project()
        vpid = make_project(title="V", team=["vic", "rita", "carol"], method_team=members_of([("vic", "viewer"), ("rita", "reviewer")]))[1]
        self.assertEqual(login("vic").post("/api/methods/%d/docs" % vpid, json={"kind": "plan"}).status_code, 403)
        self.assertEqual(login("carol").post("/api/methods/%d/docs" % vpid, json={"kind": "plan"}).status_code, 403)   # project member only
        self.assertEqual(login("bob").post("/api/methods/%d/docs" % vpid, json={"kind": "plan"}).status_code, 403)     # not a member
        self.assertEqual(login("rita").post("/api/methods/%d/docs" % vpid, json={"kind": "plan"}).status_code, 200)    # reviewer may
        self.assertEqual(login("admin").post("/api/methods/%d/docs" % vpid, json={"kind": "plan"}).status_code, 200)   # admin may
        self.assertEqual(login("vic").get("/api/methods/%d/docs" % vpid).status_code, 200)                             # viewers still list / download
        with cfg(docs_per_minute=2):
            r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
            self.assertEqual(r.status_code, 200)
            self.assertEqual(a.post("/api/methods/%d/docs" % mid, json={"kind": "report"}).status_code, 200)
            r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
            self.assertEqual(r.status_code, 429, r.text)
        did = a.get("/api/methods/%d/docs" % mid).json()[0]["id"]
        self.assertEqual(login("bob").get("/api/docs/%d/download" % did).status_code, 403)                              # outsider
        self.assertEqual(login("vic").get("/api/docs/%d/download" % did).status_code, 200)                              # project member: may read

    def test_URS97_dashboard_experiment_entries_keep_experiment_and_method_codes_apart(self):
        """Coordinator item: `SELECT e.*, p.code` made both columns `code`, so the owner's code was lost.  Now an
        entry carries the experiment `code`, the `method_code` and the `project_code`."""
        a, mid, ids = make_project()
        detail = a.get("/api/methods/%d" % mid).json()
        pcode = detail["code"]
        r = a.post("/api/methods/%d/experiments" % mid, json={"title": "我的实验 dash"})
        self.assertEqual(r.status_code, 200)
        ecode = a.get("/api/experiments/%d" % r.json()["id"]).json()["code"]
        with db.tx() as conn:       # experiments appear on the dashboard only once the plan is approved: set that state directly
            conn.execute("UPDATE stages SET status='approved' WHERE method_id=? AND key='plan'", (mid,))
        mine = [e for e in a.get("/api/dashboard").json()["my_experiments"] if e["title"] == "我的实验 dash"]
        with db.tx() as conn:
            conn.execute("UPDATE stages SET status='not_started' WHERE method_id=? AND key='plan'", (mid,))   # restore
        self.assertEqual(len(mine), 1)
        self.assertEqual(mine[0]["code"], ecode)
        self.assertEqual(mine[0]["method_code"], pcode)
        self.assertEqual(mine[0]["project_code"], detail["project"]["code"])
        self.assertEqual(len({mine[0]["code"], mine[0]["method_code"], mine[0]["project_code"]}), 3)
        self.assertTrue(pcode.startswith("AM-"))
        self.assertTrue(mine[0]["project_code"].startswith("PRJ-"))
        self.assertEqual(mine[0]["method_id"], mid)
        self.assertEqual(mine[0]["method"], detail["name"])


# ====================================================================== operations
class TestOperations(unittest.TestCase):
    def test_URS96_setup_token(self):
        tmp = tempfile.mkdtemp(prefix="mf_setup_")
        saved = config.DATA
        try:
            config.set_data(tmp)
            db.init_db()
            c = rclient()
            with cfg(setup_token="s3cret-token"):
                self.assertTrue(c.get("/api/state").json()["setup_token_required"])
                body = {"username": "boss", "display_name": "Boss", "password": PW}
                self.assertEqual(c.post("/api/setup", json=body).status_code, 403)
                self.assertEqual(c.post("/api/setup", json=dict(body, setup_token="nope")).status_code, 403)
                self.assertEqual(c.post("/api/setup", json=dict(body, setup_token=123)).status_code, 403)
                self.assertTrue(c.get("/api/state").json()["needs_setup"])
                r = c.post("/api/setup", json=dict(body, setup_token="s3cret-token"))
                self.assertEqual(r.status_code, 200, r.text)
                self.assertEqual(c.post("/api/setup", json=dict(body, setup_token="s3cret-token")).status_code, 409)
            self.assertFalse(c.get("/api/state").json()["setup_token_required"])
        finally:
            config.set_data(saved)
            shutil.rmtree(tmp, ignore_errors=True)

    def _data_copy(self, name):
        data = os.path.join(HTMP, name)
        shutil.rmtree(data, ignore_errors=True)
        os.makedirs(data)
        s = sqlite3.connect(str(config.DB_PATH)); d = sqlite3.connect(os.path.join(data, "methodflow.db")); s.backup(d); s.close(); d.close()
        for item in (".audit_key", "audit_head.json"):
            shutil.copy(str(config.DATA / item), os.path.join(data, item))
        return data

    def test_URS96_reset_admin_tool(self):
        data = self._data_copy("reset_copy")
        tool = os.path.join(ROOT, "tools", "reset_admin.py")
        free = socket.socket(); free.bind(("127.0.0.1", 0)); port = free.getsockname()[1]
        # refuses while the server answers on its port
        free.listen(1)
        p = subprocess.run([sys.executable, tool, "--data", data, "--yes", "--port", str(port)], capture_output=True)
        self.assertEqual(p.returncode, 3, p.stdout.decode("utf-8", "replace"))
        free.close()
        # refuses accounts that are not administrators
        p = subprocess.run([sys.executable, tool, "--data", data, "--yes", "--username", "alice", "--port", str(port)], capture_output=True)
        self.assertEqual(p.returncode, 1)
        p = subprocess.run([sys.executable, tool, "--data", data, "--yes", "--username", "admin", "--port", str(port)], capture_output=True)
        out = p.stdout.decode("utf-8", "replace")
        self.assertEqual(p.returncode, 0, out + p.stderr.decode("utf-8", "replace"))
        temp = [ln.strip() for ln in out.splitlines() if ln.startswith("    ")][0]
        self.assertGreaterEqual(len(temp), 12)
        saved = config.DATA
        try:
            config.set_data(data)
            auth.reset_limits()
            c = rclient()
            self.assertEqual(c.post("/api/login", json={"username": "admin", "password": PW}).status_code, 401)      # old password is dead
            r = c.post("/api/login", json={"username": "admin", "password": temp})
            self.assertEqual(r.status_code, 200, r.text)
            self.assertTrue(r.json()["user"]["must_change"])
            self.assertEqual(c.get("/api/dashboard").status_code, 403)                                                # must change first
            self.assertEqual(c.post("/api/me/password", json={"old": temp, "new": "BrandNewPassw0rd1"}).status_code, 200)
            with db.read() as conn:
                row = conn.execute("SELECT * FROM audit_log WHERE action='user.password_reset_offline'").fetchone()
                self.assertIsNotNone(row)
                self.assertEqual(row["username"], "system:reset_admin")
                res = audit.verify_chain(conn, live_anchor=True)
                self.assertTrue(res["ok"], res)
        finally:
            config.set_data(saved)
            auth.reset_limits()

    def test_URS96_audit_anchor_tool_rebuilds_a_lost_anchor_only_for_a_valid_chain(self):
        data = self._data_copy("anchor_copy")
        tool = os.path.join(ROOT, "tools", "audit_anchor.py")
        p = subprocess.run([sys.executable, tool, "--data", data], capture_output=True)
        self.assertEqual(p.returncode, 0, p.stdout.decode("utf-8", "replace"))
        os.remove(os.path.join(data, "audit_head.json"))
        p = subprocess.run([sys.executable, tool, "--data", data], capture_output=True)
        self.assertEqual(p.returncode, 2)                                     # lost anchor is reported
        p = subprocess.run([sys.executable, tool, "--data", data, "--reset", "--yes"], capture_output=True)
        self.assertEqual(p.returncode, 0, p.stdout.decode("utf-8", "replace"))
        p = subprocess.run([sys.executable, tool, "--data", data], capture_output=True)
        self.assertEqual(p.returncode, 0, p.stdout.decode("utf-8", "replace"))
        # a broken chain is never re-anchored
        con = sqlite3.connect(os.path.join(data, "methodflow.db"))
        for n in db.TRIGGER_NAMES:
            con.execute("DROP TRIGGER IF EXISTS %s" % n)
        con.execute("UPDATE audit_log SET details='{\"x\":1}' WHERE id=2")
        con.commit(); con.close()
        os.remove(os.path.join(data, "audit_head.json"))
        p = subprocess.run([sys.executable, tool, "--data", data, "--reset", "--yes"], capture_output=True)
        self.assertEqual(p.returncode, 4)

    def test_URS96_restore_tool_checks_the_manifest_and_reanchors(self):
        data = self._data_copy("restore_copy")
        a, mid, ids = make_project()
        # take a backup of the live data, then restore it into the copy
        name = workflow.backup_now("restoretest")
        shutil.copy(str(config.BACKUPS / name), os.path.join(data, name))
        shutil.copy(str(config.BACKUPS / (name + ".manifest.json")), os.path.join(data, name + ".manifest.json"))
        tool = os.path.join(ROOT, "tools", "restore.py")
        free = socket.socket(); free.bind(("127.0.0.1", 0)); port = free.getsockname()[1]; free.close()
        p = subprocess.run([sys.executable, tool, os.path.join(data, name), "--data", data, "--yes", "--port", str(port)], capture_output=True)
        out = p.stdout.decode("utf-8", "replace")
        self.assertEqual(p.returncode, 0, out + p.stderr.decode("utf-8", "replace"))
        self.assertIn("链头与旁注一致", out)
        self.assertIn("锚点已按恢复后的数据库重建", out)
        # a backup whose side-car says the log used to be longer is refused
        with open(os.path.join(data, name + ".manifest.json"), encoding="utf-8") as f:
            m = json.load(f)
        m["audit_head"]["count"] += 10
        with open(os.path.join(data, name + ".manifest.json"), "w", encoding="utf-8") as f:
            json.dump(m, f)
        p = subprocess.run([sys.executable, tool, os.path.join(data, name), "--data", data, "--yes", "--port", str(port)], capture_output=True)
        self.assertEqual(p.returncode, 4, p.stdout.decode("utf-8", "replace"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
