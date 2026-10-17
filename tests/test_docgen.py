"""Tests of the approval-level document exporter (Word + PDF): chapter-driven layout, controlled-document furniture,
signature page, draft / integrity-failure marking, settings, capabilities, PDF output.  Test names carry the URS id
(URS-120 ... URS-129, see docs/URS_需求与测试追溯.md).  Run:  python -m unittest tests.test_docgen -v
(The PDF tests need LibreOffice and are skipped when it is not installed.)"""
import contextlib
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import test_workflow as tw                       # sets METHODFLOW_DATA and imports the app (shared helpers)
from test_workflow import PW, client, login, make_project, approve_stage, fill_required  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from app import auth, config, db, docgen, main, templates_data, workflow  # noqa: E402

DTMP = tempfile.mkdtemp(prefix="mf_docgen_")
HAVE_SOFFICE = docgen.find_soffice("") is not None


def setUpModule():
    config.set_data(DTMP)
    config.override("pdf_enabled", False)           # most tests are about the Word file; PDF tests switch it on
    auth.reset_limits()
    db.init_db()
    main._seed_kb()
    c = client()
    r = c.post("/api/setup", json={"username": "admin", "display_name": "管理员", "password": PW})
    assert r.status_code == 200, r.text
    for u, name, roles in [("alice", "爱丽丝(开发)", ["author"]), ("bob", "鲍勃(开发)", ["author"]), ("rita", "丽塔(复核)", ["reviewer"]),
                           ("quinn", "奎因(QA)", ["approver"]), ("rob", "罗布(复核2)", ["reviewer"])]:
        r = c.post("/api/users", json={"username": u, "display_name": name, "password": PW, "roles": roles})
        assert r.status_code == 200, r.text
    with db.tx() as conn:
        conn.execute("UPDATE users SET must_change=0")


def tearDownModule():
    config.override("pdf_enabled", True)
    shutil.rmtree(DTMP, ignore_errors=True)


@contextlib.contextmanager
def cfg(**kw):
    old = {k: config.get(k) for k in kw}
    for k, v in kw.items():
        config.override(k, v)
    try:
        yield
    finally:
        for k, v in old.items():
            config.override(k, v)


def gen(c, mid, kind="plan"):
    r = c.post("/api/methods/%d/docs" % mid, json={"kind": kind})
    assert r.status_code == 200, r.text
    return r.json()


def load(c, res, fmt="docx"):
    d = next(x for x in res["docs"] if x["format"] == fmt)
    r = c.get("/api/docs/%d/download" % d["id"])
    assert r.status_code == 200, r.text
    return r


def docx_of(c, res):
    return Document(io.BytesIO(load(c, res).content))


def body_text(d):
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def header_text(d):
    out = []
    for s in d.sections:
        for t in s.header.tables:
            for row in t.rows:
                for cell in row.cells:
                    out.append(cell.text)
    return "\n".join(out)


def headings(d, level=None):
    out = []
    for p in d.paragraphs:
        m = re.match(r"Heading (\d)", p.style.name)
        if m and (level is None or int(m.group(1)) == level):
            out.append(p.text)
    return out


def raw(sql, args=()):
    con = sqlite3.connect(str(config.DB_PATH), timeout=15)
    con.row_factory = sqlite3.Row
    try:
        rows = con.execute(sql, args).fetchall()
        con.commit()
        return rows
    finally:
        con.close()


def approve_through_plan(a, mid):
    for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
        approve_stage(mid, k, a, with_approver=w)


class TestChapterDrivenLayout(unittest.TestCase):
    def test_URS120_plan_follows_the_template_structure_and_hides_editor_text(self):
        a, mid, _ = make_project("flow", title="结构-flow")
        d = docx_of(a, gen(a, mid, "plan"))
        h1 = headings(d, 1)
        tpl = templates_data.build_template("flow")
        want = ["%s  %s" % (c["no"], c["title"]) for c in tpl["doc_structure"]["plan"]]
        body = [h for h in h1 if re.match(r"^\d+  ", h)]
        self.assertEqual(body, want)                                  # 15 chapters, in template order, numbered
        h2 = headings(d, 2)
        self.assertIn("6.4  性能特性与判据", h2)
        self.assertIn("10.3  FMEA 风险表", h2)
        self.assertIn("11.3  实验清单与设计", h2)                       # auto block sits in the gap of chapter 11
        self.assertLess(h2.index("11.2  开发阶段计划"), h2.index("11.3  实验清单与设计"))
        self.assertLess(h2.index("11.3  实验清单与设计"), h2.index("11.4  统计分析计划"))
        for stage_title in ("立项与背景", "分析目标概况（ATP）"):
            self.assertNotIn(stage_title, [h for h in h1 if not h.startswith("附录")])   # stage titles are not chapters
        text = body_text(d)
        self.assertNotIn("内容对应方案第", text)                        # stage `desc` is editor text
        self.assertNotIn("阶段状态：", text)
        self.assertIn("附录 A  阶段检查清单状态", h1)
        self.assertIn("附录 C  引用依据", h1)

    def test_URS120_report_has_21_chapters_and_result_block(self):
        a, mid, _ = make_project("flow", title="结构-report")
        ex = next(e for e in a.get("/api/methods/%d/experiments" % mid).json() if e["stage_key"] == "explore")
        r = a.put("/api/experiments/%d" % ex["id"], json={"version": ex["version"], "conclusion": "结论-XYZ", "deviations": "偏离-ABC",
                                                           "results": [{"metric": "指标甲", "value": "0.4%", "criterion": "≤ 2%", "verdict": "达标"}]})
        self.assertEqual(r.status_code, 200, r.text)
        d = docx_of(a, gen(a, mid, "report"))
        body = [h for h in headings(d, 1) if re.match(r"^\d+  ", h)]
        self.assertEqual(len(body), 21)
        self.assertEqual(body[0], "1  摘要与结论")
        self.assertEqual(body[-1], "21  结论与批准")
        h2 = headings(d, 2)
        self.assertIn("13.4  实验结果与结论", h2)
        self.assertIn("21.2  签署页（电子签名记录）", h2)
        text = body_text(d)
        for needle in ("结论-XYZ", "偏离-ABC", "指标甲"):
            self.assertIn(needle, text)

    def test_URS121_empty_required_sections_and_example_values_are_flagged_in_drafts(self):
        a, mid, _ = make_project("hplc", title="待填写-hplc")
        d = docx_of(a, gen(a, mid, "plan"))
        text = body_text(d)
        self.assertIn(docgen.TODO_TEXT, text)                          # e.g. 1.1 目的与成功标准 is empty
        red = [r for p in d.paragraphs for r in p.runs if r.text == docgen.TODO_TEXT]
        self.assertTrue(red and all(r.font.color.rgb is not None and str(r.font.color.rgb) == "C0392B" for r in red))
        tagged = []
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        tagged += [r for r in p.runs if workflow.SAMPLE_WORD in r.text]
        self.assertTrue(tagged)                                        # seeded example values exist ...
        self.assertTrue(all(r.font.highlight_color is not None for r in tagged))     # ... and are highlighted
        self.assertTrue(docgen.DRAFT_HEADER in header_text(d))

    def test_URS121_final_documents_have_no_draft_marks(self):
        a, mid, _ = make_project("flow", title="终版-flow")
        approve_through_plan(a, mid)
        res = gen(a, mid, "plan")
        d = docx_of(a, res)
        self.assertNotIn("草稿", header_text(d))
        self.assertIn("状态：已批准版本", body_text(d))
        self.assertNotIn(docgen.TODO_TEXT, body_text(d))
        self.assertEqual(a.get("/api/methods/%d/docs" % mid).json()[0]["status_label"], "已批准版本")

    def test_URS122_tables_use_fixed_layout_with_explicit_grid_and_repeat_headers(self):
        a, mid, _ = make_project("flow", title="表格-flow")
        d = docx_of(a, gen(a, mid, "plan"))
        n_checked = 0
        for t in d.tables:
            tblPr = t._tbl.tblPr
            lay = tblPr.find(qn("w:tblLayout"))
            self.assertIsNotNone(lay)
            self.assertEqual(lay.get(qn("w:type")), "fixed")
            tw_ = int(tblPr.find(qn("w:tblW")).get(qn("w:w")))
            grid = [int(g.get(qn("w:w"))) for g in t._tbl.tblGrid.findall(qn("w:gridCol"))]
            self.assertTrue(abs(sum(grid) - tw_) <= len(grid), (sum(grid), tw_))
            n_checked += 1
        self.assertGreater(n_checked, 20)
        fm = next(t for t in d.tables if t.rows[0].cells[0].text == "#" and "失效模式" in t.rows[0].cells[1].text)
        self.assertEqual(len(fm.columns), 10)
        self.assertIsNotNone(fm.rows[0]._tr.trPr.find(qn("w:tblHeader")))
        widths = [int(g.get(qn("w:w"))) for g in fm._tbl.tblGrid.findall(qn("w:gridCol"))]
        self.assertGreater(widths[8], 3 * widths[1] // 2)                # the "控制措施" column is the widest
        # 10-column FMEA lives in a landscape section, the rest of the document is portrait
        orient = [s.orientation for s in d.sections]
        self.assertIn(1, orient)                                         # WD_ORIENT.LANDSCAPE == 1
        self.assertIn(0, orient)

    def test_URS123_fields_textarea_and_verdict_banner(self):
        config.override("enforce_stage_order", False); self.addCleanup(config.override, "enforce_stage_order", True)
        a, mid, _ = make_project("flow", title="结论-flow")
        st = a.get("/api/methods/%d/stages/report" % mid).json()
        r = a.put("/api/methods/%d/stages/report" % mid, json={"version": st["version"], "sections": {"verdict": {
            "overall": "部分达成（有处置计划）", "recommendation": "有条件进入验证", "conditions": "第一行条件\n第二行条件"}}})
        self.assertEqual(r.status_code, 200, r.text)
        d = docx_of(a, gen(a, mid, "report"))
        banner = next(t for t in d.tables if "对照 ATP 的总体结论：部分达成" in t.rows[0].cells[0].text)
        cell = banner.rows[0].cells[0]
        self.assertIn("建议：有条件进入验证", cell.text)
        self.assertIn("第一行条件", cell.text)
        self.assertGreaterEqual(len(cell.paragraphs), 2)
        # textarea field in a plain fields table keeps its line structure
        st = a.get("/api/methods/%d/stages/charter" % mid).json()
        self.assertTrue(st)

    def test_URS124_appendices_show_guidance_refs_and_reference_documents(self):
        a, mid, _ = make_project("flow", title="附录-flow")
        d = docx_of(a, gen(a, mid, "report"))
        text = body_text(d)
        self.assertIn("指引：", text)                                    # checklist guidance (appendix A)
        hdrs = [[c.text for c in t.rows[0].cells] for t in d.tables]
        self.assertIn(["检查项（含指引）", "依据", "要求", "状态", "说明/证据"], hdrs)
        self.assertIn(["类别", "文件名称", "文件编号/版本/日期", "适用章节/用途"], hdrs)     # C.1: the refdocs table
        self.assertIn("C.1  参考文件与法规依据（第 5 章）", headings(d, 2))
        self.assertIn("ICH Q14", text)

    def test_URS125_legacy_stage_definitions_without_ch_are_still_placed(self):
        a, mid, _ = make_project("flow", title="旧定义-flow")
        with db.tx() as conn:
            for s in workflow.list_stages(conn, mid):
                dfn = s["def"]
                for sec in dfn["sections"]:
                    sec.pop("ch", None)
                conn.execute("UPDATE stages SET def=? WHERE id=?", (json.dumps(dfn, ensure_ascii=False), s["id"]))
        d = docx_of(a, gen(a, mid, "plan"))
        self.assertIn("6.4  性能特性与判据", headings(d, 2))


class TestControlledDocument(unittest.TestCase):
    def test_URS126_cover_signature_page_revision_history_and_toc(self):
        a, mid, _ = make_project("flow", title="受控-flow")
        approve_through_plan(a, mid)
        stage = a.get("/api/methods/%d/stages/plan" % mid).json()
        hashes = [r["data_hash"][:12] for r in stage["reviews"]]
        self.assertEqual(len(hashes), 2)
        reason = "受控性测试：重新打开原因-QQQ"
        r = a.post("/api/methods/%d/stages/atp/reopen" % mid, json={"reason": reason, "password": PW})
        self.assertEqual(r.status_code, 200, r.text)
        d = docx_of(a, gen(a, mid, "plan"))
        text = body_text(d)
        # signature page: role / name / meaning / datetime / revision / hash(12) / integrity
        for needle in ("丽塔(复核)", "奎因(QA)", "复核：我已审阅", "批准：我批准", "R1"):
            self.assertIn(needle, text)
        for h in hashes:
            self.assertIn(h, text)
        self.assertEqual(sum(1 for t in d.tables if t.rows[0].cells[0].text == "签署角色"), 3)    # G1, G2, GP
        # revision history: stage submissions + the reopen reason
        self.assertIn(reason, text)
        self.assertIn("重新打开", text)
        self.assertIn("阶段提交", text)
        # cover table: prepared / reviewed / approved by
        cover = next(t for t in d.tables if [c.text for c in t.rows[0].cells] == ["职责", "姓名", "日期", "依据"])
        self.assertEqual([r.cells[0].text for r in cover.rows[1:]], ["编制", "复核", "批准"])
        self.assertIn("爱丽丝(开发)", cover.rows[1].cells[1].text)
        self.assertIn("丽塔(复核)", cover.rows[2].cells[1].text)
        # a TOC field + hyperlinked entries with bookmarks that exist
        xml = d.element.xml
        self.assertIn("TOC \\o", xml)
        anchors = set(re.findall(r'w:anchor="(_Toc\d+)"', xml))
        names = set(re.findall(r'w:name="(_Toc\d+)"', xml))
        self.assertTrue(anchors and anchors <= names)
        # header: company, title, doc number, version, status; footer: page X of Y
        ht = header_text(d)
        self.assertIn("分析方法开发方案", ht)
        self.assertIn("-PLAN-V1", ht)
        self.assertIn("版本：V1", ht)
        foot = "".join(s.footer._element.xml for s in d.sections)
        self.assertIn("NUMPAGES", foot)
        self.assertIn("PAGE", foot)
        self.assertIn("打印件未经控制", foot)
        # the atp reopen made the plan stale: the document is a draft now
        self.assertIn(docgen.DRAFT_HEADER, ht)

    def test_URS127_integrity_failure_is_shown_everywhere(self):
        a, mid, _ = make_project("flow", title="完整性-flow")
        approve_through_plan(a, mid)
        data = json.loads(raw("SELECT data FROM stages WHERE method_id=? AND key='atp'", (mid,))[0][0])
        data["sections"]["use"] = "DIRECT DB EDIT"
        raw("UPDATE stages SET data=? WHERE method_id=? AND key='atp'", (json.dumps(data, ensure_ascii=False), mid))
        res = gen(a, mid, "plan")
        d = docx_of(a, res)
        self.assertIn(docgen.INTEGRITY_HEADER, header_text(d))
        text = body_text(d)
        self.assertIn("完整性校验失败", text)
        self.assertNotIn("状态：已批准版本", text)
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual(lst[0]["status_label"], "完整性校验失败")
        self.assertFalse(lst[0]["signatures_ok"])

    def test_URS128_document_number_prefix_company_and_footer_come_from_settings(self):
        admin = login("admin")
        r = admin.put("/api/admin/settings", json={"company_name": "测试制药有限公司", "doc_prefix": "QC-", "footer_text": "受控文件 ZZZ"})
        self.assertEqual(r.status_code, 200, r.text)
        try:
            a, mid, _ = make_project("flow", title="设置-flow")
            code = a.get("/api/methods/%d" % mid).json()["code"]
            d = docx_of(a, gen(a, mid, "plan"))
            ht = header_text(d)
            self.assertIn("测试制药有限公司", ht)
            self.assertIn("QC-%s-PLAN-V1" % code, ht)
            self.assertIn("受控文件 ZZZ", "".join(s.footer._element.xml for s in d.sections))
            d2 = docx_of(a, gen(a, mid, "report"))
            self.assertIn("QC-%s-RPT-V1" % code, header_text(d2))
            d3 = docx_of(a, gen(a, mid, "plan"))
            self.assertIn("QC-%s-PLAN-V2" % code, header_text(d3))       # version counts documents of that kind
        finally:
            admin.put("/api/admin/settings", json={"company_name": "", "doc_prefix": "", "footer_text": "打印件未经控制"})


class TestSettingsAndCapabilities(unittest.TestCase):
    def test_URS128_settings_api_validation_permissions_and_audit(self):
        admin, alice = login("admin"), login("alice")
        self.assertEqual(alice.get("/api/admin/settings").status_code, 403)
        self.assertEqual(alice.put("/api/admin/settings", json={"company_name": "x"}).status_code, 403)
        g = admin.get("/api/admin/settings")
        self.assertEqual(g.status_code, 200)
        for k in ("company_name", "footer_text", "doc_prefix", "soffice_path", "pdf_available", "pdf_note"):
            self.assertIn(k, g.json())
        bad = [{"company_name": 5}, {"company_name": "x" * 101}, {"footer_text": ["a"]}, {"doc_prefix": "bad prefix!"}, {"doc_prefix": "A" * 21},
               {"soffice_path": "C:/definitely/not/here/soffice.exe"}, {"soffice_path": __file__}, {"unknown": "x"}, {"company_name": "a\nb"}]
        for body in bad:
            r = admin.put("/api/admin/settings", json=body)
            self.assertEqual(r.status_code, 400, (body, r.text))
        self.assertEqual(admin.put("/api/admin/settings", content=b"[1]", headers={"content-type": "application/json"}).status_code, 400)
        before = len(admin.get("/api/audit").json())
        r = admin.put("/api/admin/settings", json={"company_name": "  甲公司  ", "doc_prefix": "AB-1"})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["company_name"], "甲公司")
        self.assertEqual(r.json()["doc_prefix"], "AB-1")
        rows = [x for x in admin.get("/api/audit").json() if x["action"] == "settings.update"]
        self.assertTrue(rows)
        self.assertEqual(rows[0]["details"]["changes"]["company_name"]["new"], "甲公司")
        self.assertGreater(len(admin.get("/api/audit").json()), before)
        admin.put("/api/admin/settings", json={"company_name": "", "doc_prefix": ""})

    @unittest.skipUnless(HAVE_SOFFICE, "LibreOffice not installed")
    def test_URS128_soffice_path_accepts_an_existing_soffice(self):
        admin = login("admin")
        path = docgen.find_soffice("")
        r = admin.put("/api/admin/settings", json={"soffice_path": path})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["soffice_found"], path)
        admin.put("/api/admin/settings", json={"soffice_path": ""})

    def test_URS129_capabilities_for_any_logged_in_user(self):
        self.assertEqual(client().get("/api/system/capabilities").status_code, 401)
        a = login("alice")
        with cfg(pdf_enabled=True), mock.patch.object(docgen, "find_soffice", return_value=None):
            j = a.get("/api/system/capabilities").json()
            self.assertEqual(j["pdf"], False)
            self.assertIn("LibreOffice", j["pdf_note"])
        with cfg(pdf_enabled=True), mock.patch.object(docgen, "find_soffice", return_value="/x/soffice"):
            j = a.get("/api/system/capabilities").json()
            self.assertEqual(j, {"pdf": True, "pdf_note": ""})
        with cfg(pdf_enabled=False):
            self.assertFalse(a.get("/api/system/capabilities").json()["pdf"])


class TestPdfOutput(unittest.TestCase):
    def test_URS129_without_libreoffice_only_word_is_produced_and_nothing_crashes(self):
        a, mid, _ = make_project("flow", title="无PDF-flow")
        with cfg(pdf_enabled=True), mock.patch.object(docgen, "find_soffice", return_value=None):
            res = gen(a, mid, "plan")
        self.assertEqual([d["format"] for d in res["docs"]], ["docx"])
        self.assertFalse(res["pdf_available"])
        self.assertIn("LibreOffice", res["pdf_note"])
        self.assertEqual(res["gen_id"], res["docs"][0]["id"])
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual([x["format"] for x in lst], ["docx"])
        self.assertEqual(lst[0]["gen_id"], res["gen_id"])
        # the contents list has no page numbers, so Word is told to refresh the fields on open
        self.assertIn("updateFields", docx_of(a, res).settings.element.xml)

    def test_URS129_conversion_failure_and_timeout_are_reported_not_raised(self):
        a, mid, _ = make_project("flow", title="转换失败-flow")
        with cfg(pdf_enabled=True), mock.patch.object(docgen, "find_soffice", return_value="/x/soffice"), \
                mock.patch.object(docgen, "convert_pdf", return_value=(None, "PDF 转换失败（测试），已仅生成 Word 文件。")):
            res = gen(a, mid, "plan")
        self.assertFalse(res["pdf_available"])
        self.assertIn("转换失败", res["pdf_note"])
        self.assertEqual(len(res["docs"]), 1)
        # real convert_pdf: a program that does not exist / never finishes
        p = os.path.join(DTMP, "x.docx")
        with open(p, "wb") as f:
            f.write(b"x")
        pdf, note = docgen.convert_pdf(os.path.join(DTMP, "no-such-soffice.exe"), p)
        self.assertIsNone(pdf)
        self.assertIn("失败", note)

        class Hung(object):
            pid = 0
            returncode = None

            def __init__(self, *a, **k):
                self.n = 0

            def communicate(self, timeout=None):
                self.n += 1
                if self.n == 1:
                    import subprocess
                    raise subprocess.TimeoutExpired("soffice", timeout)
                return b"", b""
        with mock.patch.object(docgen.subprocess, "Popen", Hung), mock.patch.object(docgen, "_kill_tree"), cfg(pdf_timeout_sec=1):
            pdf, note = docgen.convert_pdf("soffice", p)
        self.assertIsNone(pdf)
        self.assertIn("超时", note)

    @unittest.skipUnless(HAVE_SOFFICE, "LibreOffice not installed")
    def test_URS129_word_and_pdf_are_produced_together(self):
        a, mid, _ = make_project("flow", title="PDF-flow")
        approve_through_plan(a, mid)
        with cfg(pdf_enabled=True):
            res = gen(a, mid, "plan")
        self.assertTrue(res["pdf_available"], res["pdf_note"])
        self.assertEqual([d["format"] for d in res["docs"]], ["docx", "pdf"])
        stems = {os.path.splitext(d["filename"])[0] for d in res["docs"]}
        self.assertEqual(len(stems), 1)                                   # same name, different extension
        pdf = load(a, res, "pdf")
        self.assertTrue(pdf.content.startswith(b"%PDF-"))
        self.assertEqual(pdf.headers["content-type"], "application/pdf")
        word = load(a, res, "docx")
        self.assertIn("wordprocessingml", word.headers["content-type"])
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual([x["format"] for x in lst], ["docx", "pdf"])
        self.assertEqual({x["gen_id"] for x in lst}, {res["gen_id"]})
        for k in ("id", "gen_id", "kind", "format", "filename", "version_label", "status_label", "signatures_ok", "created_at", "display_name", "data_hash"):
            self.assertIn(k, lst[0])
        import hashlib
        row = raw("SELECT summary FROM generated_docs WHERE id=?", (res["docs"][1]["id"],))[0][0]
        self.assertEqual(json.loads(row)["file_sha256"], hashlib.sha256(pdf.content).hexdigest())
        # the PDF text carries the chapters, the signature page and real contents page numbers
        import io as _io
        try:
            from pypdf import PdfReader
        except ImportError:
            return
        rd = PdfReader(_io.BytesIO(pdf.content))
        text = "\n".join(p.extract_text() for p in rd.pages)
        for needle in ("分析方法开发方案", "方案批准记录", "数据哈希", "目录"):
            self.assertIn(needle, text.replace(" ", ""))
        d = Document(_io.BytesIO(word.content))
        pages = re.findall(r"PAGEREF _Toc\d+ \\h </w:instrText>.*?<w:t>(\d*)</w:t>", d.element.xml, re.S)
        self.assertTrue(pages and all(p for p in pages))                  # page numbers were filled in
        # list ordering: newest generation first
        with cfg(pdf_enabled=True):
            res2 = gen(a, mid, "report")
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual([x["gen_id"] for x in lst][:2], [res2["gen_id"]] * 2)


class TestHardeningStillHolds(unittest.TestCase):
    def test_URS93_data_changed_during_rendering_is_409_and_files_are_removed(self):
        a, mid, _ = make_project("flow", title="并发-flow")
        orig = docgen.render
        b = login("alice")

        def changing(ctx):
            out = orig(ctx)
            ver = b.get("/api/methods/%d/stages/charter" % mid).json()["version"]
            b.put("/api/methods/%d/stages/charter" % mid, json={"version": ver, "sections": {"background": "changed while rendering"}})
            return out
        with mock.patch.object(docgen, "render", changing):
            r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        self.assertEqual(r.status_code, 409, r.text)
        d = config.DOCS / a.get("/api/methods/%d" % mid).json()["code"]
        self.assertEqual([f for f in os.listdir(str(d))], [])
        self.assertEqual(a.get("/api/methods/%d/docs" % mid).json(), [])

    def test_URS93_rate_limit_counts_generations_not_files(self):
        a, mid, _ = make_project("flow", title="频率-flow")
        with cfg(docs_per_minute=2):
            self.assertEqual(a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).status_code, 200)
            self.assertEqual(a.post("/api/methods/%d/docs" % mid, json={"kind": "report"}).status_code, 200)
            self.assertEqual(a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"}).status_code, 429)

    def test_URS93_other_templates_generate_for_both_kinds(self):
        for tk in ("hplc", "cell_based", "generic"):
            a, mid, _ = make_project(tk, title="全模板-" + tk)
            for kind in ("plan", "report"):
                res = gen(a, mid, kind)
                d = docx_of(a, res)
                n = len([h for h in headings(d, 1) if re.match(r"^\d+  ", h)])
                self.assertEqual(n, 15 if kind == "plan" else 21, (tk, kind))


if __name__ == "__main__":
    unittest.main()
