"""Automated tests for MethodFlow's core rules.  Each test is tagged with the URS id it supports
(see docs/URS_需求与测试追溯.md).  Run:  python -m unittest discover -s tests -v"""
import io
import json
import re
import os
import shutil
import sqlite3
import sys
import tempfile
import unittest
import zipfile

TMP = tempfile.mkdtemp(prefix="mf_test_")
os.environ["METHODFLOW_DATA"] = TMP
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient  # noqa: E402

from app import audit, auth, config, db, main  # noqa: E402

PW = "Passw0rd!x1"


def client():
    c = TestClient(main.app)
    c.headers.update({"X-Requested-With": "MethodFlow"})
    return c


def login(username, pw=PW):
    c = client()
    r = c.post("/api/login", json={"username": username, "password": pw})
    assert r.status_code == 200, r.text
    return c


def setUpModule():
    config.set_data(TMP)          # each test module owns its own database (modules share one process)
    config.override("pdf_enabled", False)      # PDF output has its own tests (test_docgen.py); here it would only slow things down
    auth.reset_limits()
    db.init_db()
    main._seed_kb()
    c = client()
    r = c.post("/api/setup", json={"username": "admin", "display_name": "管理员", "password": PW})
    assert r.status_code == 200, r.text
    for u, name, roles in [("alice", "爱丽丝(开发)", ["author"]), ("bob", "鲍勃(开发)", ["author"]),
                           ("rita", "丽塔(复核)", ["reviewer"]), ("quinn", "奎因(QA)", ["approver"]),
                           ("rob", "罗布(复核2)", ["reviewer"])]:
        r = c.post("/api/users", json={"username": u, "display_name": name, "password": PW, "roles": roles})
        assert r.status_code == 200, r.text
    with db.tx() as conn:      # accounts created by an admin must change their password first (M3); these have "done so"
        conn.execute("UPDATE users SET must_change=0")


def tearDownModule():
    shutil.rmtree(TMP, ignore_errors=True)


def uid(c, username):
    return next(u["id"] for u in c.get("/api/users").json() if u["username"] == username)


DEFAULT_TEAM = ["rita", "quinn", "rob", "dora", "carol", "vic", "multi"]     # project members besides alice (owner); bob is the outsider


def make_project(technique="flow", title="CAR-T 流式测试", team=None, method_team=None):
    """Creates a PROJECT (alice = project owner; `team` = project members) and one METHOD in it (alice = method
    owner + author, rita / rob reviewers, quinn approver unless `method_team` says otherwise).
    Returns (alice's client, METHOD id, {username: user id}); the project id is `project_id_of(a, mid)`."""
    a = login("alice")
    ids = {u["username"]: u["id"] for u in a.get("/api/users").json()}
    names = [n for n in (DEFAULT_TEAM if team is None else team) if n in ids]
    r = a.post("/api/projects", json={"name": "项目-" + title, "product": "CAR-T", "indication": "B-ALL",
                                      "members": [{"user_id": ids[n], "role": "member"} for n in names]})
    assert r.status_code == 200, r.text
    pid = r.json()["id"]
    members = method_team if method_team is not None else [{"user_id": ids["rita"], "role": "reviewer"}, {"user_id": ids["quinn"], "role": "approver"},
                                                           {"user_id": ids["rob"], "role": "reviewer"}]
    r = a.post("/api/projects/%d/methods" % pid, json={"name": title, "template_key": technique, "analyte": "CAR+ %", "matrix": "全血",
                                                        "intended_use": "放行检测", "members": members})
    assert r.status_code == 200, r.text
    return a, r.json()["id"], ids


def project_id_of(c, mid):
    return c.get("/api/methods/%d" % mid).json()["project_id"]


def fill_required(c, mid, key, extra_sections=None):
    st = c.get("/api/methods/%d/stages/%s" % (mid, key)).json()
    secs = {}
    for sec in st["def"]["sections"]:
        if not sec.get("required"):
            continue
        t = sec["type"]
        cur = st["data"]["sections"].get(sec["key"])
        if t == "text":
            secs[sec["key"]] = "内容-" + sec["key"]
        elif t == "fields":
            secs[sec["key"]] = {f["key"]: ("v-" + f["key"] if f["type"] != "select" else (f["options"][0] if f["options"] else "")) for f in sec["fields"]}
        elif t == "table":
            rows = []
            need = max(1, sec.get("min_rows", 1))
            for i in range(need):
                rows.append({cc["key"]: ((cc["options"][0] if cc["options"] else "") if cc["type"] == "select" else ("12" if cc["type"] == "number" else "值%d-%s" % (i, cc["key"]))) for cc in sec["columns"]})
            secs[sec["key"]] = rows
        elif t == "matrix":
            m = cur
            for cr in m["criteria"]:
                cr["scores"] = {o: 3 for o in m["options"]}
            m["selected"] = m["options"][0]
            secs[sec["key"]] = m
        elif t == "fmea":
            secs[sec["key"]] = [dict(r, control=r.get("control") or "控制措施") for r in cur] or [{"mode": "x", "S": 5, "O": 3, "D": 3, "control": "c"}]
    if extra_sections:
        secs.update(extra_sections)
    cl = {c_["key"]: {"status": "done", "note": ("证据-" + c_["key"]) if c_.get("evidence") else ""} for c_ in st["def"]["checklist"]}
    r = c.put("/api/methods/%d/stages/%s" % (mid, key), json={"version": st["version"], "sections": secs, "checklist": cl})
    assert r.status_code == 200, r.text
    if CONFIRM_PREFILL:
        r = confirm_prefill(c, mid, key)
    return r.json()


CONFIRM_PREFILL = True


def _scrub(value):
    """What a careful user does with template example values: replace them with product content (here: drop the marker)."""
    s = json.dumps(value, ensure_ascii=False)
    s = re.sub(r"[（(]?示例值[^）)；;\"]*[）)]?", "", s)
    s = re.sub(r"(?<![A-Za-z])XX+(?![A-Za-z])", "示范", s)       # template placeholders ("XX 实验室") are replaced by product content
    return json.loads(s)


def confirm_prefill(c, mid, key):
    """Confirm every template pre-filled section of the stage (after scrubbing the example markers); for the plan
    stage also scrub the experiments' example acceptance texts (they are frozen once the plan is submitted)."""
    url = "/api/methods/%d/stages/%s" % (mid, key)
    st = c.get(url).json()
    unc = list((st["data"].get("unconfirmed") or []))
    if unc:
        secs = {k: _scrub(st["data"]["sections"].get(k)) for k in unc}
        r = c.put(url, json={"version": st["version"], "sections": secs, "confirm_sections": unc})
        assert r.status_code == 200, r.text
    if key == "plan":
        for e in c.get("/api/methods/%d/experiments" % mid).json():
            if any("示例值" in (e.get(f) or "") for f in ("objective", "design", "acceptance")):
                cur = c.get("/api/experiments/%d" % e["id"]).json()
                body = {"version": cur["version"]}
                for f in ("objective", "design", "acceptance"):
                    if "示例值" in (e.get(f) or ""):
                        body[f] = _scrub(e[f]) or "达到预设判据"
                r = c.put("/api/experiments/%d" % e["id"], json=body)
                assert r.status_code == 200, r.text
    return c.get(url)


def sign(c, mid, key, decision="approve", comment="OK", pw=PW):
    """Sign what the reviewer currently sees: the request must carry the stage's revision and hash (H5)."""
    st = c.get("/api/methods/%d/stages/%s" % (mid, key)).json()
    return c.post("/api/methods/%d/stages/%s/review" % (mid, key),
                  json={"decision": decision, "comment": comment, "password": pw, "revision": st["revision"], "expected_hash": st["hash"]})


def approve_stage(mid, key, a, with_approver=False):
    fill_required(a, mid, key)
    r = a.post("/api/methods/%d/stages/%s/submit" % (mid, key), json={})
    assert r.status_code == 200, (key, r.text)
    st = a.get("/api/methods/%d/stages/%s" % (mid, key)).json()
    if st["status"] == "completed":
        return
    assert sign(login("rita"), mid, key).status_code == 200
    if with_approver:
        assert sign(login("quinn"), mid, key).status_code == 200
    assert a.get("/api/methods/%d/stages/%s" % (mid, key)).json()["status"] == "approved"


# ====================================================================== tests
class TestAuth(unittest.TestCase):
    def test_URS01_password_policy_and_lockout(self):
        c = client()
        self.assertEqual(c.post("/api/login", json={"username": "alice", "password": "bad"}).status_code, 401)
        # weak password cannot be set
        a = login("alice")
        r = a.post("/api/me/password", json={"old": PW, "new": "short1"})
        self.assertEqual(r.status_code, 400)
        # lockout after N failures, even if correct password afterwards
        for _ in range(config.get("max_failed_logins")):
            c.post("/api/login", json={"username": "bob", "password": "wrong-pass-1"})
        r = c.post("/api/login", json={"username": "bob", "password": PW})
        self.assertEqual(r.status_code, 423)
        with db.tx() as conn:  # unlock for other tests
            conn.execute("UPDATE users SET locked_until=NULL, failed_count=0 WHERE username='bob'")

    def test_URS02_session_required(self):
        self.assertEqual(client().get("/api/dashboard").status_code, 401)

    def test_URS03_csrf_header_required(self):
        c = TestClient(main.app)
        self.assertEqual(c.post("/api/login", json={"username": "alice", "password": PW}).status_code, 403)


class TestGateFlow(unittest.TestCase):
    def test_URS10_required_content_blocks_submission(self):
        a, mid, _ = make_project()
        r = a.post("/api/methods/%d/stages/atp/submit" % mid, json={})
        self.assertEqual(r.status_code, 422)
        self.assertTrue(r.json()["detail"]["blocks"])

    def test_URS11_review_reject_resubmit_approve_and_lock(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)          # no gate -> completed
        fill_required(a, mid, "atp")
        self.assertEqual(a.post("/api/methods/%d/stages/atp/submit" % mid, json={}).status_code, 200)
        rita = login("rita")
        # wrong password -> signature auth fails and nothing is signed
        self.assertEqual(sign(rita, mid, "atp", pw="nope-nope-1").status_code, 403)
        # reject needs a reason
        self.assertEqual(sign(rita, mid, "atp", "reject", "").status_code, 400)
        self.assertEqual(sign(rita, mid, "atp", "reject", "判据缺依据").status_code, 200)
        st = a.get("/api/methods/%d/stages/atp" % mid).json()
        self.assertEqual((st["status"], st["revision"]), ("returned", 2))
        # fix and resubmit
        r = a.put("/api/methods/%d/stages/atp" % mid, json={"version": st["version"], "sections": {"use": "修订后的用途"}})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(a.post("/api/methods/%d/stages/atp/submit" % mid, json={}).status_code, 200)
        self.assertEqual(sign(rita, mid, "atp").status_code, 200)
        st = a.get("/api/methods/%d/stages/atp" % mid).json()
        self.assertEqual(st["status"], "approved")
        self.assertEqual(len([x for x in st["reviews"] if x["decision"] == "reject"]), 1)
        self.assertTrue(all(x["intact"] for x in st["reviews"]))
        # locked: edits rejected
        r = a.put("/api/methods/%d/stages/atp" % mid, json={"version": st["version"], "sections": {"use": "偷改"}})
        self.assertEqual(r.status_code, 409)

    def test_URS12_separation_of_duties(self):
        """Somebody who contributed to a stage can never review it - even after the project roles are changed
        (a person may not hold an editing and a signing role at the same time, so the history must decide)."""
        admin = login("admin")
        a, mid, ids = make_project(team=DEFAULT_TEAM + ["bob"])         # bob must belong to the project to join the method team
        approve_stage(mid, "charter", a)
        base = [{"user_id": ids["alice"], "role": "owner"}, {"user_id": ids["alice"], "role": "author"}]
        base_url = "/api/methods/%d" % mid
        # bob is globally an author AND a reviewer (allowed globally; not in one method team at the same time)
        self.assertEqual(admin.patch("/api/users/%d" % ids["bob"], json={"roles": ["author", "reviewer"]}).status_code, 200)
        try:
            self.assertEqual(a.put(base_url + "/members", json={"members": base + [{"user_id": ids["bob"], "role": "author"}]}).status_code, 200)
            fill_required(a, mid, "atp")
            bob = login("bob")
            ver = a.get(base_url + "/stages/atp").json()["version"]
            self.assertEqual(bob.put(base_url + "/stages/atp", json={"version": ver, "sections": {"specificity": "bob 写的"}}).status_code, 200)
            # bob is moved from author to reviewer: only the history says he edited this stage
            self.assertEqual(a.put(base_url + "/members", json={"members": base + [{"user_id": ids["bob"], "role": "reviewer"}]}).status_code, 200)
            r = a.post(base_url + "/stages/atp/submit", json={})
            self.assertEqual(r.status_code, 422)  # the only reviewer (bob) contributed -> no eligible reviewer
            self.assertTrue(any("没有可用的复核人" in b for b in r.json()["detail"]["blocks"]))
        finally:
            admin.patch("/api/users/%d" % ids["bob"], json={"roles": ["author"]})

    def test_URS13_author_cannot_sign_own_gate(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)
        fill_required(a, mid, "atp")
        a.post("/api/methods/%d/stages/atp/submit" % mid, json={})
        self.assertEqual(sign(a, mid, "atp").status_code, 403)

    def test_URS14_two_step_gate_and_downstream_dependency(self):
        a, mid, _ = make_project()
        for k in ("charter",):
            approve_stage(mid, k, a)
        approve_stage(mid, "atp", a)
        approve_stage(mid, "tech", a)
        approve_stage(mid, "risk", a)
        fill_required(a, mid, "plan")
        self.assertEqual(a.post("/api/methods/%d/stages/plan/submit" % mid, json={}).status_code, 200)
        # approver cannot sign before reviewer
        self.assertEqual(sign(login("quinn"), mid, "plan").status_code, 403)
        self.assertEqual(sign(login("rita"), mid, "plan").status_code, 200)
        self.assertEqual(a.get("/api/methods/%d/stages/plan" % mid).json()["status"], "in_review")
        # same person cannot sign twice (rita is not an approver anyway)
        self.assertEqual(sign(login("rita"), mid, "plan").status_code, 403)
        self.assertEqual(sign(login("quinn"), mid, "plan").status_code, 200)
        self.assertEqual(a.get("/api/methods/%d/stages/plan" % mid).json()["status"], "approved")

    def test_URS15_prerequisites_enforced(self):
        a, mid, _ = make_project()
        # standard flow: a later step cannot even be edited before its prerequisite step passed
        st = a.get("/api/methods/%d/stages/tech" % mid).json()
        self.assertFalse(st["editable"])
        self.assertEqual([b["step"] for b in st["blocked_by"]], [2])
        r = a.put("/api/methods/%d/stages/tech" % mid, json={"version": st["version"], "sections": {"decision": "x"}})
        self.assertEqual(r.status_code, 409)
        self.assertIn("尚未开放", r.json()["error"])
        # and (defence in depth, with the order rule switched off) it cannot be submitted either
        config.override("enforce_stage_order", False)
        self.addCleanup(config.override, "enforce_stage_order", True)
        fill_required(a, mid, "tech")
        r = a.post("/api/methods/%d/stages/tech/submit" % mid, json={})
        self.assertEqual(r.status_code, 422)
        self.assertTrue(any("前置阶段" in b for b in r.json()["detail"]["blocks"]))

    def test_URS130_stage_opens_only_after_prerequisites_pass(self):
        a, mid, _ = make_project()
        m = a.get("/api/methods/%d" % mid).json()
        blocked = {s["key"]: [b["step"] for b in s["blocked_by"]] for s in m["stages"]}
        self.assertEqual(blocked["charter"], [])
        self.assertEqual(blocked["atp"], [1])
        self.assertEqual(blocked["plan"], [3, 4])
        approve_stage(mid, "charter", a)
        m = a.get("/api/methods/%d" % mid).json()
        self.assertEqual([s["blocked_by"] for s in m["stages"] if s["key"] == "atp"][0], [])
        self.assertTrue(a.get("/api/methods/%d/stages/atp" % mid).json()["editable"])
        self.assertTrue(any(s["blocked_by"] for s in m["stages"] if s["key"] == "explore"))

    def test_URS131_progressive_completion_core_blocks_rest_warns(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)
        url = "/api/methods/%d/stages/atp" % mid
        # core missing (预期用途 / 被测属性 empty) -> still blocked
        ck = a.get(url + "/checks").json()
        self.assertTrue(any("预期用途" in b for b in ck["blocks"]), ck)
        # only the core filled in: other required items are listed as incomplete but do not block
        st = a.get(url).json()
        r = a.put(url, json={"version": st["version"], "sections": {"use": "放行检测", "attribute": "CAR+ %"}})
        self.assertEqual(r.status_code, 200, r.text)
        ck = a.get(url + "/checks").json()
        self.assertTrue(any("性能特性与判据" in b and "尚未确认" in b for b in ck["blocks"]), ck)   # pre-filled core table is unconfirmed
        confirm_prefill(a, mid, "atp")
        ck = a.get(url + "/checks").json()
        self.assertEqual(ck["blocks"], [], ck)
        self.assertTrue(ck["incomplete"], ck)
        self.assertTrue(any("补充内容未填写" in w for w in ck["warns"]))
        # strict mode brings the old "everything required must be complete" behaviour back
        config.override("strict_completeness", True)
        try:
            self.assertTrue(a.get(url + "/checks").json()["blocks"])
        finally:
            config.override("strict_completeness", False)
        r = a.post(url + "/submit", json={})
        self.assertEqual(r.status_code, 200, r.text)
        with db.read() as conn:
            row = conn.execute("SELECT details FROM audit_log WHERE action='stage.submit' AND method_id=? ORDER BY id DESC LIMIT 1", (mid,)).fetchone()
        self.assertIn("incomplete", row["details"])

    def test_URS132_template_prefill_must_be_confirmed(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)
        url = "/api/methods/%d/stages/atp" % mid
        st = a.get(url).json()
        self.assertIn("criteria", st["data"]["unconfirmed"])
        before = st["progress"]
        # confirming while example markers remain is refused
        r = a.put(url, json={"version": st["version"], "confirm_sections": ["criteria"]})
        self.assertEqual(r.status_code, 400, r.text)
        self.assertIn("示例值", r.json()["error"])
        # scrub + confirm works, is audited, and the pre-fill then counts as progress
        r = confirm_prefill(a, mid, "atp")
        self.assertNotIn("criteria", r.json()["data"]["unconfirmed"])
        self.assertGreaterEqual(r.json()["progress"], before)
        with db.read() as conn:
            row = conn.execute("SELECT details FROM audit_log WHERE action='stage.save' AND method_id=? ORDER BY id DESC LIMIT 1", (mid,)).fetchone()
        self.assertIn("confirm:criteria", row["details"])
        # a stage nobody touched shows no progress for its template pre-fill
        m = a.get("/api/methods/%d" % mid).json()
        self.assertEqual([s["progress"] for s in m["stages"] if s["key"] == "risk"][0] < 20, True)

    def test_URS133_document_with_unconfirmed_content_is_not_an_approved_version(self):
        a, mid, _ = make_project()
        for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
            approve_stage(mid, k, a, with_approver=w)
        a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        self.assertEqual(a.get("/api/methods/%d/docs" % mid).json()[0]["status_label"], "已批准版本")
        # add an experiment after the plan whose acceptance still carries an example marker
        r = a.post("/api/methods/%d/experiments" % mid, json={"title": "追加", "stage_key": "explore", "design": "d", "acceptance": "SI≥5（示例值）"})
        self.assertEqual(r.status_code, 200, r.text)
        a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        self.assertIn("待确认版", a.get("/api/methods/%d/docs" % mid).json()[0]["status_label"])

    def test_URS16_fmea_high_risk_requires_control(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)
        approve_stage(mid, "atp", a)
        fill_required(a, mid, "risk")
        st = a.get("/api/methods/%d/stages/risk" % mid).json()
        rows = st["data"]["sections"]["fmea"]
        rows[0]["control"] = ""
        rows[0]["S"], rows[0]["O"], rows[0]["D"] = 9, 6, 5
        a.put("/api/methods/%d/stages/risk" % mid, json={"version": st["version"], "sections": {"fmea": rows}})
        r = a.post("/api/methods/%d/stages/risk/submit" % mid, json={})
        self.assertEqual(r.status_code, 422)
        self.assertTrue(any("必须填写控制措施" in b for b in r.json()["detail"]["blocks"]))

    def test_URS17_optimistic_concurrency(self):
        a, mid, _ = make_project()
        st = a.get("/api/methods/%d/stages/charter" % mid).json()
        r1 = a.put("/api/methods/%d/stages/charter" % mid, json={"version": st["version"], "sections": {"background": "A"}})
        r2 = a.put("/api/methods/%d/stages/charter" % mid, json={"version": st["version"], "sections": {"background": "B"}})
        self.assertEqual((r1.status_code, r2.status_code), (200, 409))


class TestReopenAndStale(unittest.TestCase):
    def test_URS18_reopen_invalidates_dependents_and_recertify(self):
        a, mid, _ = make_project()
        approve_stage(mid, "charter", a)
        approve_stage(mid, "atp", a)
        approve_stage(mid, "tech", a)
        # reopen atp (owner) with reason + password
        r = a.post("/api/methods/%d/stages/atp/reopen" % mid, json={"reason": "判据收紧", "password": PW})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertTrue(any(x["stage"] == "tech" for x in r.json()["affected"]))
        tech = a.get("/api/methods/%d/stages/tech" % mid).json()
        self.assertTrue(tech["stale"])
        # draft doc flags unapproved content; later gates cannot be submitted
        self.assertEqual(a.post("/api/methods/%d/stages/atp/reopen" % mid, json={"reason": "", "password": PW}).status_code, 400)
        # re-approve atp then recertify tech
        a.put("/api/methods/%d/stages/atp" % mid, json={"version": a.get("/api/methods/%d/stages/atp" % mid).json()["version"], "sections": {"use": "新用途"}})
        a.post("/api/methods/%d/stages/atp/submit" % mid, json={})
        sign(login("rita"), mid, "atp")
        r = a.post("/api/methods/%d/stages/tech/recertify" % mid, json={})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(sign(login("rita"), mid, "tech").status_code, 200)
        self.assertFalse(a.get("/api/methods/%d/stages/tech" % mid).json()["stale"])


class TestExperiments(unittest.TestCase):
    def test_URS20_experiments_need_approved_plan_and_conclusion(self):
        a, mid, _ = make_project()
        exps = a.get("/api/methods/%d/experiments" % mid).json()
        self.assertGreaterEqual(len(exps), 5)
        e = exps[0]
        r = a.post("/api/experiments/%d/status" % e["id"], json={"status": "running"})
        self.assertEqual(r.status_code, 409)       # plan not approved
        for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
            approve_stage(mid, k, a, with_approver=w)
        self.assertEqual(a.post("/api/experiments/%d/status" % e["id"], json={"status": "running"}).status_code, 200)
        self.assertEqual(a.post("/api/experiments/%d/status" % e["id"], json={"status": "done"}).status_code, 400)  # no conclusion
        cur = a.get("/api/experiments/%d" % e["id"]).json()
        a.put("/api/experiments/%d" % e["id"], json={"version": cur["version"], "conclusion": "达标", "results": [{"metric": "SI", "value": "46", "criterion": "≥5", "verdict": "符合"}]})
        self.assertEqual(a.post("/api/experiments/%d/status" % e["id"], json={"status": "done"}).status_code, 200)

    def test_URS21_stage_needs_all_experiments_closed(self):
        a, mid, _ = make_project()
        for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
            approve_stage(mid, k, a, with_approver=w)
        fill_required(a, mid, "explore")
        r = a.post("/api/methods/%d/stages/explore/submit" % mid, json={})
        self.assertEqual(r.status_code, 422)
        self.assertTrue(any("尚未完成或取消" in b for b in r.json()["detail"]["blocks"]))


class TestAuditAndFiles(unittest.TestCase):
    def test_URS30_audit_chain_detects_tampering(self):
        a, mid, _ = make_project()
        admin = login("admin")
        self.assertTrue(admin.get("/api/audit/verify").json()["ok"])
        con = sqlite3.connect(str(config.DB_PATH))
        for name in db.TRIGGER_NAMES:      # the append-only triggers stop ordinary SQL; an attacker would drop them first
            con.execute("DROP TRIGGER IF EXISTS %s" % name)
        con.execute("UPDATE audit_log SET details='{\"x\":1}' WHERE id=(SELECT MIN(id) FROM audit_log WHERE action='project.create')")
        con.commit(); con.close()
        res = admin.get("/api/audit/verify").json()
        self.assertFalse(res["ok"])
        self.assertIsNotNone(res["first_bad_id"])
        # restore chain validity for other tests: recompute by rebuilding from scratch is not possible -> use fresh DB copy
        TestAuditAndFiles._tampered = True

    def test_URS31_attachments_hash_and_lock(self):
        a, mid, _ = make_project()
        body = b"flow-data-" * 1000
        r = a.post("/api/methods/%d/attachments?filename=run1.fcs&stage_key=charter" % mid, content=body)
        self.assertEqual(r.status_code, 200, r.text)
        aid = r.json()["id"]
        d = a.get("/api/attachments/%d/download" % aid)
        self.assertEqual(d.content, body)
        self.assertTrue(a.post("/api/attachments/%d/verify" % aid).json()["intact"])
        # tamper on disk -> verification fails
        with db.read() as conn:
            path = conn.execute("SELECT path FROM attachments WHERE id=?", (aid,)).fetchone()["path"]
        (config.DATA / path).write_bytes(b"changed")
        self.assertFalse(a.post("/api/attachments/%d/verify" % aid).json()["intact"])
        # deletion needs a reason; locked stage blocks upload
        self.assertEqual(a.delete("/api/attachments/%d" % aid).status_code, 400)
        approve_stage(mid, "charter", a)
        r = a.post("/api/methods/%d/attachments?filename=late.txt&stage_key=charter" % mid, content=b"x")
        self.assertEqual(r.status_code, 409)

    def test_URS32_access_control_between_projects(self):
        a, mid, _ = make_project()
        bob = login("bob")                                                   # member of neither the project nor the method
        self.assertEqual(bob.get("/api/methods/%d" % mid).status_code, 403)
        self.assertEqual(bob.get("/api/projects/%d" % project_id_of(a, mid)).status_code, 403)
        self.assertEqual(bob.get("/api/methods/%d/stages/atp" % mid).status_code, 403)
        self.assertEqual(login("admin").get("/api/methods/%d" % mid).status_code, 200)
        self.assertEqual(login("admin").get("/api/projects/%d" % project_id_of(a, mid)).status_code, 200)


class TestDocuments(unittest.TestCase):
    def test_URS40_document_draft_and_final_status(self):
        a, mid, _ = make_project()
        r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        self.assertEqual(r.status_code, 200, r.text)
        did = r.json()["id"]
        content = a.get("/api/docs/%d/download" % did).content
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(content)))
        from docx import Document
        d = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        self.assertIn("草稿", text)
        self.assertIn("分析目标概况", text)
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual(lst[0]["status_label"], "草稿（含未复核内容）")
        self.assertEqual(lst[0]["format"], "docx")
        # approve everything in the plan phase -> final status
        for k, w in (("charter", False), ("atp", False), ("tech", False), ("risk", False), ("plan", True)):
            approve_stage(mid, k, a, with_approver=w)
        r = a.post("/api/methods/%d/docs" % mid, json={"kind": "plan"})
        lst = a.get("/api/methods/%d/docs" % mid).json()
        self.assertEqual(lst[0]["status_label"], "已批准版本")
        d = Document(io.BytesIO(a.get("/api/docs/%d/download" % r.json()["id"]).content))
        text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        self.assertIn("已批准版本", text)

    def test_URS41_all_templates_instantiate(self):
        for tk in ("flow", "hplc", "bioassay", "qpcr", "elisa", "cell_based", "generic"):
            a, mid, _ = make_project(tk, title="t-" + tk)
            p = a.get("/api/methods/%d" % mid).json()
            self.assertEqual(len(p["stages"]), 9, tk)


class TestKbAndBackup(unittest.TestCase):
    def test_URS50_kb_filter_and_lessons(self):
        a, mid, _ = make_project()
        items = a.get("/api/methods/%d/kb?stage=tech" % mid).json()
        self.assertTrue(any("串色" in i["title"] for i in items))
        self.assertFalse(any("qPCR 验证设计" in i["title"] for i in items))   # technique filter
        r = a.post("/api/methods/%d/lessons" % mid, json={"title": "测试经验", "body": "内容", "stage_key": "tech"})
        self.assertEqual(r.status_code, 200)
        allkb = a.get("/api/kb?q=测试经验").json()
        self.assertEqual(len(allkb), 1)

    def test_URS51_backup(self):
        admin = login("admin")
        name = admin.post("/api/admin/backup").json()["name"]
        self.assertTrue((config.BACKUPS / name).exists())
        con = sqlite3.connect(str(config.BACKUPS / name))
        self.assertGreater(con.execute("SELECT COUNT(*) FROM users").fetchone()[0], 0)
        con.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
