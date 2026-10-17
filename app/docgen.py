"""Approval-level document generation: 《分析方法开发方案》 (plan) and 《分析方法开发报告》 (report) as Word + PDF.

The documents are *template driven*: the chapter list comes from ``template["doc_structure"]`` and every
section's position from its ``ch[kind]`` number (see templates_data.py); the generator never hard-codes chapter
order.  Automatic blocks (experiment list / results, Gantt chart, signature page) are inserted at the chapters
that ask for them.  Stage titles, stage status lines and the editor-facing stage ``desc`` are NOT part of the
body of a controlled document.

generate() works in phases so that the slow parts (charts, the .docx, the LibreOffice conversion) never run while
a database write lock is held:
  1. collect   read-only snapshot (stages, signatures, experiments, settings ...)
  2. render    .docx file on disk, then the .pdf via LibreOffice (``soffice --headless --convert-to pdf``);
               the PDF is converted twice when a PDF reader library (pypdf or PyMuPDF) is available: the first
               PDF tells us on which page each heading landed, the contents list is then filled in and the
               document converted again, so the page numbers of the contents list are right in Word and PDF
  3. record    one short write transaction that re-checks the snapshot is still current, then inserts one
               ``generated_docs`` row per format (shared ``gen_id`` = id of the .docx row) and the audit record.
"""
import datetime as dt
import hashlib
import io
import json
import os
import pathlib
import re
import shutil
import subprocess
import threading

import matplotlib
matplotlib.use("Agg")
from matplotlib import font_manager as fm
from matplotlib.figure import Figure
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import audit, config, db, i18n_doc, templates_data, validate, workflow

SANS = "Microsoft YaHei"          # headings, tables, header / footer
SERIF = "SimSun"                  # running text (CJK); Latin text uses Times New Roman
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x55, 0x97)
GRAY = RGBColor(0x59, 0x59, 0x59)
LGRAY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xC0, 0x39, 0x2B)
GREEN = RGBColor(0x0B, 0x7A, 0x4B)
AMBER = RGBColor(0x9A, 0x62, 0x00)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)

INTEGRITY_FAIL_LABEL = "完整性校验失败"
DRAFT_LABEL = "草稿（含未复核内容）"
FINAL_LABEL = "已批准版本"
PENDING_LABEL = "待确认版（含未确认内容，不得作为批准文件）"
DRAFT_HEADER = "草稿—不得作为批准文件使用"
INTEGRITY_HEADER = "完整性校验失败—不得作为批准文件使用"
TODO_TEXT = "【待填写】"

PORTRAIT_W = 17.0                 # usable text width in cm
LANDSCAPE_W = 26.1
WIDE_COLS = 9                     # tables with at least this many columns go into a landscape section
FILL_HEAD = "DCE6F2"
FILL_BAND = "EEF3FB"
FILL_BAD = "FBE9E7"
FILL_OK = "E3F3F2"
FILL_WARN = "FFF4DC"
BORDER = "8C98AB"

_font_ready = False
_font_lock = threading.Lock()
_XML_BAD = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_TAG_RE = re.compile("[（(]?\\s*" + re.escape(workflow.SAMPLE_WORD) + "\\s*[）)]?")


def _x(s):
    """Text that is safe to put into a WordprocessingML run (M7): no XML-1.0-illegal control characters."""
    s = "" if s is None else str(s)
    return _XML_BAD.sub("", s.replace("\r\n", "\n").replace("\r", "\n"))


def _s(v):
    """Any stored value -> display text (tolerates data written before validation existed)."""
    if v is None:
        return ""
    if isinstance(v, str):
        return v
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, int):
        return str(v) if abs(v) < 10 ** 18 else "（数值过大）"
    if isinstance(v, float):
        return "%g" % v
    if isinstance(v, (list, tuple)):
        return "；".join(_s(i) for i in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


# ====================================================================== settings (company name, footer, prefix, soffice)
SETTING_KEYS = ("company_name", "footer_text", "doc_prefix", "soffice_path")
SETTING_DEFAULTS = {"company_name": "", "footer_text": "打印件未经控制", "doc_prefix": "", "soffice_path": ""}
SETTING_LABELS = {"company_name": "公司名称", "footer_text": "文件页脚文字", "doc_prefix": "文件编号前缀", "soffice_path": "soffice 路径"}
_SETTING_MAX = {"company_name": 100, "footer_text": 200, "doc_prefix": 20, "soffice_path": 500}
_PREFIX_RE = re.compile(r"^[A-Za-z0-9._/-]*$")
_SOFFICE_NAMES = ("soffice", "soffice.exe", "soffice.bin", "libreoffice", "libreoffice.exe")


def get_settings(conn):
    """Effective document settings: stored value if the key exists, else the default.  `soffice_path` falls back
    to config.json when the setting was never saved."""
    out = dict(SETTING_DEFAULTS)
    if config.get("soffice_path"):
        out["soffice_path"] = str(config.get("soffice_path"))
    for r in conn.execute("SELECT key, value FROM settings WHERE key IN (%s)" % ",".join("?" * len(SETTING_KEYS)), SETTING_KEYS):
        out[r["key"]] = r["value"] if r["value"] is not None else ""
    return out


def update_settings(conn, user, payload):
    """Validate + store (inside the caller's write transaction) and audit.  Returns the effective settings."""
    if not isinstance(payload, dict):
        raise workflow.WFError("请求内容无效")
    unknown = [k for k in payload if k not in SETTING_KEYS]
    if unknown:
        raise workflow.WFError("未知设置项：%s" % "、".join(str(k) for k in unknown))
    old = get_settings(conn)
    new = {}
    for k in SETTING_KEYS:
        if k not in payload:
            continue
        v = payload[k]
        if v is None:
            v = ""
        if not isinstance(v, str):
            raise workflow.WFError("%s 必须是文本" % SETTING_LABELS[k])
        v = v.strip()
        if k == "soffice_path" and len(v) >= 2 and v[0] == v[-1] and v[0] in "\"'":
            v = v[1:-1].strip()
        if len(v) > _SETTING_MAX[k]:
            raise workflow.WFError("%s 过长（最多 %d 个字符）" % (SETTING_LABELS[k], _SETTING_MAX[k]))
        if _XML_BAD.search(v) or "\n" in v or "\r" in v:
            raise workflow.WFError("%s 含有非法字符" % SETTING_LABELS[k])
        if k == "doc_prefix" and not _PREFIX_RE.match(v):
            raise workflow.WFError("文件编号前缀只能含字母、数字及 . _ / -")
        if k == "soffice_path" and v:
            if os.path.basename(v).lower() not in _SOFFICE_NAMES:
                raise workflow.WFError("soffice 路径应指向 LibreOffice 的 soffice 程序（soffice / soffice.exe）")
            if not os.path.isfile(v):
                raise workflow.WFError("找不到文件：%s" % v)
        new[k] = v
    changes = {}
    for k, v in new.items():
        if old.get(k) != v:
            changes[k] = {"old": old.get(k), "new": v}
        conn.execute("INSERT INTO settings(key, value) VALUES(?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (k, v))
    if changes:
        audit.log(conn, user, "settings.update", "settings", None, None, {"changes": changes})
    return get_settings(conn)


# ====================================================================== LibreOffice / PDF
_LO_LOCK = threading.Lock()          # one conversion at a time: they share one private profile directory

NO_SOFFICE_NOTE = ("未检测到 LibreOffice（soffice），本次只生成了 Word 文件。安装 LibreOffice（https://www.libreoffice.org）后重新生成即可得到 PDF；"
                   "若已安装在非默认位置，请管理员在“管理 › 系统设置”中填写 soffice 路径。")
PDF_OFF_NOTE = "PDF 输出已在配置中关闭（pdf_enabled=false），仅生成 Word 文件。"


def find_soffice(configured=""):
    """soffice executable: configured path (setting / config.json) -> PATH -> default install locations; else None."""
    cands = []
    if configured:
        cands.append(str(configured))
    for name in ("soffice", "soffice.exe", "libreoffice"):
        w = shutil.which(name)
        if w:
            cands.append(w)
    cands += [r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              "/usr/bin/soffice", "/usr/local/bin/soffice", "/opt/libreoffice/program/soffice",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for c in cands:
        try:
            if os.path.isfile(c):
                return c
        except OSError:
            continue
    return None


def capabilities(conn):
    """{"pdf": bool, "pdf_note": str} - whether PDF output can be produced on this server right now."""
    if not config.get("pdf_enabled"):
        return {"pdf": False, "pdf_note": PDF_OFF_NOTE}
    if find_soffice(get_settings(conn)["soffice_path"]):
        return {"pdf": True, "pdf_note": ""}
    return {"pdf": False, "pdf_note": NO_SOFFICE_NOTE}


def _kill_tree(proc):
    try:
        if os.name == "nt":
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(proc.pid)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=20)
        else:
            proc.kill()
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass


def convert_pdf(soffice, docx_path):
    """docx -> pdf next to it, with a private LibreOffice profile and a timeout.  Returns (pdf path or None, note)."""
    docx_path = pathlib.Path(docx_path)
    pdf_path = docx_path.with_suffix(".pdf")
    profile = config.DATA / ".lo_profile"
    timeout = int(config.get("pdf_timeout_sec"))
    try:
        profile.mkdir(parents=True, exist_ok=True)
        try:
            os.remove(str(pdf_path))
        except OSError:
            pass
        cmd = [soffice, "--headless", "--norestore", "--nolockcheck", "--nologo", "--nodefault",
               "-env:UserInstallation=" + profile.resolve().as_uri(),
               "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(docx_path.parent), str(docx_path)]
        kw = {"creationflags": getattr(subprocess, "CREATE_NO_WINDOW", 0)} if os.name == "nt" else {}
        with _LO_LOCK:
            proc = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, stdin=subprocess.DEVNULL, **kw)
            try:
                proc.communicate(timeout=timeout)
            except subprocess.TimeoutExpired:
                _kill_tree(proc)
                try:
                    proc.communicate(timeout=10)
                except Exception:
                    pass
                shutil.rmtree(str(profile), ignore_errors=True)       # a hung instance may have damaged the profile
                return None, "PDF 转换超时（%d 秒），已仅生成 Word 文件。" % timeout
        if not pdf_path.exists() or pdf_path.stat().st_size < 100:
            return None, "PDF 转换失败（LibreOffice 未产生文件，退出码 %s），已仅生成 Word 文件。" % proc.returncode
        with open(str(pdf_path), "rb") as f:
            if f.read(5) != b"%PDF-":
                os.remove(str(pdf_path))
                return None, "PDF 转换失败（输出不是有效的 PDF），已仅生成 Word 文件。"
        return pdf_path, ""
    except Exception as e:                      # never let PDF problems break the Word output
        try:
            os.remove(str(pdf_path))
        except OSError:
            pass
        return None, "PDF 转换失败（%s），已仅生成 Word 文件。" % type(e).__name__


def pdf_heading_pages(pdf_path):
    """[(title, page number 1-based)] of the PDF outline, in document order ([] when no PDF library is available)."""
    try:
        from pypdf import PdfReader
        rd = PdfReader(str(pdf_path))
        out = []

        def walk(items):
            for it in items:
                if isinstance(it, list):
                    walk(it)
                else:
                    out.append((it.title or "", rd.get_destination_page_number(it) + 1))
        walk(rd.outline)
        return out
    except ImportError:
        pass
    except Exception:
        return []
    try:
        import fitz
        d = fitz.open(str(pdf_path))
        try:
            return [(t, p) for _, t, p in d.get_toc(simple=True)]
        finally:
            d.close()
    except Exception:
        return []


def _norm_title(t):
    return re.sub(r"\s+", "", t or "")


def map_pages(entries, outline):
    """entries [(level, text, bookmark)] in order; outline [(title, page)] -> {bookmark: page}."""
    pages, i = {}, 0
    norm = [(_norm_title(t), p) for t, p in outline]
    for _, text, bm in entries:
        want = _norm_title(text)
        j = i
        while j < len(norm) and norm[j][0] != want:
            j += 1
        if j < len(norm):
            pages[bm] = norm[j][1]
            i = j + 1
    return pages


# ====================================================================== docx helpers
_PPR = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
        "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
        "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
        "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
_TCPR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing",
          "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
_SETTINGS_ORDER = ["updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "attachedSchema",
                   "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade", "captions",
                   "readModeInkLockDown", "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"]


def _local_name(el):
    return el.tag.split("}")[-1]


def _ins(parent, child, order):
    """Insert `child` into `parent` at its schema position (Word is strict about child order)."""
    name = _local_name(child)
    for ex in parent.findall(child.tag):
        parent.remove(ex)
    later = set(order[order.index(name) + 1:])
    for i, ch in enumerate(parent):
        if _local_name(ch) in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _el(tag, **attrs):
    """_el("w:shd", w_val="clear") - attribute names use '_' for the namespace colon."""
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k.replace("_", ":", 1)), str(v))
    return e


def _rfonts(rPr, face):
    """Set ascii / hAnsi / eastAsia / cs fonts on an rPr element and drop theme font attributes."""
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in list(rf.attrib):
        if "Theme" in a:
            del rf.attrib[a]
    ascii_face = "Times New Roman" if face == SERIF else face
    rf.set(qn("w:ascii"), ascii_face)
    rf.set(qn("w:hAnsi"), ascii_face)
    rf.set(qn("w:eastAsia"), face)
    rf.set(qn("w:cs"), ascii_face)


def _fmt(run, size=None, bold=None, color=None, italic=None, face=SERIF, highlight=None):
    rPr = run._element.get_or_add_rPr()
    _rfonts(rPr, face)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if highlight is not None:
        run.font.highlight_color = highlight


def _shade_cell(cell, fill):
    _ins(cell._tc.get_or_add_tcPr(), _el("w:shd", w_val="clear", w_color="auto", w_fill=fill), _TCPR)


def _par_border(par, side="bottom", color="1F3864", sz=8, space=1):
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        _ins(pPr, bdr, _PPR)
    bdr.append(_el("w:" + side, w_val="single", w_sz=sz, w_space=space, w_color=color))


def _add_runs(par, text, size=10.5, bold=False, color=None, italic=False, face=SERIF, hl=False):
    """Text (line breaks flattened); example-value tags are highlighted when `hl` (drafts)."""
    text = _x(text).replace("\n", " ")
    if not (hl and _TAG_RE.search(text)):
        if text:
            _fmt(par.add_run(text), size, bold, color, italic, face)
        return
    pos = 0
    for m in _TAG_RE.finditer(text):
        if m.start() > pos:
            _fmt(par.add_run(text[pos:m.start()]), size, bold, color, italic, face)
        _fmt(par.add_run(m.group(0)), size, True, RED, italic, face, WD_COLOR_INDEX.YELLOW)
        pos = m.end()
    if pos < len(text):
        _fmt(par.add_run(text[pos:]), size, bold, color, italic, face)


_LIST_RE = re.compile(r"^\s*((?:\d{1,2}[\.\)、）])|(?:[（(]\d{1,2}[)）])|(?:[-•·●]))\s*")


def _para(container, text="", size=10.5, bold=False, color=None, align=None, italic=False, after=3, before=0, line=1.3,
          face=SERIF, hl=False, keep_next=False, left=None, hanging=None):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if align is not None:
        p.alignment = align
    if left is not None:
        pf.left_indent = Cm(left)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if text:
        _add_runs(p, text, size, bold, color, italic, face, hl)
    return p


class Cx(object):
    """Table cell value with formatting; `paras` = [(text, {size,bold,color,italic})] for multi-style cells."""

    def __init__(self, text="", bold=None, color=None, fill=None, align=None, italic=None, paras=None):
        self.text, self.bold, self.color, self.fill, self.align, self.italic, self.paras = text, bold, color, fill, align, italic, paras


def _cell_fill(cell, value, size=9, bold=False, color=None, align=None, face=SANS, hl=False, fill=None, italic=False):
    spec = value if isinstance(value, Cx) else Cx(value)
    if spec.bold is not None:
        bold = spec.bold
    if spec.color is not None:
        color = spec.color
    if spec.align is not None:
        align = spec.align
    if spec.italic is not None:
        italic = spec.italic
    fill = spec.fill or fill
    if fill:
        _shade_cell(cell, fill)
    lines = []                                   # [(line text, overrides)]
    for text, over in (spec.paras if spec.paras else [(_s(spec.text), {})]):
        for ln in _x(text).split("\n"):
            lines.append((ln, over))
    for i, (ln, over) in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_after = Pt(1)
        pf.space_before = Pt(0)
        pf.line_spacing = 1.12
        if align is not None:
            p.alignment = align
        _add_runs(p, ln, over.get("size", size), over.get("bold", bold), over.get("color", color), over.get("italic", italic), face, hl)


def _new_table(container, rows, cols, width_cm):
    if hasattr(container, "is_linked_to_previous"):             # header / footer
        return container.add_table(rows, cols, Cm(width_cm))
    return container.add_table(rows, cols)


def _grid(tbl, widths_cm, border_color=BORDER, cell_margin=(30, 70)):
    """Fixed layout with explicit tblW / tblGrid / tcW so that Word *and* LibreOffice honour the column widths."""
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    _ins(tblPr, _el("w:tblW", w_w=int(round(sum(widths_cm) * 567)), w_type="dxa"), _TBLPR)
    _ins(tblPr, _el("w:tblLayout", w_type="fixed"), _TBLPR)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el("w:" + side, w_val="single", w_sz=4, w_space=0, w_color=border_color))
    _ins(tblPr, b, _TBLPR)
    m = OxmlElement("w:tblCellMar")
    tb, lr = cell_margin
    for side, v in (("top", tb), ("left", lr), ("bottom", tb), ("right", lr)):
        m.append(_el("w:" + side, w_w=v, w_type="dxa"))
    _ins(tblPr, m, _TBLPR)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(round(w * 567))))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def _row_flags(row, header=False, cant_split=True):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(_el("w:cantSplit", w_val="true"))
    if header:
        trPr.append(_el("w:tblHeader", w_val="true"))


def _scale(widths, total):
    s = float(sum(widths)) or 1.0
    return [total * w / s for w in widths]


def _text_len(row):
    n = 0
    for v in row:
        if isinstance(v, Cx):
            n += len(_s(v.text)) + sum(len(t) for t, _ in (v.paras or []))
        else:
            n += len(_s(v))
    return n


def _add_field(par, instr, size=8, color=GRAY, face=SANS, result="1"):
    """Complex field (PAGE / NUMPAGES ...) with a cached result."""
    for kind in ("begin", "instr", "separate", "text", "end"):
        run = par.add_run()
        _fmt(run, size, color=color, face=face)
        if kind in ("begin", "separate", "end"):
            run._r.append(_el("w:fldChar", w_fldCharType=kind))
        elif kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " %s " % instr
            run._r.append(it)
        else:
            run.text = result


def _local(ts):
    if not ts:
        return ""
    try:
        d = dt.datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=dt.timezone.utc).astimezone()
    except ValueError:
        return str(ts)
    return d.strftime("%Y-%m-%d %H:%M")


def _local_date(ts):
    return _local(ts)[:10]


def _date(s):
    try:
        return dt.date.fromisoformat(s)
    except (TypeError, ValueError):
        return None


# ====================================================================== charts (matplotlib Figure API, no pyplot)
def _mpl_font():
    global _font_ready
    with _font_lock:
        if _font_ready:
            return
        for path in ("C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf", "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
                     "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"):
            try:
                fm.fontManager.addfont(path)
                matplotlib.rcParams["font.family"] = fm.FontProperties(fname=path).get_name()
                break
            except Exception:
                continue
        matplotlib.rcParams["axes.unicode_minus"] = False
        _font_ready = True


def _png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170)
    buf.seek(0)
    return buf


def _gantt_png(stages, T=None):
    T = T or (lambda x: x)
    _mpl_font()
    today = dt.date.today()
    fig = Figure(figsize=(9, 0.42 * len(stages) + 1.0))
    ax = fig.subplots()
    colors = {"approved": "#0E7C7B", "completed": "#0E7C7B", "in_review": "#B7791F", "in_progress": "#2E5597", "returned": "#C0392B", "not_started": "#9aa5b8"}
    for i, s in enumerate(stages):
        ps, pe = _date(s["planned_start"]), _date(s["planned_end"])
        if ps is None or pe is None:
            continue
        ax.barh(i, (pe - ps).days + 1, left=ps.toordinal(), height=0.5, color="#dfe5ee", edgecolor="#9aa5b8")
        a0d = _date(s["actual_start"])
        if a0d:
            a0 = a0d.toordinal()
            a1d = _date(s["actual_end"])
            a1 = a1d.toordinal() if a1d else today.toordinal()
            ax.barh(i, max(a1 - a0, 1), left=a0, height=0.22, color=colors.get(s["status"], "#2E5597"))
    ax.axvline(today.toordinal(), color="#C0392B", lw=1, ls="--")
    ax.set_yticks(range(len(stages)))
    ax.set_yticklabels([T(s["def"]["title"]) for s in stages], fontsize=8)
    ax.invert_yaxis()
    ticks = ax.get_xticks()
    ax.set_xticks(ticks)
    ax.set_xticklabels([dt.date.fromordinal(int(t)).strftime("%m-%d") if t > 1 else "" for t in ticks], fontsize=7)
    ax.set_title(T("阶段计划（浅色=计划，深色=实际，红虚线=今日）"), fontsize=9)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.tight_layout()
    return _png(fig)


def _fmea_png(rows, T=None):
    T = T or (lambda x: x)
    _mpl_font()
    rr = []
    for r in rows:
        trio = workflow._sod(r) if isinstance(r, dict) else None
        if trio is None or not all(1 <= v <= 10 for v in trio):
            continue
        rr.append((_x(r.get("mode") or "?"), trio[0] * trio[1] * trio[2], trio[0]))
    if not rr:
        return None
    rr.sort(key=lambda x: -x[1])
    fig = Figure(figsize=(8.5, 0.38 * len(rr) + 1.0))
    ax = fig.subplots()
    thr, sthr = config.get("rpn_threshold"), config.get("severity_threshold")
    cols = ["#C0392B" if (v >= thr or s >= sthr) else "#2E5597" for _, v, s in rr]
    ax.barh(range(len(rr)), [v for _, v, _ in rr], color=cols)
    ax.set_yticks(range(len(rr)))
    ax.set_yticklabels([m[:22] for m, _, _ in rr], fontsize=8)
    ax.invert_yaxis()
    for i, (_, v, s) in enumerate(rr):
        ax.text(v + 2, i, "%d (S=%d)" % (v, s), va="center", fontsize=7.5)
    ax.axvline(thr, color="#B7791F", ls="--", lw=1)
    ax.set_xlabel(T("RPN = S×O×D（红色：RPN≥%d 或 S≥%d）") % (thr, sthr), fontsize=8)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    fig.tight_layout()
    return _png(fig)


# ====================================================================== document outline (what goes where)
_CH_RE = re.compile(r"^(\d+)\.(\d+)$")


def _ch_of(stage_key, sec, kind, tpl_ch):
    """(chapter, sub) of a section in document `kind`, or None when it is not part of it.  Stage definitions
    stored before template 1.1 have no `ch`: the current template of the method's technique is consulted."""
    ch = sec.get("ch")
    if not (isinstance(ch, dict) and ch):
        ch = tpl_ch.get((stage_key, sec.get("key"))) or {}
    v = ch.get(kind)
    m = _CH_RE.match(v) if isinstance(v, str) else None
    return (int(m.group(1)), int(m.group(2))) if m else None


def _sec_has_content(sec, val):
    t = sec.get("type")
    if t == "text":
        return not workflow._blank(val)
    if t == "fields":
        return isinstance(val, dict) and not workflow._blank(val)
    if t in ("table", "fmea"):
        return any(not workflow._blank(r) for r in workflow._dict_rows(val))
    if t == "matrix":
        m = val if isinstance(val, dict) else {}
        return bool(m.get("options")) and bool(workflow._dict_rows(m.get("criteria")))
    return not workflow._blank(val)


def _sec_wide(sec):
    t = sec.get("type")
    if t == "fmea":
        return True
    if t == "table":
        return len(sec.get("columns") or []) >= WIDE_COLS
    return False


def _stage_sections(s):
    v = s["data"].get("sections") if isinstance(s["data"], dict) else None
    return v if isinstance(v, dict) else {}


def plan_outline(ctx):
    """Flat list of render units + the contents-list entries [(level, text, bookmark)].  Headings take the
    orientation (portrait / landscape) of the content unit that follows them."""
    kind = ctx["kind"]
    struct = ctx["structure"][kind]
    tpl_ch = ctx["tpl_ch"]
    by_ch, orphans = {}, []
    for s in ctx["included"]:
        vals = _stage_sections(s)
        for sec in s["def"].get("sections", []):
            pos = _ch_of(s["key"], sec, kind, tpl_ch)
            legacy_unplaced = not (isinstance(sec.get("ch"), dict) and sec.get("ch")) and (s["key"], sec.get("key")) not in tpl_ch
            if pos is None:
                if legacy_unplaced and _sec_has_content(sec, vals.get(sec["key"])):
                    orphans.append((s, sec))
                continue
            by_ch.setdefault(pos[0], []).append((pos[1], s, sec))
    chapters = [dict(c) for c in struct]
    known = {int(c["no"]) for c in chapters}
    for n in sorted(set(by_ch) - known):
        chapters.append({"no": str(n), "title": "其他内容", "auto": []})
    units, toc, counter = [], [], [0]

    def heading(level, text, page_break=False):
        counter[0] += 1
        b = "_Toc%05d" % counter[0]
        units.append({"u": "h%d" % level, "text": text, "bm": b, "wide": False, "pb": page_break})
        toc.append((level, text, b))

    for ch in chapters:
        n = int(ch["no"])
        items = sorted(by_ch.get(n, []), key=lambda x: x[0])
        children = [{"sub": sub, "u": "section", "s": s, "sec": sec} for sub, s, sec in items]
        taken = {c["sub"] for c in children}

        def free_slot():
            mx = max(taken) if taken else 0
            for k in range(1, mx + 1):
                if k not in taken:
                    return k
            return mx + 1
        want_gantt = False
        for name in ch.get("auto", []):
            if name == "gantt":
                want_gantt = True
                continue
            k = free_slot() if name == "experiments" else (max(taken) if taken else 0) + 1
            taken.add(k)
            title = {"experiments": "实验清单与设计", "experiment_results": "实验结果与结论", "approval": "签署页（电子签名记录）"}[name]
            children.append({"sub": k, "u": "auto", "name": name, "title": title})
        children.sort(key=lambda c: c["sub"])
        bare_approval = not items and any(c["u"] == "auto" and c["name"] == "approval" for c in children)
        heading(1, "%d  %s" % (n, ch["title"]), page_break=bare_approval)
        gantt_done = False
        for c in children:
            if c["u"] == "section":
                heading(2, "%d.%d  %s" % (n, c["sub"], c["sec"]["title"]))
                units.append({"u": "section", "s": c["s"], "sec": c["sec"], "wide": _sec_wide(c["sec"])})
                if want_gantt and not gantt_done:
                    units.append({"u": "gantt", "wide": False})
                    gantt_done = True
            else:
                if not (c["name"] == "approval" and bare_approval):
                    heading(2, "%d.%d  %s" % (n, c["sub"], c["title"]), page_break=(c["name"] == "approval"))
                units.append({"u": "auto", "name": c["name"], "wide": c["name"] in ("experiments", "approval")})
        if want_gantt and not gantt_done:
            units.append({"u": "gantt", "wide": False})
    # appendices
    heading(1, "附录 A  阶段检查清单状态", page_break=True)
    for i, s in enumerate(ctx["included"], 1):
        heading(2, "A.%d  %s" % (i, s["def"]["title"]))
        units.append({"u": "app_check", "s": s, "wide": False})
    heading(1, "附录 B  附件清单")
    units.append({"u": "app_atts", "wide": False})
    heading(1, "附录 C  引用依据")
    refno = next((c["no"] for c in chapters if c["title"].startswith("参考文件")), "4")
    heading(2, "C.1  参考文件与法规依据（第 %s 章）" % refno)
    units.append({"u": "app_refdocs", "wide": False})
    heading(2, "C.2  检查清单引用的条款")
    units.append({"u": "app_checkrefs", "wide": False})
    if orphans:
        heading(1, "附录 D  未纳入章节结构的内容", page_break=True)
        units.append({"u": "app_orphans", "items": orphans, "wide": False})
    nxt = False
    for u in reversed(units):
        if u["u"] in ("h1", "h2"):
            u["wide"] = nxt
        else:
            nxt = u["wide"]
    return units, toc


# ====================================================================== render state + page furniture
class R(object):
    """Rendering state: the python-docx document, current orientation, counters and the draft/final switch."""
    _bm_id = 100

    def __init__(self, ctx):
        self.ctx = ctx
        self.doc = Document()
        self.draft = not ctx["final"]
        self.T = i18n_doc.translate if ctx.get("lang") == "en" else (lambda x: x)
        self.wide = False
        self.fig = 0
        self.W = PORTRAIT_W
        self._styles()
        self.setup_section(self.doc.sections[0], False)

    # ---- styles
    def _restyle(self, name, face, size, bold, color, before=None, after=None, keep_next=None, line=None):
        st = self.doc.styles[name]
        _rfonts(st.element.get_or_add_rPr(), face)
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.font.color.rgb = color
        pf = st.paragraph_format
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if keep_next is not None:
            pf.keep_with_next = keep_next
        if line is not None:
            pf.line_spacing = line
        return st

    def _styles(self):
        doc = self.doc
        n = doc.styles["Normal"]
        rPr = n.element.get_or_add_rPr()
        _rfonts(rPr, SERIF)
        lang = rPr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rPr.append(lang)
        lang.set(qn("w:val"), "en-US")
        lang.set(qn("w:eastAsia"), "zh-CN")
        n.font.size = Pt(10.5)
        n.paragraph_format.space_after = Pt(3)
        h1 = self._restyle("Heading 1", SANS, 15, True, NAVY, 16, 8, True, 1.1)
        bdr = OxmlElement("w:pBdr")
        bdr.append(_el("w:bottom", w_val="single", w_sz=8, w_space=2, w_color="1F3864"))
        _ins(h1.element.get_or_add_pPr(), bdr, _PPR)
        self._restyle("Heading 2", SANS, 12, True, NAVY, 10, 4, True, 1.1)
        self._restyle("Heading 3", SANS, 10.5, True, RGBColor(0x33, 0x33, 0x33), 8, 3, True, 1.1)
        for nm, size, bold, ind in (("TOC 1", 10, True, 0.0), ("TOC 2", 9, False, 0.6)):
            try:
                st = doc.styles[nm]
            except KeyError:
                st = doc.styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = doc.styles["Normal"]
            _rfonts(st.element.get_or_add_rPr(), SANS)
            st.font.size = Pt(size)
            st.font.bold = bold
            st.font.color.rgb = BLACK if bold else GRAY
            st.paragraph_format.space_after = Pt(1.5)
            st.paragraph_format.space_before = Pt(4 if bold else 0)
            st.paragraph_format.left_indent = Cm(ind)
            st.paragraph_format.line_spacing = 1.1
            st.paragraph_format.tab_stops.add_tab_stop(Cm(PORTRAIT_W), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # ---- sections
    def setup_section(self, sec, wide):
        sec.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = (Cm(29.7), Cm(21.0)) if wide else (Cm(21.0), Cm(29.7))
        sec.left_margin = sec.right_margin = Cm(1.8 if wide else 2.0)
        sec.top_margin = Cm(3.0 if (self.draft or self.ctx["bad_sig"]) else 2.6)
        sec.bottom_margin = Cm(2.0)
        sec.header_distance = Cm(0.8)
        sec.footer_distance = Cm(0.8)
        self.W = LANDSCAPE_W if wide else PORTRAIT_W
        self.build_header(sec, self.W)
        self.build_footer(sec, self.W)

    def set_orient(self, wide):
        if wide == self.wide:
            return
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self.wide = wide
        self.setup_section(sec, wide)

    def build_header(self, sec, W):
        ctx = self.ctx
        hdr = sec.header
        hdr.is_linked_to_previous = False
        for el in list(hdr._element):
            hdr._element.remove(el)
        draft_row = self.draft or bool(ctx["bad_sig"])
        tbl = _new_table(hdr, 2 if draft_row else 1, 3, W)
        _grid(tbl, _scale([3.6, 7.6, 5.8], W))
        c0, c1, c2 = tbl.rows[0].cells
        co = ctx["settings"]["company_name"]
        _cell_fill(c0, Cx(co or "（公司名称未设置）", bold=True, color=NAVY if co else LGRAY, align=WD_ALIGN_PARAGRAPH.CENTER), 9.5)
        _cell_fill(c1, Cx(paras=[(ctx["doc_title"], {"size": 11.5, "bold": True, "color": NAVY}),
                                 ("%s　%s" % (ctx["p"]["code"], ctx["p"]["name"]), {"size": 8, "color": GRAY})], align=WD_ALIGN_PARAGRAPH.CENTER), 9)
        _cell_fill(c2, Cx(paras=[("文件编号：%s" % ctx["doc_no"], {"size": 8}),
                                 ("版本：%s　状态：%s" % (ctx["version"], ctx["status_short"]),
                                  {"size": 8, "bold": True, "color": RED if draft_row else GREEN})]), 8)
        if draft_row:
            m = tbl.rows[1].cells[0].merge(tbl.rows[1].cells[2])
            _cell_fill(m, Cx(INTEGRITY_HEADER if ctx["bad_sig"] else DRAFT_HEADER, bold=True, color=RED, fill=FILL_BAD,
                             align=WD_ALIGN_PARAGRAPH.CENTER), 9)
        tail = hdr.add_paragraph()
        tail.paragraph_format.space_after = Pt(0)
        tail.paragraph_format.line_spacing = Pt(2)
        _fmt(tail.add_run(""), 1)

    def build_footer(self, sec, W):
        ctx = self.ctx
        ftr = sec.footer
        ftr.is_linked_to_previous = False
        for el in list(ftr._element):
            ftr._element.remove(el)
        p = ftr.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.tab_stops.add_tab_stop(Cm(W / 2.0), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(W), WD_TAB_ALIGNMENT.RIGHT)
        _par_border(p, "top", "8C98AB", 4, 4)
        _fmt(p.add_run(_x(ctx["settings"]["footer_text"])), 8, color=GRAY, face=SANS)
        _fmt(p.add_run("\t第 "), 8, color=GRAY, face=SANS)
        _add_field(p, "PAGE")
        _fmt(p.add_run(" 页 / 共 "), 8, color=GRAY, face=SANS)
        _add_field(p, "NUMPAGES")
        _fmt(p.add_run(" 页\t生成时间：%s" % ctx["generated_at"]), 8, color=GRAY, face=SANS)

    # ---- blocks
    def heading(self, level, text, bm=None, page_break=False):
        p = self.doc.add_paragraph(style="Heading %d" % level)
        if page_break:
            p.paragraph_format.page_break_before = True
        _fmt(p.add_run(_x(text)), face=SANS)
        if bm:
            R._bm_id += 1
            s = _el("w:bookmarkStart", w_id=R._bm_id, w_name=bm)
            e = _el("w:bookmarkEnd", w_id=R._bm_id)
            if p._p.pPr is not None:
                p._p.pPr.addnext(s)
            else:
                p._p.insert(0, s)
            p._p.append(e)
        return p

    def spacer(self, pt=4):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = Pt(pt)
        _fmt(p.add_run(""), 1)
        return p

    def body(self, text, **kw):
        return _para(self.doc, text, hl=self.draft, **kw)

    def note(self, text, color=GRAY, italic=True, size=9):
        return _para(self.doc, text, size, color=color, italic=italic, after=3, face=SANS)

    def todo(self):
        p = _para(self.doc, "", after=3)
        if self.draft:
            _fmt(p.add_run(TODO_TEXT), 10.5, True, RED, face=SANS, highlight=WD_COLOR_INDEX.YELLOW)
        else:
            _fmt(p.add_run(TODO_TEXT), 10.5, True, RED, face=SANS)       # approved with a gap: still clearly visible
        return p

    def table(self, headers, rows, widths, size=9, aligns=None, row_fills=None, header_fill=FILL_HEAD, hl=None, first_col_bold=False):
        """Fixed-layout table that fits the current text width; the header row repeats on every page."""
        hl = self.draft if hl is None else hl
        widths = _scale(widths, self.W)
        tbl = _new_table(self.doc, 1 + len(rows), len(headers), self.W)
        tbl.style = "Table Grid"
        _grid(tbl, widths)
        aligns = aligns or [None] * len(headers)
        for i, h in enumerate(headers):
            _cell_fill(tbl.rows[0].cells[i], Cx(h, align=WD_ALIGN_PARAGRAPH.CENTER if aligns[i] is not None else None), size, True, NAVY, fill=header_fill)
        _row_flags(tbl.rows[0], header=True)
        for ri, r in enumerate(rows):
            cells = tbl.rows[ri + 1].cells
            for i in range(len(headers)):
                v = r[i] if i < len(r) else ""
                _cell_fill(cells[i], v, size, bool(first_col_bold and i == 0), None, aligns[i], hl=hl,
                           fill=(row_fills[ri] if row_fills and row_fills[ri] else None))
            _row_flags(tbl.rows[ri + 1], cant_split=_text_len(r) < 500)
        self.spacer()
        return tbl

    def kv_table(self, pairs, widths=(4.2, 12.8), size=9.5, hl=None, label_fill=FILL_BAND):
        hl = self.draft if hl is None else hl
        widths = _scale(list(widths), self.W)
        tbl = _new_table(self.doc, len(pairs), 2, self.W)
        tbl.style = "Table Grid"
        _grid(tbl, widths)
        for i, (k, v) in enumerate(pairs):
            _cell_fill(tbl.rows[i].cells[0], Cx(k), size, True, NAVY, fill=label_fill)
            _cell_fill(tbl.rows[i].cells[1], v, size, False, None, hl=hl)
            _row_flags(tbl.rows[i], cant_split=_text_len([v]) < 500)
        self.spacer()
        return tbl

    def banner(self, text, fill, color, size=10, bold=True, sub=None):
        tbl = _new_table(self.doc, 1, 1, self.W)
        tbl.style = "Table Grid"
        _grid(tbl, [self.W], border_color="B0B8C8", cell_margin=(70, 120))
        paras = [(text, {"size": size, "bold": bold, "color": color})]
        if sub:
            paras += [(t, {"size": size - 1, "bold": False, "color": GRAY}) for t in sub]
        _cell_fill(tbl.rows[0].cells[0], Cx(paras=paras, fill=fill), size)
        self.spacer()
        return tbl

    def picture(self, buf, width_cm, caption):
        self.fig += 1
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(buf, width=Cm(min(width_cm, self.W)))
        _para(self.doc, "图 %d  %s" % (self.fig, caption), 9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, face=SANS)


# ====================================================================== section renderers
def _align_for(col):
    t = (col or {}).get("type")
    return WD_ALIGN_PARAGRAPH.RIGHT if t == "number" else (WD_ALIGN_PARAGRAPH.CENTER if t == "date" else None)


def _render_text(r, sec, val):
    if not isinstance(val, str):
        val = _s(val)
    for ln in _x(val).split("\n"):
        if not ln.strip():
            continue
        if _LIST_RE.match(ln):
            r.body(ln.strip(), left=0.6, hanging=0.6, after=2)
        else:
            r.body(ln, after=3)


def _verdict_style(val):
    txt = " ".join(_s(v) for v in val.values()) if isinstance(val, dict) else ""
    if any(w in txt for w in ("未达成", "需继续开发", "终止")):
        return FILL_BAD, RED
    if any(w in txt for w in ("部分达成", "有条件")):
        return FILL_WARN, AMBER
    return FILL_OK, GREEN


def _render_fields(r, sec, val):
    flds = sec.get("fields") or []
    val = val if isinstance(val, dict) else {}
    if sec["key"] == "verdict":                                   # conclusion summary banner (D7)
        fill, color = _verdict_style(val)
        head = "　".join("%s：%s" % (f["label"], _s(val.get(f["key"]))) for f in flds
                        if f.get("type") != "textarea" and not workflow._blank(val.get(f["key"])))
        sub = ["%s：%s" % (f["label"], _s(val.get(f["key"]))) for f in flds
               if f.get("type") == "textarea" and not workflow._blank(val.get(f["key"]))]
        r.banner(head or "（结论性意见未填写）", fill, color, 11, sub=sub)
        return
    r.kv_table([(f["label"], Cx(_s(val.get(f["key"])))) for f in flds])


def _render_table(r, sec, val):
    cols = sec.get("columns") or []
    rows = [x for x in workflow._dict_rows(val) if not workflow._blank(x)]
    n = len(cols)
    size = 9 if n <= 6 else (8.5 if n == 7 else 8)
    if r.wide:
        size = 8.5
    out = [[_s(x.get(c["key"])) for c in cols] for x in rows]
    r.table([c["label"] for c in cols], out, [c.get("width") or 100 for c in cols], size, [_align_for(c) for c in cols])


def _render_matrix(r, sec, val):
    m = val if isinstance(val, dict) else {}
    opts = m.get("options") if isinstance(m.get("options"), list) else []
    crit = workflow._dict_rows(m.get("criteria"))
    rows = []
    totals = {o: 0.0 for o in opts}
    for c in crit:
        w = validate.fnum(c.get("weight"))
        scores = c.get("scores") if isinstance(c.get("scores"), dict) else {}
        row = [_s(c.get("name")), _s(c.get("weight"))]
        for o in opts:
            sc = scores.get(o)
            row.append("" if sc in (None, "") else _s(sc))
            totals[o] += w * validate.fnum(sc)
        rows.append(row)
    rows.append([Cx("加权总分", bold=True), ""] + [Cx("%g" % totals[o], bold=True) for o in opts])
    fills = [None] * (len(rows) - 1) + [FILL_BAND]
    ra = WD_ALIGN_PARAGRAPH.CENTER
    r.table(["准则", "权重"] + [_s(o) for o in opts], rows, [5, 1.2] + [2.2] * len(opts), 9, [None, ra] + [ra] * len(opts), row_fills=fills)
    r.body("选定方案：%s" % (_s(m.get("selected")) or "（未选定）"), bold=True)


def _render_fmea(r, sec, val):
    rows = [x for x in workflow._dict_rows(val) if not workflow._blank(x)]
    thr, sthr = config.get("rpn_threshold"), config.get("severity_threshold")
    out, fills = [], []
    for i, x in enumerate(rows, 1):
        trio = workflow._sod(x)
        if trio is not None and all(1 <= v <= 10 for v in trio):
            rpn = trio[0] * trio[1] * trio[2]
            high = rpn >= thr or trio[0] >= sthr
        else:
            rpn, high = "", False
        out.append([str(i), _s(x.get("mode")), _s(x.get("cause")), _s(x.get("effect")), _s(x.get("S")), _s(x.get("O")), _s(x.get("D")),
                    Cx(_s(rpn), bold=high, color=RED if high else None), _s(x.get("control")), _s(x.get("owner"))])
        fills.append(FILL_BAD if high else None)
    c = WD_ALIGN_PARAGRAPH.CENTER
    r.table(["#", "失效模式", "原因", "影响", "S", "O", "D", "RPN", "控制措施", "责任人"], out,
            [0.7, 3.4, 3.2, 3.2, 0.8, 0.8, 0.8, 1.1, 6.0, 2.0], 8.5, [c, None, None, None, c, c, c, c, None, None], row_fills=fills)
    r.note("RPN = S×O×D；红底：RPN ≥ %d 或 S ≥ %d，必须有控制措施。" % (thr, sthr), size=8.5)
    img = _fmea_png(rows, r.T)
    if img:
        r.picture(img, 15.5, "FMEA 风险排序（RPN）")


def render_section(r, sec, val, stage=None):
    t = sec.get("type")
    if not _sec_has_content(sec, val):
        if sec.get("required"):
            r.todo()
        else:
            r.note("（未填写）", LGRAY)
        return
    if t == "text":
        _render_text(r, sec, val)
    elif t == "fields":
        _render_fields(r, sec, val)
    elif t == "table":
        _render_table(r, sec, val)
    elif t == "matrix":
        _render_matrix(r, sec, val)
    elif t == "fmea":
        _render_fmea(r, sec, val)
    else:
        r.body(_s(val))


# ====================================================================== automatic blocks
def _exp_period(e):
    a, b = e.get("planned_start") or "", e.get("planned_end") or ""
    return "%s\n~ %s" % (a, b) if (a or b) else ""


def _experiments_table(r, exps, stage_titles):
    if not exps:
        r.note("（暂无实验）", LGRAY)
        return
    rows = [[_s(e["code"]), _s(e["title"]), stage_titles.get(e.get("stage_key"), ""), _s(e.get("objective")), _s(e.get("design")),
             _s(e.get("acceptance")), _exp_period(e), workflow.EXP_STATUS.get(e["status"], _s(e["status"]))] for e in exps]
    r.table(["编号", "实验", "所属阶段", "目的", "设计与分析", "接受标准", "计划日期", "状态"], rows, [1.4, 3.2, 2.4, 4.2, 6.0, 5.0, 2.4, 1.5], 8.5,
            first_col_bold=True)


def _auto_experiments(r):
    ctx = r.ctx
    titles = {s["key"]: s["def"]["title"] for s in ctx["stages"]}
    early = [e for e in ctx["exps"] if not e.get("added_after_plan")]
    late = [e for e in ctx["exps"] if e.get("added_after_plan")]
    _experiments_table(r, early, titles)
    if late:
        r.heading(3, "方案外追加的实验（方案批准后新增，不在方案签名范围内）")
        _experiments_table(r, late, titles)


def _result_rows(e):
    return [x for x in workflow._dict_rows(e.get("results")) if not workflow._blank(x)]


def _auto_experiment_results(r):
    ctx = r.ctx
    exps = ctx["exps"]
    if not exps:
        r.note("（暂无实验）", LGRAY)
        return
    titles = {s["key"]: s["def"]["title"] for s in ctx["stages"]}
    order = [s["key"] for s in ctx["stages"]]
    rows = []
    for e in exps:
        concl = _s(e.get("conclusion"))
        rows.append([_s(e["code"]), _s(e["title"]), titles.get(e.get("stage_key"), ""), workflow.EXP_STATUS.get(e["status"], _s(e["status"])),
                     concl if len(concl) <= 120 else concl[:118] + "…"])
    r.table(["编号", "实验", "所属阶段", "状态", "结论摘要"], rows, [1.4, 4.2, 3.0, 1.6, 7.0], 8.5, first_col_bold=True)
    merged = {}
    for e in exps:
        merged.setdefault(e.get("stage_key") or "", []).append(e)
    for k in sorted(merged, key=lambda k: order.index(k) if k in order else 99):
        es = merged[k]
        shown = [e for e in es if e["status"] in ("done", "running", "cancelled") or _result_rows(e) or e.get("conclusion") or e.get("deviations")]
        not_shown = [e for e in es if e not in shown]
        r.heading(3, "%s（%d 个实验）" % (titles.get(k, "未归属阶段"), len(es)))
        for e in shown:
            _para(r.doc, "%s  %s" % (_s(e["code"]), _s(e["title"])), 10.5, True, NAVY, after=2, before=4, face=SANS, keep_next=True)
            meta = "状态：%s　实际：%s ~ %s" % (workflow.EXP_STATUS.get(e["status"], _s(e["status"])), e.get("actual_start") or "—", e.get("actual_end") or "—")
            _para(r.doc, meta, 8.5, color=GRAY, after=2, face=SANS, keep_next=True)
            res = _result_rows(e)
            if res:
                r.table(["指标/观察", "结果", "判据", "判定"], [[_s(x.get("metric")), _s(x.get("value")), _s(x.get("criterion")), _s(x.get("verdict"))] for x in res],
                        [4.5, 3.5, 5.5, 1.8], 8.5)
            r.body("结论：" + (_s(e.get("conclusion")) or "（未填写）"), after=2)
            if e.get("deviations"):
                r.body("偏离/说明：" + _s(e["deviations"]), color=RED, after=3)
        if not_shown:
            r.note("尚未开始/无结果记录：" + "、".join(_s(e["code"]) for e in not_shown), LGRAY)


def _auto_gantt(r):
    r.picture(_gantt_png(r.ctx["stages"], r.T), 15.5, "阶段计划（来自各阶段计划/实际日期）")


def _meaning(rv):
    if rv.get("meaning"):
        return rv["meaning"]
    return {"reviewer": "复核", "approver": "批准"}.get(rv["role"], "") + ("：通过" if rv["decision"] == "approve" else "：退回")


def _auto_approval(r):
    ctx = r.ctx
    r.body("下列记录由 MethodFlow 依据各评审门的电子签名生成。签名绑定到提交时的阶段数据哈希（含该阶段实验与附件）；签署后数据若被改动，哈希将不一致，"
           "本页相应行会显示“完整性校验失败”。", size=9.5, color=GRAY, after=4)
    c = WD_ALIGN_PARAGRAPH.CENTER
    for s in ctx["included"]:
        gate = s["def"].get("gate")
        if not gate:
            continue
        lab = workflow.STATUS_LABELS[s["status"]] + ("，已失效(需重新复核)" if s["stale"] else "")
        ok = s["status"] in workflow.LOCKED and not s["stale"] and ctx["sig"][s["id"]][0]
        _para(r.doc, "%s　（阶段：%s　修订 R%d　状态：%s）" % (gate["name"], s["def"]["title"], s["revision"], lab), 10, True, GREEN if ok else RED,
              after=2, before=6, face=SANS, keep_next=True)
        rows, fills = [], []
        rv = ctx["reviews"].get(s["id"], [])
        for x in rv:
            intact = ctx["intact"].get(x["id"], False)
            rows.append([workflow.ROLE_LABELS.get(x["role"], x["role"]), _s(x["display_name"]), _meaning(x),
                         Cx("通过" if x["decision"] == "approve" else "退回", bold=True, color=GREEN if x["decision"] == "approve" else RED),
                         _local(x["signed_at"]), "R%d" % x["revision"], _s(x["data_hash"])[:12],
                         Cx("有效" if intact else INTEGRITY_FAIL_LABEL, bold=True, color=GREEN if intact else RED)])
            fills.append(None if intact else FILL_BAD)
        have = {x["role"] for x in rv if x["decision"] == "approve" and x["revision"] == s["revision"]}
        for role in gate["roles"]:
            if role not in have:
                rows.append([workflow.ROLE_LABELS.get(role, role), Cx("—"), Cx("未签署" if not rv else "本修订尚无该角色的通过签名", color=RED), "", "",
                             "R%d" % s["revision"], "", ""])
                fills.append(FILL_WARN)
        r.table(["签署角色", "姓名", "签署含义", "决定", "签署日期时间", "修订号", "数据哈希(前12位)", "签名完整性"], rows,
                [2.6, 2.4, 9.0, 1.3, 3.0, 1.3, 3.4, 2.4], 8.5, [None, None, None, c, c, c, c, c], row_fills=fills)
        if ctx["sig"][s["id"]][1] and s["status"] in workflow.LOCKED:
            r.note("；".join(ctx["sig"][s["id"]][1]), RED, False, 8.5)
    nogate = [s for s in ctx["included"] if not s["def"].get("gate")]
    if nogate:
        _para(r.doc, "无评审门的阶段（提交即完成）", 10, True, NAVY, after=2, before=6, face=SANS, keep_next=True)
        rows = []
        for s in nogate:
            sub = ctx["submissions"].get(s["id"])
            rows.append([s["def"]["title"], workflow.STATUS_LABELS[s["status"]] + ("，已失效" if s["stale"] else ""), "R%d" % s["revision"],
                         s.get("actual_end") or "", _s(sub["display_name"]) if sub else ""])
        r.table(["阶段", "状态", "修订", "完成日期", "提交人"], rows, [6, 4, 2, 3, 4], 8.5)
    r.note("本程序的电子签名是否满足企业 21 CFR Part 11 / EU GMP Annex 11 的要求，须由 QA 按计算机化系统验证流程评估。", size=8.5)


_AUTO = {"experiments": _auto_experiments, "experiment_results": _auto_experiment_results, "approval": _auto_approval}


# ====================================================================== appendices
def _app_check(r, s):
    rows = []
    cl = s["data"].get("checklist") if isinstance(s["data"].get("checklist"), dict) else {}
    for c in s["def"].get("checklist", []):
        it = cl.get(c["key"], {})
        it = it if isinstance(it, dict) else {}
        st = it.get("status", "todo")
        lab = {"todo": "未完成", "done": "已完成", "na": "不适用"}.get(st, "")
        paras = [(c["text"], {})]
        if c.get("guidance"):
            paras.append(("指引：" + c["guidance"], {"size": 7.5, "color": LGRAY, "italic": True}))
        rows.append([Cx(paras=paras), _s(c.get("ref")) or "—", "必需" if c.get("required") else "建议",
                     Cx(lab, color=GREEN if st == "done" else (RED if st == "todo" and c.get("required") else None)), _s(it.get("note"))])
    r.table(["检查项（含指引）", "依据", "要求", "状态", "说明/证据"], rows, [6.2, 3.4, 1.1, 1.5, 4.2], 8.5)


def _app_atts(r):
    att = r.ctx["atts"]
    if not att:
        r.note("（无附件）", LGRAY)
        return
    r.table(["#", "文件名", "归属", "大小(KB)", "SHA-256(前12位)", "上传人/时间"],
            [[str(a["id"]), _s(a["filename"]), _s(a["stage_key"]) + ((" / 实验%d" % a["experiment_id"]) if a["experiment_id"] else ""),
              "%.1f" % (a["size"] / 1024.0), _s(a["sha256"])[:12], "%s %s" % (a["display_name"] or "", _local(a["uploaded_at"]))] for a in att],
            [0.8, 5.0, 2.6, 1.6, 3.0, 3.6], 8.5, [WD_ALIGN_PARAGRAPH.CENTER, None, None, WD_ALIGN_PARAGRAPH.RIGHT, None, None])


def _app_refdocs(r):
    st = next((s for s in r.ctx["included"] if s["key"] == "charter"), None)
    rows = []
    if st:
        for x in workflow._dict_rows(_stage_sections(st).get("refdocs")):
            if not workflow._blank(x):
                rows.append([_s(x.get("category")), _s(x.get("title")), _s(x.get("number")), _s(x.get("relevance"))])
    if not rows:
        r.note("（未填写）", LGRAY)
        return
    r.table(["类别", "文件名称", "文件编号/版本/日期", "适用章节/用途"], rows, [2.0, 6.2, 4.0, 4.8], 8.5)


def _app_checkrefs(r):
    refs = []
    for s in r.ctx["included"]:
        for c in s["def"].get("checklist", []):
            for x in re.split(r"[；;]", c.get("ref") or ""):
                x = x.strip()
                if x and x not in refs:
                    refs.append(x)
    if not refs:
        r.note("（无）", LGRAY)
    for i, x in enumerate(refs, 1):
        _para(r.doc, "%d. %s" % (i, x), 9, after=1, face=SANS, left=0.6, hanging=0.6)
    r.note("条款号按所引用文件的版本核对；标注“待核对”的条款在引用前须对照原文确认。", size=8.5)


def _app_orphans(r, items):
    r.note("下列内容来自较早模板版本创建的阶段定义，其章节在当前文档结构中没有对应编号，按原样列出。")
    for s, sec in items:
        r.heading(3, "%s › %s" % (s["def"]["title"], sec["title"]))
        render_section(r, sec, _stage_sections(s).get(sec["key"]), s)


# ====================================================================== cover, control page, contents
def _status_banner(r):
    ctx = r.ctx
    if ctx["bad_sig"]:
        names = "；".join("%s（%s）" % (s["def"]["title"], "、".join(ctx["sig"][s["id"]][1][:2])) for s in ctx["bad_sig"])
        r.banner("状态：完整性校验失败", FILL_BAD, RED, 14)
        r.banner("完整性校验失败 —— 下列已通过阶段的电子签名与当前数据不一致（数据可能在签署后被改动）。本文件不得作为批准文件使用，请 QA 调查：" + names,
                 FILL_BAD, RED)
    if ctx["final"]:
        gap_txt = ""
        if ctx.get("gaps"):
            gap_txt = "　注意：仍有 %d 个必填章节未填写（已在正文标为红色【待填写】）：%s" % (len(ctx["gaps"]), "；".join(ctx["gaps"][:8]) + ("；……" if len(ctx["gaps"]) > 8 else ""))
        r.banner("状态：已批准版本 —— 本文包含的全部阶段均已通过评审门并锁定，且电子签名完整性校验通过。" + gap_txt, FILL_OK if not gap_txt else FILL_BAD, GREEN if not gap_txt else RED)
    elif ctx["unapproved"]:
        names = "、".join("%s（%s）" % (s["def"]["title"], workflow.STATUS_LABELS[s["status"]] + ("，已失效" if s["stale"] else "")) for s in ctx["unapproved"])
        r.banner("状态：草稿 —— 下列阶段尚未通过复核/完成，内容可能变化，不得作为批准文件使用：" + names, FILL_BAD, RED)
    elif ctx.get("pending"):
        r.banner("状态：待确认版 —— 各阶段均已通过评审，但仍有未确认的模板预置内容或示例值，不得作为批准文件使用：" + "；".join(ctx["pending"][:12])
                 + ("；……另有 %d 项" % (len(ctx["pending"]) - 12) if len(ctx["pending"]) > 12 else ""), FILL_BAD, RED)


def _cover(r):
    ctx, doc = r.ctx, r.doc
    p, project = ctx["p"], ctx["project"]
    co = ctx["settings"]["company_name"]
    _para(doc, "", after=14)
    _para(doc, co or "（公司名称未设置，可在“管理 › 系统设置”中配置）", 15, True, NAVY if co else LGRAY, WD_ALIGN_PARAGRAPH.CENTER, after=26, face=SANS)
    _para(doc, ctx["doc_title"], 28, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=10, face=SANS)
    _para(doc, p["name"], 18, True, BLUE, WD_ALIGN_PARAGRAPH.CENTER, after=6, face=SANS)
    _para(doc, "%s　%s" % (ctx["doc_no"], ctx["version"]), 12, False, GRAY, WD_ALIGN_PARAGRAPH.CENTER, after=16, face=SANS)
    eff = ctx["effective"] if ctx["final"] and ctx["effective"] else "—（草稿，尚未批准）"
    r.kv_table([
        ("文件编号", ctx["doc_no"]), ("文件名称", "%s — %s" % (ctx["doc_title"], p["name"])), ("版本", ctx["version"]),
        ("方法编号", p["code"]), ("所属项目", "%s　%s" % (project["code"], project["name"])),
        ("被测属性 / 样品基质", "%s / %s" % (p["analyte"] or "—", p["matrix"] or "—")),
        ("技术平台", "%s（模板 %s）" % (p["meta"].get("template_label", p["technique"]), p["template_version"] or "")),
        ("文件状态", ctx["status_label"]), ("生效/批准日期", eff)], widths=(4.0, 13.0), size=10)
    _para(doc, "", after=2)
    rows = []
    for role, key in (("编制", "author"), ("复核", "reviewer"), ("批准", "approver")):
        who = ctx["cover_people"].get(key)
        rows.append([Cx(role, bold=True), who["name"] if who else Cx("—", color=LGRAY), who["date"] if who else Cx("—", color=LGRAY),
                     who["note"] if who else Cx("待提交" if key == "author" else "待签署", color=LGRAY)])
    r.table(["职责", "姓名", "日期", "依据"], rows, [2.4, 4.4, 3.2, 7.0], 10, first_col_bold=True)
    _status_banner(r)


def _revision_rows(ctx):
    ev = []
    for d in ctx["prev_docs"]:
        ev.append((d["created_at"] or "", "文档版本", d["version_label"] or "", "生成《%s》%s（%s）" % (ctx["doc_title"], d["version_label"] or "", d["status_label"] or ""),
                   d["display_name"] or ""))
    for s in ctx["included"]:
        for x in ctx["stage_revs"].get(s["id"], []):
            ev.append((x["created_at"] or "", "阶段提交", "%s R%d" % (s["def"]["title"], x["revision"]), x["reason"] or "提交", x["display_name"] or ""))
        for x in ctx["reviews"].get(s["id"], []):
            if x["decision"] != "approve":
                ev.append((x["signed_at"] or "", "评审退回", "%s R%d" % (s["def"]["title"], x["revision"]), x["comment"] or "退回", x["display_name"] or ""))
    stage_title = {s["key"]: s["def"]["title"] for s in ctx["included"]}
    for a in ctx["audit_changes"]:
        d = a["details"] if isinstance(a["details"], dict) else {}
        if d.get("stage") not in stage_title:
            continue
        what = {"stage.reopen": "重新打开", "stage.withdraw": "撤回提交"}[a["action"]]
        ev.append((a["ts"] or "", what, "%s R%s" % (stage_title[d.get("stage")], d.get("revision", "")), a["reason"] or "", a["display_name"] or a["username"] or ""))
    ev.sort(key=lambda e: e[0])
    rows = [[str(i), t, _local(ts), ver, desc, who] for i, (ts, t, ver, desc, who) in enumerate(ev, 1)]
    rows.append([Cx(str(len(rows) + 1), bold=True), Cx("本版", bold=True), ctx["generated_at"], Cx(ctx["version"], bold=True),
                 Cx("生成《%s》%s（%s）" % (ctx["doc_title"], ctx["version"], ctx["status_label"]), bold=True), ctx["user"]["display_name"]])
    return rows


def _control_page(r):
    ctx = r.ctx
    p, project = ctx["p"], ctx["project"]
    r.heading(1, "文件控制", page_break=True)
    r.heading(2, "文件信息")
    r.kv_table([
        ("文件编号 / 版本", "%s / %s" % (ctx["doc_no"], ctx["version"])),
        ("方法编号 / 名称", "%s / %s" % (p["code"], p["name"])),
        ("所属项目", "%s / %s" % (project["code"], project["name"])),
        ("产品 / 适应症", "%s / %s" % (project["product"] or "—", project["indication"] or "—")),
        ("被测属性 / 样品基质 / 预期用途", "%s / %s / %s" % (p["analyte"] or "—", p["matrix"] or "—", p["intended_use"] or "—")),
        ("技术模板", "%s（模板版本 %s）" % (p["meta"].get("template_label", p["technique"]), p["template_version"] or "")),
        ("方法版本 / 生命周期", "%s / %s" % (p["method_version"], workflow.LIFECYCLE_LABELS.get(p["lifecycle"], p["lifecycle"]))),
        ("方法计划周期", "%s ~ %s" % (p["planned_start"] or "—", p["planned_end"] or "—")),
        ("文件状态", ctx["status_label"]),
        ("生成时间 / 生成人", "%s / %s" % (ctx["generated_at"], ctx["user"]["display_name"])),
        ("内容数据哈希（SHA-256）", ctx["doc_hash"][:32] + "…"),
    ], widths=(5.0, 12.0), size=9.5)
    r.heading(2, "修订历史")
    rows = _revision_rows(ctx)
    if len(rows) > 150:
        rows = rows[:20] + [["…", "", "", "", Cx("（中间 %d 条从略，完整记录见审计日志）" % (len(rows) - 100), italic=True), ""]] + rows[-80:]
    r.table(["序号", "类型", "日期时间", "版本/修订", "说明/原因", "人员"], rows, [1.0, 2.0, 3.0, 3.8, 5.6, 2.2], 8.5,
            [WD_ALIGN_PARAGRAPH.CENTER, None, None, None, None, None])
    r.note("修订历史汇总：已生成的文档版本、各阶段每次提交的修订号与说明、评审退回意见、阶段重新打开/撤回的原因（来自审计记录）。", size=8.5)


def _toc_par_runs(par, entry, page, first, last):
    """One contents entry: hyperlink to the heading bookmark + dot leader + PAGEREF field.  The first / last
    paragraph also carry the begin / end of the surrounding TOC field, so Word can update the whole list."""
    lvl, text, bm = entry

    def fc(kind):
        run = OxmlElement("w:r")
        run.append(_el("w:fldChar", w_fldCharType=kind))
        return run

    def instr(txt):
        run = OxmlElement("w:r")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = txt
        run.append(it)
        return run
    if first:
        par._p.append(fc("begin"))
        par._p.append(instr(' TOC \\o "1-2" \\h \\z \\u '))
        par._p.append(fc("separate"))
    h = _el("w:hyperlink", w_anchor=bm, w_history="1")
    r1 = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = _x(text)
    r1.append(t)
    h.append(r1)
    r2 = OxmlElement("w:r")
    r2.append(OxmlElement("w:tab"))
    h.append(r2)
    h.append(fc("begin"))
    h.append(instr(" PAGEREF %s \\h " % bm))
    h.append(fc("separate"))
    r4 = OxmlElement("w:r")
    t4 = OxmlElement("w:t")
    t4.text = str(page) if page else ""
    r4.append(t4)
    h.append(r4)
    h.append(fc("end"))
    par._p.append(h)
    if last:
        par._p.append(fc("end"))


def _contents(r, toc, pages):
    doc = r.doc
    t = _para(doc, "目　录", 16, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=10, face=SANS)
    t.paragraph_format.page_break_before = True
    for i, e in enumerate(toc):
        par = doc.add_paragraph(style="TOC 1" if e[0] == 1 else "TOC 2")
        _toc_par_runs(par, e, pages.get(e[2]) if pages else None, i == 0, i == len(toc) - 1)


# ====================================================================== collect (phase 1)
def _change_signature(conn, mid):
    """Cheap fingerprint of everything a document is built from; compared again right before the INSERT."""
    m = conn.execute("SELECT updated_at, project_id FROM methods WHERE id=?", (mid,)).fetchone()
    return db.canon({
        "stages": [[r["key"], r["revision"], r["version"], r["status"], r["stale"]] for r in
                   conn.execute("SELECT key, revision, version, status, stale FROM stages WHERE method_id=? ORDER BY ord", (mid,))],
        "exps": [[r["id"], r["version"], r["status"]] for r in conn.execute("SELECT id, version, status FROM experiments WHERE method_id=? ORDER BY id", (mid,))],
        "atts": [r["id"] for r in conn.execute("SELECT id FROM attachments WHERE method_id=? AND deleted=0 ORDER BY id", (mid,))],
        "reviews": conn.execute("SELECT COUNT(*) FROM reviews WHERE method_id=?", (mid,)).fetchone()[0],
        "method": [m["updated_at"], conn.execute("SELECT updated_at FROM projects WHERE id=?", (m["project_id"],)).fetchone()[0]],
    })


def _check_rate(conn, mid):
    since = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(seconds=60)).strftime("%Y-%m-%dT%H:%M:%SZ")
    n = conn.execute("SELECT COUNT(*) FROM generated_docs WHERE method_id=? AND created_at>=? AND format='docx'", (mid, since)).fetchone()[0]
    if n >= config.get("docs_per_minute"):
        raise workflow.WFError("生成文档过于频繁（每个方法每分钟最多 %d 次），请稍后再试" % config.get("docs_per_minute"), 429)


def collect(conn, user, mid, kind="plan", lang="zh"):
    T = i18n_doc.translate if lang == "en" else (lambda x: x)
    if not isinstance(kind, str) or kind not in ("plan", "report"):
        raise workflow.WFError("未知文档类型")
    p = workflow.get_method(conn, mid)
    project = workflow.get_project(conn, p["project_id"])
    roles = workflow.method_roles(conn, mid, user["id"])
    if not ("admin" in user["roles"] or roles & {"owner", "author", "reviewer", "approver"}):
        raise workflow.WFError("无权生成文档（仅方法的负责人/开发人员/复核人/批准人或管理员可以）", 403)
    _check_rate(conn, mid)
    settings = get_settings(conn)
    stages = workflow.list_stages(conn, mid)
    exps = workflow.list_experiments(conn, mid)
    by = {s["key"]: s for s in stages}
    plan_keys = ["charter", "atp", "tech", "risk", "plan"]
    report_keys = plan_keys + ["explore", "robust", "control", "report"]
    keys = plan_keys if kind == "plan" else report_keys
    included = [by[k] for k in keys if k in by]
    unapproved = [s for s in included if s["status"] not in workflow.LOCKED or s["stale"]]
    sig = {}          # stage id -> (ok, problems)
    intact = {}       # review id -> bool
    for s in included:
        ok_map, problems = workflow._signature_eval(conn, mid, s)
        intact.update(ok_map)
        sig[s["id"]] = (not problems, problems)
    bad_sig = [s for s in included if s["status"] in workflow.LOCKED and not sig[s["id"]][0]]
    pending = []      # template pre-fill / example values nobody confirmed: such a document is never an "approved version"
    for s in included:
        for k in (s["data"].get("unconfirmed") or []):
            title = next((x["title"] for x in s["def"]["sections"] if x["key"] == k), k)
            pending.append("%s › %s（模板预置内容，尚未确认）" % (s["def"]["title"], title))
        nmark = workflow._marked_items(s["data"].get("sections", {}))
        if nmark:
            pending.append("%s：%d 处仍含“%s”标记" % (s["def"]["title"], nmark, workflow.SAMPLE_WORD))
    nexp = sum(1 for e in exps if e.get("status") != "cancelled" and workflow.SAMPLE_WORD in (e.get("acceptance") or ""))
    if nexp:
        pending.append("实验接受标准：%d 个仍含“%s”标记" % (nexp, workflow.SAMPLE_WORD))
    gaps = []         # required sections nobody filled (allowed by "progressive completion"): disclosed on the cover
    for s in included:
        unc = set(s["data"].get("unconfirmed") or [])
        for sec in s["def"]["sections"]:
            if sec.get("required") and sec["key"] not in unc and not workflow._section_filled(sec, s["data"].get("sections", {}).get(sec["key"])):
                gaps.append("%s › %s" % (s["def"]["title"], sec["title"]))
    final = not unapproved and not bad_sig and not pending
    n = conn.execute("SELECT COUNT(*) FROM generated_docs WHERE method_id=? AND kind=? AND format='docx'", (mid, kind)).fetchone()[0] + 1
    status_label = INTEGRITY_FAIL_LABEL if bad_sig else (FINAL_LABEL if final else (PENDING_LABEL if (not unapproved and pending) else DRAFT_LABEL))
    doc_hash = db.sha256(db.canon({"stages": [[s["key"], s["revision"], s["data"]] for s in included],
                                  "experiments": [[e["code"], e["version"], e["status"]] for e in exps]}))
    reviews, stage_revs, submissions = {}, {}, {}
    for s in included:
        if s["def"].get("gate"):
            reviews[s["id"]] = [dict(r) for r in conn.execute(
                "SELECT r.*, u.display_name FROM reviews r JOIN users u ON u.id=r.user_id WHERE r.stage_id=? ORDER BY r.id", (s["id"],))]
        rows = [dict(r) for r in conn.execute(
            "SELECT sr.revision, sr.created_at, sr.reason, u.display_name FROM stage_revisions sr LEFT JOIN users u ON u.id=sr.created_by "
            "WHERE sr.stage_id=? ORDER BY sr.id", (s["id"],))]
        stage_revs[s["id"]] = rows
        if rows:
            submissions[s["id"]] = rows[-1]
    audit_changes = []
    for r in conn.execute("SELECT a.ts, a.action, a.details, a.reason, a.username, u.display_name FROM audit_log a LEFT JOIN users u ON u.id=a.user_id "
                          "WHERE a.method_id=? AND a.action IN ('stage.reopen','stage.withdraw') ORDER BY a.id", (mid,)):
        d = dict(r)
        d["details"] = db.jl(d["details"], {}) or {}
        audit_changes.append(d)
    prev_docs = [dict(r) for r in conn.execute(
        "SELECT d.created_at, d.version_label, d.status_label, u.display_name FROM generated_docs d LEFT JOIN users u ON u.id=d.created_by "
        "WHERE d.method_id=? AND d.kind=? AND d.format='docx' ORDER BY d.id", (mid, kind))]
    atts = [dict(r) for r in conn.execute(
        "SELECT a.*, u.display_name FROM attachments a LEFT JOIN users u ON u.id=a.uploaded_by WHERE a.method_id=? AND a.deleted=0 ORDER BY a.id", (mid,))]
    # cover: prepared by (last submission of the document's gate stage), reviewed / approved by (e-signatures of that stage)
    gate_stage = by.get("plan" if kind == "plan" else "report")
    people, effective = {}, ""
    if gate_stage:
        sub = submissions.get(gate_stage["id"])
        if sub:
            people["author"] = {"name": sub["display_name"] or "", "date": _local_date(sub["created_at"]), "note": "提交 R%d" % sub["revision"]}
        for role in ("reviewer", "approver"):
            rec = [x for x in reviews.get(gate_stage["id"], []) if x["role"] == role and x["decision"] == "approve" and x["revision"] == gate_stage["revision"]]
            if rec:
                x = rec[-1]
                people[role] = {"name": x["display_name"] or "", "date": _local_date(x["signed_at"]),
                                "note": "电子签名 R%d　哈希 %s%s" % (x["revision"], x["data_hash"][:12], "" if intact.get(x["id"]) else "　（校验失败）")}
                effective = _local_date(x["signed_at"])
    try:
        tpl = templates_data.build_template(p["template_key"])
        structure = tpl["doc_structure"]
        tpl_ch = {(st["key"], sec["key"]): sec.get("ch") for st in tpl["stages"] for sec in st["sections"]}
    except KeyError:
        structure, tpl_ch = templates_data.doc_structure(), {}
    doc_title = "分析方法开发方案" if kind == "plan" else "分析方法开发报告"
    version = "V%d" % n
    doc_no = "%s%s-%s-%s" % (settings["doc_prefix"], p["code"], "PLAN" if kind == "plan" else "RPT", version)
    return {"p": p, "project": project, "kind": kind, "stages": stages, "exps": exps, "included": included, "unapproved": unapproved,
            "bad_sig": bad_sig, "sig": sig, "intact": intact, "final": final, "n": n, "version": version, "pending": pending, "gaps": gaps,
            "status_label": T(status_label), "status_short": T("完整性校验失败" if bad_sig else ("已批准" if final else ("待确认" if (not unapproved and pending) else "草稿"))), "doc_hash": doc_hash, "lang": lang,
            "reviews": reviews, "stage_revs": stage_revs, "submissions": submissions, "audit_changes": audit_changes, "prev_docs": prev_docs, "atts": atts,
            "cover_people": people, "effective": effective, "settings": settings, "structure": structure, "tpl_ch": tpl_ch,
            "doc_title": T(doc_title), "doc_no": doc_no, "generated_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "soffice": settings["soffice_path"],
            "user": {"id": user["id"], "display_name": user["display_name"]}, "change_sig": _change_signature(conn, mid)}


# ====================================================================== render (phase 2)
def _set_update_fields(doc):
    _ins(doc.settings.element, _el("w:updateFields", w_val="true"), _SETTINGS_ORDER)


def build_document(ctx, pages=None):
    """Build the python-docx Document.  `pages` {bookmark: page} fills the contents list (second pass); without it
    the document asks Word to refresh the fields (contents page numbers) when it is opened."""
    r = R(ctx)
    doc = r.doc
    cp = doc.core_properties
    cp.title = "%s %s" % (ctx["doc_title"], ctx["p"]["name"])
    cp.subject = "%s %s" % (ctx["doc_no"], ctx["status_label"])
    cp.author = ctx["user"]["display_name"]
    cp.comments = "由 MethodFlow 生成；内容数据哈希 %s" % ctx["doc_hash"]
    units, toc = plan_outline(ctx)
    if ctx.get("lang") == "en":                # headings and the contents list first, so that the PDF page lookup matches
        for u in units:
            if isinstance(u.get("text"), str):
                u["text"] = i18n_doc.translate(u["text"])
        toc = [(lv, i18n_doc.translate(t), b) for lv, t, b in toc]
        cp.title = i18n_doc.translate(cp.title)
        cp.comments = "Generated by MethodFlow; content data hash %s" % ctx["doc_hash"]
    ctx["_toc"] = toc
    _cover(r)
    _control_page(r)
    _contents(r, toc, pages)
    if units:
        units[0]["pb"] = True              # the body starts on a new page
    for u in units:
        r.set_orient(bool(u["wide"]))
        k = u["u"]
        if k in ("h1", "h2"):
            r.heading(int(k[1]), u["text"], u["bm"], u.get("pb", False))
        elif k == "section":
            render_section(r, u["sec"], _stage_sections(u["s"]).get(u["sec"]["key"]), u["s"])
        elif k == "auto":
            _AUTO[u["name"]](r)
        elif k == "gantt":
            _auto_gantt(r)
        elif k == "app_check":
            _app_check(r, u["s"])
        elif k == "app_atts":
            _app_atts(r)
        elif k == "app_refdocs":
            _app_refdocs(r)
        elif k == "app_checkrefs":
            _app_checkrefs(r)
        elif k == "app_orphans":
            _app_orphans(r, u["items"])
    if pages is None:
        _set_update_fields(doc)
    if ctx.get("lang") == "en":
        i18n_doc.localize(doc)
    return doc


def render(ctx):
    """First pass: the .docx file.  Returns (path, filename)."""
    p, kind = ctx["p"], ctx["kind"]
    out_dir = config.DOCS / p["code"]
    out_dir.mkdir(parents=True, exist_ok=True)
    fname = "%s_%s_%s_%s.docx" % (p["code"], ("Plan" if kind == "plan" else "Report") if ctx.get("lang") == "en" else ("方案" if kind == "plan" else "报告"), ctx["version"], dt.datetime.now().strftime("%Y%m%d%H%M%S"))
    path = out_dir / fname
    build_document(ctx, None).save(str(path))
    return path, fname


def _make_pdf(ctx, docx_path):
    """(pdf path or None, note).  Never raises."""
    if not config.get("pdf_enabled"):
        return None, PDF_OFF_NOTE
    soffice = find_soffice(ctx["soffice"])
    if not soffice:
        return None, NO_SOFFICE_NOTE
    pdf, note = convert_pdf(soffice, docx_path)
    if pdf is None:
        return None, note
    try:
        pages = map_pages(ctx["_toc"], pdf_heading_pages(pdf))
        if pages:             # second pass: contents list with real page numbers (layout is identical), then the final PDF
            build_document(ctx, pages).save(str(docx_path))
            pdf, note = convert_pdf(soffice, docx_path)
    except Exception:
        pass
    return pdf, note


# ====================================================================== record (phase 3)
def _file_sha(path):
    h = hashlib.sha256()
    with open(str(path), "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _commit(conn, user, mid, ctx, files):
    """files: [(format, path, filename)] - the .docx first.  Returns (gen_id, [{id, format, filename, sha256}])."""
    if _change_signature(conn, mid) != ctx["change_sig"]:
        raise workflow.WFError("生成文档期间方法数据已被修改，请重新生成", 409)
    n = conn.execute("SELECT COUNT(*) FROM generated_docs WHERE method_id=? AND kind=? AND format='docx'", (mid, ctx["kind"])).fetchone()[0] + 1
    if n != ctx["n"]:
        raise workflow.WFError("同时有另一份同类文档在生成，请稍后重试", 409)
    _check_rate(conn, mid)
    base = {"unapproved": [s["key"] for s in ctx["unapproved"]], "final": ctx["final"], "signatures_ok": not ctx["bad_sig"],
            "bad_signatures": [s["key"] for s in ctx["bad_sig"]], "doc_no": ctx["doc_no"]}
    out, gen_id = [], None
    for fmt, path, fname in files:
        summ = dict(base, file_sha256=_file_sha(path), size=os.path.getsize(str(path)), gen_id=gen_id)
        cur = conn.execute("INSERT INTO generated_docs(method_id,kind,format,filename,path,version_label,status_label,data_hash,created_by,created_at,summary) "
                           "VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                           (mid, ctx["kind"], fmt, fname, str(path.relative_to(config.DATA)), ctx["version"], ctx["status_label"], ctx["doc_hash"], user["id"],
                            db.now(), json.dumps(summ, ensure_ascii=False)))
        if gen_id is None:
            gen_id = cur.lastrowid
            summ["gen_id"] = gen_id
            conn.execute("UPDATE generated_docs SET summary=? WHERE id=?", (json.dumps(summ, ensure_ascii=False), gen_id))
        out.append({"id": cur.lastrowid, "format": fmt, "filename": fname, "sha256": summ["file_sha256"]})
    audit.log(conn, user, "document.generate", "document", gen_id, mid,
              {"kind": ctx["kind"], "version": ctx["version"], "doc_no": ctx["doc_no"], "status": ctx["status_label"], "hash": ctx["doc_hash"],
               "filename": files[0][2], "formats": {o["format"]: o["sha256"] for o in out}, "signatures_ok": not ctx["bad_sig"]})
    return gen_id, out


def generate(user, mid, kind="plan", lang="zh"):
    """Word + PDF in one go.  Returns {gen_id, docs:[{id, format, filename}], pdf_available, pdf_note}."""
    with db.read() as conn:
        workflow.require_view(conn, user, mid)
        if lang not in ("zh", "en"):
            raise workflow.WFError("文档语言须为 zh 或 en")
        ctx = collect(conn, user, mid, kind, lang)
    path, fname = render(ctx)            # no database connection / lock is held while rendering
    created = [path, path.with_suffix(".pdf")]
    try:
        pdf, note = _make_pdf(ctx, path)
        files = [("docx", path, fname)]
        if pdf is not None:
            files.append(("pdf", pdf, pdf.name))
        with db.tx() as conn:
            gen_id, out = _commit(conn, user, mid, ctx, files)
    except BaseException:
        for fp in created:
            try:
                os.remove(str(fp))
            except OSError:
                pass
        raise
    return {"gen_id": gen_id, "docs": [{"id": o["id"], "format": o["format"], "filename": o["filename"]} for o in out],
            "pdf_available": pdf is not None, "pdf_note": note}
